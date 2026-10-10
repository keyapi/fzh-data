"""Write the private monthly workbook. Business rows stay outside Git."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sellfox_settlement.validation_paths import private_output


def _cell(value):
    if isinstance(value, (dict, list)):
        value = str(value)
    if isinstance(value, str) and value[:1] in "=+@":
        return "'" + value
    return value


def write_workbook(tables, path):
    path = private_output(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    header_fill = PatternFill("solid", fgColor="22384D")
    note_fill = PatternFill("solid", fgColor="E9F1F7")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    body_font = Font(name="Arial", size=10)
    for source in tables["sheets"]:
        sheet = book.create_sheet(source["name"][:31])
        cols = len(source["headers"])
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
        note = sheet.cell(1, 1, _cell(source["note"]))
        note.alignment = Alignment(wrap_text=True, vertical="center")
        note.fill = note_fill
        note.font = Font(name="Arial", size=11, color="22384D")
        sheet.row_dimensions[1].height = 48
        for col, header in enumerate(source["headers"], start=1):
            cell = sheet.cell(3, col, _cell(header))
            cell.fill = header_fill
            cell.font = header_font
        for row_index, row in enumerate(source["rows"], start=4):
            for col, value in enumerate(row, start=1):
                cell = sheet.cell(row_index, col, _cell(value))
                cell.font = body_font
        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = f"A3:{get_column_letter(cols)}{max(3, 3 + len(source['rows']))}"
        for col in range(1, cols + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 42 if col == 1 else 22
    book.save(path)
    return path
