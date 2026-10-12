"""Verify one private run, its source hashes and optionally its exported workbook.

Verification proves artifact integrity, not finance approval or production sync.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from sellfox_settlement.validation_paths import private_output


REQUIRED = {'technical_month_report.json', 'monthly_validation.json',
            'source_manifest.json', 'workbook_tables.json'}


def load_json(path):
    return json.loads(path.read_text(encoding='utf-8'))


def digest(path):
    hasher = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            hasher.update(chunk)
    return hasher.hexdigest()


def contained_path(root, name):
    path = (root / name).resolve()
    if not path.is_relative_to(root.resolve()):
        raise ValueError('artifact path escapes report directory')
    return path


def same_cell(actual, expected):
    if isinstance(expected, (dict, list)):
        expected = json.dumps(expected, ensure_ascii=False, separators=(',', ':'))
    if isinstance(expected, str) and expected[:1] in {'=', '+', '@'}:
        expected = "'" + expected
    if expected is None or expected == '':
        return actual is None or actual == ''
    if isinstance(expected, bool) or isinstance(actual, bool):
        return type(actual) is type(expected) and actual == expected
    return actual == expected


def verify_workbook(path, tables):
    from openpyxl import load_workbook

    # Some valid exporters omit worksheet dimension metadata. Normal mode
    # computes the actual bounds instead of treating absent metadata as failure.
    book = load_workbook(path, read_only=False, data_only=False)
    try:
        if book.sheetnames != [s['name'] for s in tables['sheets']]:
            raise ValueError('workbook sheet identity mismatch')
        cells = 0
        for source in tables['sheets']:
            sheet = book[source['name']]
            if sheet.max_row != 3 + len(source['rows']) or sheet.max_column != len(source['headers']):
                raise ValueError('workbook dimensions mismatch')
            if not same_cell(sheet.cell(1, 1).value, source['note']):
                raise ValueError('workbook evidence note mismatch')
            expected_rows = [source['headers'], *source['rows']]
            for row_number, (actual_row, expected_row) in enumerate(zip(
                    sheet.iter_rows(min_row=3, max_row=sheet.max_row), expected_rows, strict=True), 3):
                if len(expected_row) != len(source['headers']):
                    raise ValueError('table row width mismatch')
                for column, (cell, expected) in enumerate(zip(actual_row, expected_row, strict=True), 1):
                    if cell.data_type in {'f', 'e'} or not same_cell(cell.value, expected):
                        raise ValueError(f'workbook cell mismatch: {source["name"]} row {row_number} column {column}')
                    cells += 1
        return {'sheets': len(tables['sheets']), 'verified_cells': cells}
    finally:
        book.close()


def verify_run(root, *, require_workbook=False, register_workbook=False):
    root = private_output(root).resolve()
    run = load_json(root / 'run_status.json')
    if run.get('status') != 'succeeded' or not run.get('run_id'):
        raise ValueError('current run has not succeeded')
    artifacts = run.get('artifacts') or []
    names = [a['path'] for a in artifacts]
    if len(set(names)) != len(names) or not REQUIRED.issubset(names):
        raise ValueError('required artifacts missing or duplicated')
    for artifact in artifacts:
        if digest(contained_path(root, artifact['path'])) != artifact['sha256']:
            raise ValueError('artifact hash mismatch')
    report = load_json(root / 'technical_month_report.json')
    tables = load_json(root / 'workbook_tables.json')
    monthly = load_json(root / 'monthly_validation.json')
    if (report.get('run_id') != run['run_id'] or report.get('status') != 'succeeded'
            or any(item.get('month') != run['month'] for item in [report, tables, monthly['summary']])
            or monthly['summary']['fatal_validation_errors']):
        raise ValueError('run identity, month or monthly validation mismatch')
    inputs = [*load_json(root / 'source_manifest.json')['files'], *report.get('snapshot_manifest', [])]
    if not inputs:
        raise ValueError('input evidence manifest missing')
    for source in inputs:
        if digest(Path(source['path'])) != source['sha256']:
            raise ValueError('input hash mismatch')
    result = {'status': 'verified_artifacts', 'run_id': run['run_id'], 'month': run['month'],
              'verified_artifacts': len(artifacts), 'verified_inputs': len(inputs),
              'workbook_status': 'not_requested'}
    if require_workbook or register_workbook:
        name = f'{run["month"]}-technical-workbook.xlsx'
        workbook = contained_path(root, name)
        table_hash = digest(root / 'workbook_tables.json')
        exports = [e for e in run.get('external_exports', []) if e.get('path') == name]
        direct = [a for a in artifacts if a['path'] == name]
        if not register_workbook and not direct:
            if (len(exports) != 1 or exports[0].get('source_tables_sha256') != table_hash
                    or exports[0].get('sha256') != digest(workbook)):
                raise ValueError('workbook is not a matching current export')
        result.update(verify_workbook(workbook, tables), workbook_status='verified')
        if register_workbook:
            run['external_exports'] = [{'path': name, 'sha256': digest(workbook),
                                        'source_tables_sha256': table_hash,
                                        'verification': 'workbook_verification.json'}]
            (root / 'workbook_verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
            (root / 'run_status.json').write_text(json.dumps(run, ensure_ascii=False, indent=2), encoding='utf-8')
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--reports', type=Path, required=True)
    parser.add_argument('--require-workbook', action='store_true')
    parser.add_argument('--register-workbook', action='store_true', help='Verify cells before registering an external export')
    args = parser.parse_args()
    try:
        result = verify_run(args.reports, require_workbook=args.require_workbook, register_workbook=args.register_workbook)
    except (OSError, ValueError, KeyError) as exc:
        parser.error(str(exc))
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    main()
