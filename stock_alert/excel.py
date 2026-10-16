# -*- coding: utf-8 -*-
"""低库存明细 Excel：**按仓库分 sheet**，每行一个「可用库存 < 阈值」的 SKU。"""

from __future__ import annotations

from pathlib import Path

from .scan import ScanResult

# Excel sheet 名限制：≤31 字符，禁用 : \ / ? * [ ]
_INVALID = set(':\\/?*[]')


def _sheet_name(name: str) -> str:
    cleaned = "".join(c for c in (name or "") if c not in _INVALID)
    return (cleaned or "Sheet")[:31]


def _num(value: float):
    return int(value) if float(value).is_integer() else value


def write_low_stock_workbook(result: ScanResult, path: str | Path) -> Path:
    """写明细 Excel；返回落盘路径。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)  # 去掉默认空 sheet

    for wh in result.warehouses:
        ws = wb.create_sheet(_sheet_name(wh.warehouse))
        ws.append(["SKU", "可用库存", "在途库存"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for item in wh.low:
            ws.append([item.sku, _num(item.available), _num(item.intransit)])
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 12

    if not wb.sheetnames:
        wb.create_sheet("空")

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
