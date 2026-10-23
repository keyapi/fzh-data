from __future__ import annotations

from openpyxl import load_workbook

from stock_alert.client import StockRow
from stock_alert.excel import write_low_stock_workbook
from stock_alert.scan import scan


def test_workbook_has_one_sheet_per_warehouse(tmp_path):
    rows = {
        "美中-FZH-DANEEY": [
            StockRow("美中-FZH-DANEEY", "A", 0, 3),
            StockRow("美中-FZH-DANEEY", "B", 9, 0),  # 不低于 5，不写入
        ],
        "美东-CENTRADE": [StockRow("美东-CENTRADE", "C", 1, 2)],
    }
    result = scan(rows, threshold=5)
    path = write_low_stock_workbook(result, tmp_path / "out.xlsx")

    wb = load_workbook(path)
    assert wb.sheetnames == ["美中-FZH-DANEEY", "美东-CENTRADE"]

    ws = wb["美中-FZH-DANEEY"]
    assert [c.value for c in ws[1]] == ["SKU", "可用库存", "在途库存"]
    data = [tuple(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    assert data == [("A", 0, 3)]

    ws2 = wb["美东-CENTRADE"]
    assert [tuple(r) for r in ws2.iter_rows(min_row=2, values_only=True)] == [("C", 1, 2)]


def test_workbook_sanitizes_sheet_name(tmp_path):
    rows = {r"bad:name[1]": [StockRow("bad:name[1]", "A", 0, 0)]}
    result = scan(rows, threshold=5)
    path = write_low_stock_workbook(result, tmp_path / "o.xlsx")
    wb = load_workbook(path)
    assert "badname1" in wb.sheetnames
