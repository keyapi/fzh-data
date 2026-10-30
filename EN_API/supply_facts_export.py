#!/usr/bin/env python3
"""Batch machine entry for the supply facts contract (P2).

Calls `key_test.supply_facts.get_supply_facts` on the EN test/prod site and
writes a versioned, reproducible snapshot: JSON datasets + Excel + manifest
with file checksums. Row-level data stays local; nothing is pushed.

Examples:
  uv run python EN_API/supply_facts_export.py --codes "TT0031038K0062927,CENCord995-DKGrey-97" --test
  uv run python EN_API/supply_facts_export.py --so SO-26-00026 --test
  uv run python EN_API/supply_facts_export.py --all --test --page-size 500
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import sys
import time
from datetime import datetime
from pathlib import Path

EN_API = Path(__file__).resolve().parent
sys.path.insert(0, str(EN_API))
sys.stdout.reconfigure(encoding="utf-8")

import item_shipment_status as iss  # noqa: E402

ROOT = EN_API.parent
sys.path.insert(0, str(ROOT))
from supply_chain_visibility.receipt_bridge import (  # noqa: E402
    SnapshotRow, build_inventory_bridge,
)
from supply_chain_visibility.receipt_reconciliation import (  # noqa: E402
    ReceiptEvent, SnapshotWindow,
)
from supply_chain_visibility.receipt_state import (  # noqa: E402
    ReceiptLine, ShipmentLink, reconcile_signed_receipts,
)

METHOD = "/api/method/key_test.supply_facts.get_supply_facts"
METHOD_DELIVERY = "/api/method/key_test.delivery_facts.get_delivery_facts"
DATASETS = ("lines", "stage_balances", "milestones", "availability_schedule",
            "links", "production_evidence", "shipment_progress",
            "receipt_reconciliation", "inventory_bridge")
DELIVERY_DATASETS = ("delivery_lines", "association_edges",
                     "coverage_exceptions", "so_fulfillment", "supply_ledger",
                     "physical_chain")
EVIDENCE_CLASSES = ("first_process_done", "all_process_done", "qc_performed",
                    "qc_passed", "stock_in_posted")
DEFAULT_OUT = EN_API / "out" / "supply_facts"
SCHEMA_VERSION = "supply-facts/v3"
RECEIPT_STATE_SCHEMA = "supply-facts/receipt-state/v1"
INVENTORY_BRIDGE_SCHEMA = "supply-facts/inventory-bridge/v1"


def _call(base_url, key, secret, filters):
    import urllib.parse
    path = METHOD + "?" + urllib.parse.urlencode(
        {"filters": json.dumps(filters, ensure_ascii=False)})
    status, body = iss.api_get_status(path, base_url, key, secret, timeout=300)
    if status != 200:
        raise iss.QueryError("get_supply_facts HTTP %s %s" % (status, str(body)[:200]))
    message = body.get("message")
    if not isinstance(message, dict):
        raise iss.QueryError("get_supply_facts returned %s" % type(message).__name__)
    return message


def _call_delivery(base_url, key, secret, filters):
    import urllib.parse
    path = METHOD_DELIVERY + "?" + urllib.parse.urlencode(
        {"filters": json.dumps(filters, ensure_ascii=False)})
    status, body = iss.api_get_status(path, base_url, key, secret, timeout=300)
    if status != 200:
        raise iss.QueryError("get_delivery_facts HTTP %s %s"
                             % (status, str(body)[:200]))
    message = body.get("message")
    if not isinstance(message, dict):
        raise iss.QueryError("get_delivery_facts returned %s" % type(message).__name__)
    return message


MAX_PAGES = 200


def _paginate(base_url, key, secret, base_filters, page_size, max_pages=MAX_PAGES):
    """Fetch all pages for one scope; validate continuity, totals and exhaustion."""
    offset = int(base_filters.get("offset") or 0)
    page = 0
    batches = []
    seen = set()
    expected_total = None
    while page < max_pages:
        filters = dict(base_filters, limit=page_size, offset=offset)
        batch = _call(base_url, key, secret, filters)
        lines = batch.get("lines") or []
        for line in lines:
            line_key = line.get("so_line") or line.get("so_line_id")
            if not line_key:
                raise iss.QueryError("line without so_line at offset %s" % offset)
            if line_key in seen:
                raise iss.QueryError("duplicate so_line across pages: %s" % line_key)
            seen.add(line_key)
        pagination = batch.get("pagination") or {}
        if pagination:
            if "limit" in pagination and int(pagination["limit"]) != page_size:
                raise iss.QueryError(
                    "pagination limit mismatch: %s != %s" % (pagination["limit"], page_size))
            if "offset" in pagination and int(pagination["offset"]) != offset:
                raise iss.QueryError(
                    "pagination offset mismatch: %s != %s" % (pagination["offset"], offset))
            if "returned" in pagination and int(pagination["returned"]) != len(lines):
                raise iss.QueryError(
                    "pagination returned mismatch: %s != %s"
                    % (pagination["returned"], len(lines)))
            if "total" in pagination:
                total = int(pagination["total"])
                if expected_total is None:
                    expected_total = total
                elif total != expected_total:
                    raise iss.QueryError(
                        "pagination total changed: %s != %s" % (total, expected_total))
        batches.append(batch)
        exhausted = pagination.get("exhausted")
        if exhausted is True:
            if expected_total is not None and len(seen) != expected_total:
                raise iss.QueryError(
                    "exhausted with %s unique lines, expected %s" % (len(seen), expected_total))
            return batches, "exhausted"
        if exhausted is False:
            next_offset = pagination.get("next_offset")
            expected_next = offset + len(lines)
            if next_offset is None or int(next_offset) != expected_next:
                raise iss.QueryError(
                    "pagination gap: offset=%s rows=%s next=%s expected_next=%s"
                    % (offset, len(lines), next_offset, expected_next))
            if not lines:
                raise iss.QueryError("empty page before exhausted at offset %s" % offset)
            offset = int(next_offset)
            page += 1
            continue
        # Backward compatibility with a server that predates pagination metadata.
        if len(lines) < page_size:
            if expected_total is None:
                coverage = (batch.get("manifest") or {}).get("coverage") or {}
                if coverage.get("so_lines_total") is not None:
                    expected_total = int(coverage["so_lines_total"])
            if expected_total is not None and len(seen) != expected_total:
                raise iss.QueryError(
                    "exhausted with %s unique lines, expected %s" % (len(seen), expected_total))
            return batches, "exhausted"
        if not lines:
            raise iss.QueryError("empty page before exhausted at offset %s" % offset)
        offset += len(lines)
        page += 1
    return batches, "page_limit"


def _paginate_delivery(base_url, key, secret, page_size, *,
                       include_inactive=False, base_filters=None,
                       max_pages=MAX_PAGES):
    offset = 0
    page = 0
    batches = []
    seen = set()
    expected_total = None
    while page < max_pages:
        filters = dict(base_filters or {})
        filters.update({"limit": page_size, "offset": offset,
                        "include_unlinked": 1})
        if include_inactive:
            filters["include_inactive"] = 1
        batch = _call_delivery(base_url, key, secret, filters)
        lines = batch.get("delivery_lines") or []
        for line in lines:
            key_value = line.get("delivery_note_item")
            if not key_value:
                raise iss.QueryError("delivery line without identity at offset %s"
                                     % offset)
            if key_value in seen:
                raise iss.QueryError("duplicate delivery_note_item across pages: %s"
                                     % key_value)
            seen.add(key_value)
        pagination = batch.get("pagination") or {}
        if int(pagination.get("limit", page_size)) != page_size:
            raise iss.QueryError("delivery pagination limit mismatch")
        if int(pagination.get("offset", offset)) != offset:
            raise iss.QueryError("delivery pagination offset mismatch")
        if int(pagination.get("returned", len(lines))) != len(lines):
            raise iss.QueryError("delivery pagination returned mismatch")
        total = int(pagination.get("total", 0))
        if expected_total is None:
            expected_total = total
        elif total != expected_total:
            raise iss.QueryError("delivery pagination total changed: %s != %s"
                                 % (total, expected_total))
        batches.append(batch)
        if pagination.get("exhausted") is True:
            if len(seen) != expected_total:
                raise iss.QueryError("delivery exhausted with %s rows, expected %s"
                                     % (len(seen), expected_total))
            return batches, "exhausted"
        next_offset = pagination.get("next_offset")
        if next_offset is None or int(next_offset) != offset + len(lines):
            raise iss.QueryError("delivery pagination gap at offset %s" % offset)
        if not lines:
            raise iss.QueryError("empty delivery page before exhausted")
        offset = int(next_offset)
        page += 1
    return batches, "page_limit"


class BatchVersionError(Exception):
    """Batch manifest is missing or conflicts with already merged data."""


VERSION_FIELDS = ("schema_version", "rule_version", "implementation_version", "as_of")


def _check_batch_version(target, manifest):
    incoming = {
        "schema_version": (manifest.get("schema_version") or "").strip(),
        "rule_version": (manifest.get("rule_version") or "").strip(),
        "implementation_version": (manifest.get("implementation_version") or "").strip(),
        "as_of": (manifest.get("as_of") or "")[:10],
    }
    missing = [key for key, value in incoming.items() if not value]
    if missing:
        raise BatchVersionError("批次缺少版本/快照字段: " + ", ".join(missing))
    for key, value in incoming.items():
        current = (target.get(key) or "")[:10] if key == "as_of" else (target.get(key) or "")
        if not current:
            target[key] = value
        elif current != value:
            raise BatchVersionError(
                "批次 %s 不一致: 收到 %s，已合并 %s" % (key, value, current))


def _coverage_state(target):
    state = target.get("_coverage_state")
    if state is None:
        state = {"so_ids": set(), "so_line_ids": set(), "dn_ids": set(),
                 "shipment_ids": set(), "scope_totals": {}}
        target["_coverage_state"] = state
    return state


def _merge(target, batch, scope_key=None):
    manifest = batch.get("manifest") or {}
    _check_batch_version(target, manifest)
    state = _coverage_state(target)
    for line in batch.get("lines") or []:
        if line.get("so"):
            state["so_ids"].add(line["so"])
        line_key = line.get("so_line") or line.get("so_line_id")
        if line_key:
            state["so_line_ids"].add(line_key)
    for link in batch.get("links") or []:
        for side in ("source", "target"):
            doctype = link.get(side + "_doctype")
            name = link.get(side + "_name")
            if not name:
                continue
            if doctype == "Delivery Note":
                state["dn_ids"].add(name)
            elif doctype == "Shipment" and link.get("target_is_record") is True:
                state["shipment_ids"].add(name)
    pagination = batch.get("pagination") or {}
    total = pagination.get("total")
    if total is None:
        total = (manifest.get("coverage") or {}).get("so_lines_total")
    if total is not None:
        key = scope_key or (manifest.get("scope") or {}).get("sales_order") or "__all__"
        state["scope_totals"][key] = int(total)
    for key in DATASETS:
        target[key].extend(batch.get(key) or [])
    target["partial"] = target["partial"] or bool(manifest.get("partial"))
    for key, value in (manifest.get("counts") or {}).items():
        target["counts"][key] = target["counts"].get(key, 0) + int(value or 0)
    for key, value in (manifest.get("coverage") or {}).items():
        if key in ("sales_orders", "so_lines", "so_lines_total", "delivery_notes", "shipments"):
            # These are recomputed by identity in _finalize_coverage; never
            # sum per-page counts (a 501-line SO would become 1002).
            continue
        if key == "sources":
            for name, ok in (value or {}).items():
                target["sources"][name] = bool(target["sources"].get(name, True)) and bool(ok)
        elif key == "restricted":
            for name, flag in (value or {}).items():
                target["restricted"][name] = bool(target["restricted"].get(name)) or bool(flag)
        elif key == "external_sources":
            for name, state in (value or {}).items():
                old = target["external_sources"].get(name)
                if old == "not_integrated":
                    continue
                if state != "integrated":
                    target["external_sources"][name] = state or "not_integrated"
                else:
                    target["external_sources"][name] = "integrated"
        elif key == "source_timestamps":
            for name, stamp in (value or {}).items():
                old = target["source_timestamps"].get(name)
                if old is None or (stamp and stamp < old):
                    target["source_timestamps"][name] = stamp
        elif isinstance(value, int):
            target["coverage"][key] = target["coverage"].get(key, 0) + value
        elif key == "empty_reason" and value:
            target["coverage"]["empty_reason"] = value
        elif key == "source_status":
            merged = target["coverage"].setdefault("source_status", {})
            for name, state in (value or {}).items():
                current = merged.get(name)
                if current is None:
                    merged[name] = dict(state)
                elif current.get("available", True) and not state.get("available", True):
                    merged[name] = dict(state)
    for note in manifest.get("notes") or []:
        if note not in target["notes"]:
            target["notes"].append(note)


def _finalize_coverage(target):
    """Replace per-page sums with identity-deduped/range totals."""
    state = target.pop("_coverage_state", None)
    if not state:
        return
    coverage = target.setdefault("coverage", {})
    coverage["sales_orders"] = len(state["so_ids"])
    coverage["so_lines"] = len(state["so_line_ids"])
    coverage["so_lines_total"] = (
        sum(state["scope_totals"].values())
        if state["scope_totals"] else len(state["so_line_ids"])
    )
    coverage["delivery_notes"] = len(state["dn_ids"])
    coverage["shipments"] = len(state["shipment_ids"])


def _resolve_sos(codes, base_url, key, secret):
    names = []
    for code in codes:
        found = [r["name"] for r in iss.find_sales_orders_by_items([code], base_url, key, secret)]
        if not found:
            found = [r["name"] for r in iss.find_sales_orders_by_customer_code(
                code, base_url, key, secret)]
        if not found:
            print("  ! 未命中销售订单: %s" % code, file=sys.stderr)
        names.extend(found)
    return sorted(set(names))


def _qty_status(evidence):
    status = evidence.get("qty_status")
    if status:
        return status
    confidence = evidence.get("confidence")
    if evidence.get("qty_proven") is not None and confidence in ("proven", "mixed"):
        return "known"
    if evidence.get("qty_estimated") is not None and confidence in ("estimated", "mixed"):
        return "estimated"
    return "unknown"


def _flatten_production_evidence(rows):
    """One stable row per SO line × WO × evidence class."""
    flat = []
    for line in rows or []:
        line_item = line.get("line_item_code")
        for wo in line.get("work_orders") or []:
            item = wo.get("production_item")
            role = ("subassembly" if line_item and item and item != line_item
                    else "primary")
            for evidence_class in EVIDENCE_CLASSES:
                evidence = wo.get(evidence_class) or {}
                components = evidence.get("components") or {}
                flat.append({
                    "so": line.get("so"),
                    "so_line_id": line.get("so_line"),
                    "line_item_code": line_item,
                    "production_item": item,
                    "wo": wo.get("wo"),
                    "wo_qty": wo.get("qty"),
                    "wo_docstatus": wo.get("docstatus"),
                    "wo_cancelled": bool(wo.get("wo_cancelled")),
                    "line_role": role,
                    "evidence_class": evidence_class,
                    "qty_proven": evidence.get("qty_proven"),
                    "qty_estimated": evidence.get("qty_estimated"),
                    "confidence": evidence.get("confidence"),
                    "qty_status": _qty_status(evidence),
                    "started_at": evidence.get("started_at"),
                    "completed_at": evidence.get("completed_at"),
                    "time_status": evidence.get("time_status"),
                    "observed_at": evidence.get("observed_at") or line.get("observed_at"),
                    "gross_qty": components.get("gross_qty"),
                    "return_qty": components.get("return_qty"),
                    "net_qty": components.get("net_qty"),
                    "source_records": json.dumps(evidence.get("source_records") or [],
                                                 ensure_ascii=False, default=str),
                    "reasons": json.dumps(evidence.get("reasons") or [],
                                          ensure_ascii=False, default=str),
                    "conflicts": json.dumps(evidence.get("conflicts") or [],
                                            ensure_ascii=False, default=str),
                })
    return flat


def _flatten_shipment_progress(rows):
    out = []
    for row in rows or []:
        item = dict(row)
        for key in ("conflict_reasons", "progress_candidates", "source_status",
                    "shipment_candidates"):
            if key in item and not isinstance(item[key], str):
                item[key] = json.dumps(item[key], ensure_ascii=False, default=str)
        out.append(item)
    return out


def _load_json(path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


class LocalInputError(Exception):
    """Local read-only input violates the strict schema/coverage contract."""


COVERAGE_KEYS = {"query_range_complete", "requested_windows", "truncated",
                 "not_collected", "source_as_of", "covered_po_codes",
                 "coverage_mode", "po_full_history_proven",
                 "po_full_history_source", "po_full_history_as_of"}
REQUIRED_COVERAGE_KEYS = {"query_range_complete", "requested_windows",
                          "truncated", "not_collected", "source_as_of",
                          "covered_po_codes"}
LINK_FIELDS = {"shipment", "dn", "po_codes", "signed_at", "qty", "uom", "sku",
               "warehouse", "identity_confidence", "status_semantics",
               "signed_status_proven"}
RECEIPT_LINE_FIELDS = {"receipt_id", "line_id", "po_code", "sku", "warehouse",
                       "uom", "qty", "effective_at", "quantity_semantics",
                       "uom_proven", "identity_proven"}
SNAPSHOT_FIELDS = {"inventory_snapshot_id", "sku", "warehouse", "uom",
                   "available", "as_of", "window_start", "window_end"}
RECEIPT_EVENT_FIELDS = {"receipt_id", "line_id", "po_code", "sku", "warehouse",
                        "uom", "qty", "receipt_effective_at", "uom_proven",
                        "quantity_semantics_proven", "po_line_id",
                        "dn_line_id", "shipment_id", "inclusion_proof"}
RECEIPT_EVENT_REQUIRED = RECEIPT_EVENT_FIELDS - {
    "po_line_id", "dn_line_id", "shipment_id", "inclusion_proof"}


def _check_keys(row, allowed, required, label, index):
    if not isinstance(row, dict):
        raise LocalInputError("%s[%s] must be an object" % (label, index))
    unknown = sorted(set(row) - allowed)
    if unknown:
        raise LocalInputError("%s[%s] unknown field(s): %s"
                              % (label, index, ", ".join(unknown)))
    missing = sorted(required - set(row))
    if missing:
        raise LocalInputError("%s[%s] missing field(s): %s"
                              % (label, index, ", ".join(missing)))


def _require_bool(value, label, name):
    if not isinstance(value, bool):
        raise LocalInputError("%s.%s must be a JSON boolean (true/false)" % (label, name))
    return value


def _require_aware_time(value, label, name, allow_none=False):
    if value is None and allow_none:
        return None
    if not isinstance(value, str) or not value.strip():
        raise LocalInputError("%s.%s must be an ISO-8601 string with timezone" % (label, name))
    try:
        parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
    except ValueError:
        raise LocalInputError("%s.%s is not a valid ISO-8601 time" % (label, name))
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise LocalInputError("%s.%s must carry an explicit timezone offset" % (label, name))
    return value


def _validate_window(value, label, name):
    _check_keys(value, {"start", "end"}, {"start", "end"}, label, name)
    _require_aware_time(value["start"], label + "." + name, "start")
    _require_aware_time(value["end"], label + "." + name, "end")


def _validate_coverage(value):
    if not isinstance(value, dict):
        raise LocalInputError("coverage must be an object")
    unknown = sorted(set(value) - COVERAGE_KEYS)
    if unknown:
        raise LocalInputError("coverage unknown field(s): %s" % ", ".join(unknown))
    missing = sorted(REQUIRED_COVERAGE_KEYS - set(value))
    if missing:
        raise LocalInputError("coverage missing field(s): %s" % ", ".join(missing))
    _require_bool(value["query_range_complete"], "coverage", "query_range_complete")
    _require_bool(value["truncated"], "coverage", "truncated")
    mode = value.get("coverage_mode") or "time_range"
    if mode not in ("time_range", "po_full_history"):
        raise LocalInputError("coverage.coverage_mode must be time_range or po_full_history")
    for name in ("requested_windows", "not_collected"):
        windows = value[name]
        if not isinstance(windows, list):
            raise LocalInputError("coverage.%s must be a list" % name)
        for index, window in enumerate(windows):
            _validate_window(window, "coverage", "%s[%s]" % (name, index))
    source_as_of = _require_aware_time(value["source_as_of"], "coverage",
                                       "source_as_of")
    source_dt = datetime.fromisoformat(source_as_of.strip().replace("Z", "+00:00"))
    codes = value["covered_po_codes"]
    if not isinstance(codes, list) or not all(
            isinstance(code, str) and code.strip() for code in codes):
        raise LocalInputError(
            "coverage.covered_po_codes must be a list of non-empty strings")
    if mode == "time_range":
        for name in ("po_full_history_proven", "po_full_history_source",
                     "po_full_history_as_of"):
            if value.get(name) is not None:
                raise LocalInputError(
                    "coverage.%s must not be used with coverage_mode=time_range"
                    % name)
        if not value["requested_windows"]:
            raise LocalInputError(
                "coverage.requested_windows must contain at least one window")
        for index, window in enumerate(value["requested_windows"]):
            start_dt = datetime.fromisoformat(
                window["start"].strip().replace("Z", "+00:00"))
            end_dt = datetime.fromisoformat(
                window["end"].strip().replace("Z", "+00:00"))
            if start_dt > end_dt:
                raise LocalInputError(
                    "coverage.requested_windows[%s] start is after end" % index)
            if end_dt > source_dt:
                raise LocalInputError(
                    "coverage.source_as_of is earlier than requested_windows[%s].end"
                    % index)
    else:
        if value.get("po_full_history_proven") is not True:
            raise LocalInputError(
                "coverage.po_full_history_proven must be JSON true")
        if not isinstance(value.get("po_full_history_source"), str) or not \
                value["po_full_history_source"].strip():
            raise LocalInputError(
                "coverage.po_full_history_source must be a non-empty string")
        _require_aware_time(value.get("po_full_history_as_of"), "coverage",
                            "po_full_history_as_of")
        history_dt = datetime.fromisoformat(
            value["po_full_history_as_of"].strip().replace("Z", "+00:00"))
        if history_dt != source_dt:
            raise LocalInputError(
                "coverage.po_full_history_as_of must equal coverage.source_as_of")


def _validate_duplicate_identity(rows, identity_fn, payload_fn, label):
    groups = {}
    for index, row in enumerate(rows):
        key = identity_fn(row)
        if key in (None, "", (None, None)):
            continue
        groups.setdefault(key, []).append((index, row))
    for key, entries in groups.items():
        payloads = {json.dumps(payload_fn(row), sort_keys=True, default=str)
                    for _, row in entries}
        if len(payloads) > 1:
            raise LocalInputError("%s duplicate identity %s has different payloads"
                                  % (label, key))


def _validate_link_row(row, index):
    label = "links[%s]" % index
    _check_keys(row, LINK_FIELDS, LINK_FIELDS - {"signed_status_proven"},
                "links", index)
    for name in ("shipment", "dn", "sku", "warehouse", "uom",
                 "identity_confidence", "status_semantics"):
        if not isinstance(row[name], str) or not row[name].strip():
            raise LocalInputError("%s.%s must be a non-empty string" % (label, name))
    if not isinstance(row["po_codes"], list) or not all(
            isinstance(code, str) for code in row["po_codes"]):
        raise LocalInputError("%s.po_codes must be a list of strings" % label)
    _require_aware_time(row["signed_at"], label, "signed_at", allow_none=True)
    if "signed_status_proven" in row:
        _require_bool(row["signed_status_proven"], label, "signed_status_proven")
    if row["qty"] is None or isinstance(row["qty"], bool):
        raise LocalInputError("%s.qty must be a number" % label)
    return row


def _validate_receipt_line(row, index):
    label = "receipts[%s]" % index
    _check_keys(row, RECEIPT_LINE_FIELDS, RECEIPT_LINE_FIELDS, "receipts", index)
    for name in ("receipt_id", "line_id", "po_code", "sku", "warehouse", "uom",
                 "quantity_semantics"):
        if not isinstance(row[name], str) or not row[name].strip():
            raise LocalInputError("%s.%s must be a non-empty string" % (label, name))
    _require_bool(row["uom_proven"], label, "uom_proven")
    _require_bool(row["identity_proven"], label, "identity_proven")
    _require_aware_time(row["effective_at"], label, "effective_at", allow_none=True)
    if row["qty"] is None or isinstance(row["qty"], bool):
        raise LocalInputError("%s.qty must be a number" % label)
    return row


def _validate_snapshot_row(row, index):
    label = "snapshots[%s]" % index
    _check_keys(row, SNAPSHOT_FIELDS, SNAPSHOT_FIELDS, "snapshots", index)
    for name in ("inventory_snapshot_id", "sku", "warehouse", "uom"):
        if not isinstance(row[name], str) or not row[name].strip():
            raise LocalInputError("%s.%s must be a non-empty string" % (label, name))
    for name in ("as_of", "window_start", "window_end"):
        _require_aware_time(row[name], label, name)
    if row["available"] is None or isinstance(row["available"], bool):
        raise LocalInputError("%s.available must be a number" % label)
    return row


def _validate_receipt_event(row, index):
    label = "receipts[%s]" % index
    _check_keys(row, RECEIPT_EVENT_FIELDS, RECEIPT_EVENT_REQUIRED,
                "receipts", index)
    for name in ("receipt_id", "line_id", "po_code", "sku", "warehouse", "uom"):
        if not isinstance(row[name], str) or not row[name].strip():
            raise LocalInputError("%s.%s must be a non-empty string" % (label, name))
    _require_bool(row["uom_proven"], label, "uom_proven")
    _require_bool(row["quantity_semantics_proven"], label,
                  "quantity_semantics_proven")
    _require_aware_time(row["receipt_effective_at"], label, "receipt_effective_at",
                        allow_none=True)
    if row["qty"] is None or isinstance(row["qty"], bool):
        raise LocalInputError("%s.qty must be a number" % label)
    for name in ("po_line_id", "dn_line_id", "shipment_id"):
        if row.get(name) is not None and not isinstance(row[name], str):
            raise LocalInputError("%s.%s must be a string or null" % (label, name))
    proof = row.get("inclusion_proof")
    if proof is not None:
        if not isinstance(proof, dict):
            raise LocalInputError("%s.inclusion_proof must be an object" % label)
        _check_keys(proof, {"proof_type", "source_record", "receipt_identity",
                            "snapshot_id", "source_as_of"},
                    {"proof_type", "source_record", "receipt_identity",
                     "snapshot_id", "source_as_of"},
                    label + ".inclusion_proof", index)
        identity = proof["receipt_identity"]
        if not isinstance(identity, (dict, list, tuple)):
            raise LocalInputError("%s.inclusion_proof.receipt_identity invalid" % label)
        _require_aware_time(proof["source_as_of"], label + ".inclusion_proof",
                            "source_as_of")
    return row


def _load_local_input(path, kind):
    try:
        payload = _load_json(path)
    except (OSError, json.JSONDecodeError) as exc:
        raise LocalInputError("cannot read %s: %s" % (path, exc))
    if kind == "receipt_state":
        allowed = {"schema_version", "coverage", "links", "receipts"}
        required = {"schema_version", "coverage", "receipts"}
        schema = RECEIPT_STATE_SCHEMA
    else:
        allowed = {"schema_version", "coverage", "snapshots", "receipts"}
        required = allowed
        schema = INVENTORY_BRIDGE_SCHEMA
    unknown = sorted(set(payload) - allowed) if isinstance(payload, dict) else []
    if not isinstance(payload, dict):
        raise LocalInputError("%s must be a JSON object" % kind)
    if unknown:
        raise LocalInputError("%s unknown top-level field(s): %s"
                              % (kind, ", ".join(unknown)))
    missing = sorted(required - set(payload))
    if missing:
        raise LocalInputError("%s missing top-level field(s): %s"
                              % (kind, ", ".join(missing)))
    if payload["schema_version"] != schema:
        raise LocalInputError("%s schema_version must be %s (got %r)"
                              % (kind, schema, payload["schema_version"]))
    _validate_coverage(payload["coverage"])
    receipts = payload["receipts"]
    if not isinstance(receipts, list):
        raise LocalInputError("%s.receipts must be a list" % kind)
    if kind == "receipt_state":
        links = payload.get("links", [])
        if not isinstance(links, list):
            raise LocalInputError("links must be a list")
        for index, row in enumerate(links):
            _validate_link_row(row, index)
        covered = {code.strip() for code in payload["coverage"]["covered_po_codes"]}
        for index, row in enumerate(links):
            uncovered = sorted({code for code in row["po_codes"]
                                if code not in covered})
            if uncovered:
                raise LocalInputError(
                    "links[%s].po_codes not covered by coverage.covered_po_codes: %s"
                    % (index, ", ".join(uncovered)))
        for index, row in enumerate(receipts):
            _validate_receipt_line(row, index)
        _validate_duplicate_identity(
            receipts, lambda row: (row.get("receipt_id"), row.get("line_id")),
            lambda row: {k: v for k, v in row.items()
                         if k not in ("receipt_id", "line_id")},
            "receipts")
    else:
        snapshots = payload["snapshots"]
        if not isinstance(snapshots, list):
            raise LocalInputError("snapshots must be a list")
        for index, row in enumerate(snapshots):
            _validate_snapshot_row(row, index)
        for index, row in enumerate(receipts):
            _validate_receipt_event(row, index)
        _validate_duplicate_identity(
            receipts, lambda row: (row.get("receipt_id"), row.get("line_id")),
            lambda row: {k: v for k, v in row.items()
                         if k not in ("receipt_id", "line_id")},
            "receipts")
        _validate_duplicate_identity(
            snapshots,
            lambda row: (row.get("inventory_snapshot_id"), row.get("sku"),
                         row.get("warehouse"), row.get("uom")),
            lambda row: row, "snapshots")
    return payload


def _auto_shipment_links(result):
    """Build Shipment -> DN -> PO links from the server's own evidence."""
    balances = {}
    for row in result.get("stage_balances") or []:
        if row.get("stage") != "signed_pending_tongtool_receipt":
            continue
        dn = row.get("batch_key")
        if not dn:
            continue
        key = (dn, row.get("item_code"), row.get("uom"), row.get("warehouse"))
        balances.setdefault(key, []).append(row)
    links, seen = [], set()
    for row in result.get("shipment_progress") or []:
        if row.get("stage") != "signed_pending_tongtool_receipt":
            continue
        if row.get("identity_conflict"):
            continue
        shipment = row.get("shipment")
        dn = row.get("dn")
        if not shipment or not dn:
            continue
        po_codes = tuple(code for code in [row.get("tongtool_po")] if code)
        key = (shipment, dn, po_codes)
        if key in seen:
            continue
        seen.add(key)
        candidates = balances.get(
            (dn, row.get("item_code"), row.get("uom"), row.get("warehouse")), [])
        qty = (candidates[0].get("qty")
               if len(candidates) == 1 and candidates[0].get("qty") is not None
               else None)
        links.append(ShipmentLink(
            shipment=shipment, dn=dn, po_codes=po_codes,
            signed_at=row.get("shipment_signed_at"), qty=qty,
            uom=row.get("uom"), sku=row.get("item_code"),
            warehouse=row.get("warehouse"),
            identity_confidence=("proven" if po_codes and qty is not None
                                 else "unproven"),
            status_semantics="unproven",
            signed_status_proven=bool(
                row.get("signed_status_proven") or row.get("is_final"))))
    return links


def _receipt_event(row):
    return ReceiptEvent(
        receipt_id=row.get("receipt_id"), line_id=row.get("line_id"),
        po_code=row.get("po_code"), sku=row.get("sku"),
        warehouse=row.get("warehouse"), uom=row.get("uom"),
        quantity=row.get("qty", row.get("quantity")),
        stock_effective_at=row.get("receipt_effective_at",
                                   row.get("stock_effective_at")),
        uom_proven=row.get("uom_proven") is True,
        quantity_semantics_proven=row.get("quantity_semantics_proven") is True,
        po_line_id=row.get("po_line_id"), dn_line_id=row.get("dn_line_id"),
        shipment_id=row.get("shipment_id"),
        inclusion_proof_record=row.get("inclusion_proof"))


def _receipt_line(row):
    return ReceiptLine(
        receipt_id=row.get("receipt_id"), line_id=row.get("line_id"),
        po_code=row.get("po_code"), sku=row.get("sku"),
        warehouse=row.get("warehouse"), uom=row.get("uom"),
        qty=row.get("qty", row.get("quantity")),
        effective_at=row.get("effective_at", row.get("receipt_effective_at")),
        quantity_semantics=row.get("quantity_semantics", "unknown"),
        uom_proven=bool(row.get("uom_proven")),
        identity_proven=bool(row.get("identity_proven")))


def _shipment_link(row):
    return ShipmentLink(
        shipment=row.get("shipment"), dn=row.get("dn"),
        po_codes=tuple(row.get("po_codes") or ()),
        signed_at=row.get("signed_at"), qty=row.get("qty"),
        uom=row.get("uom"), sku=row.get("sku"),
        warehouse=row.get("warehouse"),
        identity_confidence=row.get("identity_confidence", "unproven"),
        status_semantics=row.get("status_semantics", "unproven"),
        signed_status_proven=row.get("signed_status_proven") is True)


def _snapshot_row(row):
    return SnapshotRow(
        inventory_snapshot_id=row.get("inventory_snapshot_id"),
        sku=row.get("sku"), warehouse=row.get("warehouse"),
        uom=row.get("uom"), available=row.get("available"),
        as_of=row.get("as_of"), window_start=row.get("window_start"),
        window_end=row.get("window_end"))


def _write_xlsx(path, result):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    columns = {
        "lines": list(result["lines"][0].keys()) if result["lines"] else [],
        "stage_balances": ["so", "so_line_id", "item_code", "warehouse", "uom", "stage",
                           "stage_detail", "qty", "qty_status", "qty_basis",
                           "cumulative_excess_qty", "batch_key", "source_record",
                           "exclusion_reason", "included_in_stock", "demand_status",
                           "scope", "role"],
        "milestones": ["so", "so_line_id", "item_code", "uom", "node", "baseline_planned_at",
                       "committed_due_at", "latest_expected_at", "actual_at", "qty",
                       "qty_status", "date_precision", "timezone_status", "observed_at",
                       "source_record", "note", "rule_version"],
        "availability_schedule": ["tranche_id", "so", "so_line_id", "item_code", "warehouse",
                                  "uom", "qty", "qty_status", "qty_basis", "stage",
                                  "available_date", "latest_expected_at", "date_kind",
                                  "date_confidence", "scenario", "eligible",
                                  "exclusion_reason", "batch_key", "source_record",
                                  "demand_status", "rule_version"],
        "links": ["so", "so_line_id", "source_doctype", "source_name", "source_line_id",
                  "relation", "target_doctype", "target_name", "target_is_record",
                  "qty", "uom", "proof", "allocation"],
        "production_evidence": ["so", "so_line_id", "line_item_code", "production_item",
                                "wo", "wo_qty", "wo_docstatus", "wo_cancelled",
                                "line_role", "evidence_class", "qty_proven",
                                "qty_estimated", "confidence", "qty_status",
                                "started_at", "completed_at", "time_status",
                                "observed_at", "gross_qty", "return_qty", "net_qty",
                                "source_records", "reasons", "conflicts"],
        "shipment_progress": ["so", "so_line_id", "item_code", "dn", "shipment",
                              "logistic_number", "shipment_candidates",
                              "identity_conflict", "warehouse", "uom", "status11",
                              "status4", "status4_normalized",
                              "status4_normalization", "status_semantics", "eta",
                              "atd", "ata", "dlvr", "latest_progress_status",
                              "latest_progress_at", "latest_progress_kind",
                              "progress_source", "is_final", "conflict_flag",
                              "conflict_reasons", "progress_candidates", "stage",
                              "shipment_signed_at", "tongtool_first_receipt_at",
                              "tongtool_last_receipt_at", "tongtool_receipt_at",
                              "receipt_lag_days", "receipt_completion",
                              "received_qty", "pending_qty",
                              "received_evidence_status", "pending_evidence_status",
                              "receipt_basis", "identity_confidence",
                              "signed_status_proven",
                              "receipt_source_status", "inventory_inclusion_status",
                              "source_status"],
        "receipt_reconciliation": ["shipment", "dn", "po_codes", "sku", "warehouse",
                                   "uom", "status", "status_semantics",
                                   "shipment_signed_at", "signed_status_proven",
                                   "signed_time_status",
                                   "tongtool_first_receipt_at",
                                   "tongtool_last_receipt_at", "tongtool_receipt_at",
                                   "receipt_time_status", "receipt_lag_seconds",
                                   "receipt_lag_hours", "receipt_lag_days",
                                   "receipt_lag_status", "receipt_completion",
                                   "received_qty", "pending_qty", "over_receipt_qty",
                                   "received_evidence_status",
                                   "pending_evidence_status", "receipt_conflict",
                                   "receipt_basis",
                                   "identity_confidence", "coverage_status",
                                   "receipts_outside_coverage",
                                   "outside_coverage_qty", "exclusion_reason",
                                   "receipt_duplicates", "unmatched_receipt_rows"],
        "inventory_bridge": ["receipt_id", "line_id", "po_code", "po_line_id",
                             "dn_line_id", "shipment_id", "sku", "warehouse", "uom",
                             "qty", "receipt_effective_at",
                             "receipt_inclusion_status", "inclusion_proof",
                             "inclusion_proof_type",
                             "post_snapshot_delta_qty", "conditional_delta_qty",
                             "eligible", "count_in_inventory", "inventory_snapshot_id",
                             "source_as_of", "snapshot_window_start",
                             "snapshot_window_end", "snapshot_candidates",
                             "exclusion_reason", "duplicate_of"],
    }
    for name in DELIVERY_DATASETS:
        columns[name] = (list(result[name][0].keys()) if result.get(name) else [])
    for name in DATASETS + DELIVERY_DATASETS:
        ws = wb.create_sheet(name)
        headers = columns[name]
        if name == "production_evidence":
            rows = _flatten_production_evidence(result[name])
        elif name == "shipment_progress":
            rows = _flatten_shipment_progress(result[name])
        else:
            rows = result[name]
        ws.append(headers)
        for row in rows:
            ws.append([
                json.dumps(value, ensure_ascii=False, default=str)
                if isinstance(value, (dict, list)) else value
                for value in (row.get(key) for key in headers)])
        ws.freeze_panes = "A2"
    covered = wb.create_sheet("Coverage")
    covered.append(["key", "value"])
    for key, value in (result["manifest"].get("coverage") or {}).items():
        covered.append([key, json.dumps(value, ensure_ascii=False)
                        if isinstance(value, (dict, list)) else value])
    man = wb.create_sheet("Manifest")
    man.append(["key", "value"])
    for key, value in result["manifest"].items():
        if key == "scope":
            value = json.dumps(value, ensure_ascii=False, default=str)
        elif isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False)
        man.append([key, value])
    wb.save(path)


def _sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_json(path, payload):
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=1, default=str)


def main(argv=None):
    parser = argparse.ArgumentParser(prog="supply_facts_export",
                                     description="供应事实批量导出（结构化 JSON/Excel + manifest）")
    parser.add_argument("--codes", help="客户码或 EN 物料码，逗号分隔")
    parser.add_argument("--so", help="销售订单号，逗号分隔")
    parser.add_argument("--all", action="store_true", help="全部未发完 SO 行（分页）")
    parser.add_argument("--test", action="store_true", help="使用测试系统")
    parser.add_argument("--env-file", help=".env 路径")
    parser.add_argument("--as-of", help="as_of 日期 YYYY-MM-DD（默认今天）")
    parser.add_argument("--include-delivered", action="store_true", help="含已发完 SO 行")
    parser.add_argument("--page-size", type=int, default=500)
    parser.add_argument("--out", default=str(DEFAULT_OUT))
    parser.add_argument("--no-excel", action="store_true")
    parser.add_argument("--receipt-state-json",
                        help="本地只读：{links:[...],receipts:[...]} 生成 receipt_reconciliation")
    parser.add_argument("--inventory-bridge-json",
                        help="本地只读：{snapshots:[...],receipts:[...]} 生成 inventory_bridge")
    parser.add_argument("--delivery-all", action="store_true",
                        help="只读扫描全部可读 DN 明细，生成 delivery_lines/association_edges/coverage_exceptions 等")
    parser.add_argument("--delivery-include-inactive", action="store_true",
                        help="交付扫描包含草稿/取消 DN（仅留痕，不计实际交付）")
    parser.add_argument("--delivery-customer", help="交付扫描客户过滤")
    parser.add_argument("--delivery-item", help="交付扫描物料过滤")
    args = parser.parse_args(argv)

    if not (args.codes or args.so or args.all or args.delivery_all):
        parser.error("至少给一个输入: --codes / --so / --all / --delivery-all")
    if not (1 <= args.page_size <= 2000):
        parser.error("--page-size 必须在 1..2000")
    env = "test" if args.test else "prod"
    key, secret = iss.load_credentials(env, iss.resolve_env_files(args.env_file))
    if not key or not secret or "your_" in key:
        print("✗ API 凭证未配置。请检查 EN_API/.env", file=sys.stderr)
        return 1
    base_url = iss.ENV_URLS[env]
    started = time.perf_counter()

    receipt_payload = bridge_payload = None
    try:
        if args.receipt_state_json:
            receipt_payload = _load_local_input(
                args.receipt_state_json, "receipt_state")
        if args.inventory_bridge_json:
            bridge_payload = _load_local_input(
                args.inventory_bridge_json, "inventory_bridge")
    except LocalInputError as exc:
        print("✗ 本地输入校验失败，未创建新版本: %s" % exc, file=sys.stderr)
        return 2

    result = {"lines": [], "stage_balances": [], "milestones": [],
              "availability_schedule": [], "links": [], "production_evidence": [],
              "shipment_progress": [], "receipt_reconciliation": [],
              "inventory_bridge": [],
              "delivery_lines": [], "association_edges": [],
              "coverage_exceptions": [], "so_fulfillment": [],
              "supply_ledger": [], "physical_chain": [],
              "delivery_coverage": {}, "delivery_pagination": None,
              "delivery_versions": {},
              "partial": False,
              "counts": {}, "coverage": {}, "sources": {}, "source_timestamps": {},
              "restricted": {}, "external_sources": {},
              "notes": [], "schema_version": "", "rule_version": "",
              "implementation_version": "", "as_of": args.as_of or ""}
    today_str = time.strftime("%Y-%m-%d")
    if args.delivery_all and not result["as_of"]:
        result["as_of"] = today_str
    if args.as_of and args.as_of != today_str:
        print("✗ 历史 as_of 未实现：只能导出当前快照，拒绝 --as-of %s" % args.as_of,
              file=sys.stderr)
        return 2
    base_filters = {"as_of": today_str}
    terminated_by = "exhausted"
    try:
        def _export_scope(label, scope_filters):
            batches, term = _paginate(
                base_url, key, secret, scope_filters, args.page_size)
            for batch in batches:
                _merge(result, batch, scope_key=label)
            rows = sum(len(batch.get("lines") or []) for batch in batches)
            print("  %s: %d 行（%s）" % (label, rows, term))
            return term

        terms = []
        if args.codes:
            codes = [c.strip() for c in args.codes.replace(";", ",").split(",") if c.strip()]
            t0 = time.perf_counter()
            sos = _resolve_sos(codes, base_url, key, secret)
            print("解析 %d 个码 → %d 张 SO（%.2fs）" % (len(codes), len(sos), time.perf_counter() - t0))
            if not sos:
                raise iss.QueryError("没有命中的销售订单")
            for so in sos:
                filters = dict(base_filters, sales_order=so, include_delivered=1)
                terms.append(_export_scope(so, filters))
        elif args.so:
            for so in [s.strip() for s in args.so.split(",") if s.strip()]:
                filters = dict(base_filters, sales_order=so, include_delivered=1)
                terms.append(_export_scope(so, filters))
        elif args.all:
            filters = dict(base_filters)
            if args.include_delivered:
                filters["include_delivered"] = 1
            terms.append(_export_scope("all", filters))
        terminated_by = "page_limit" if "page_limit" in terms else "exhausted"
        if terminated_by == "page_limit":
            print("  ! 达到 200 页上限，结果可能截断（manifest.truncated=true）",
                  file=sys.stderr)
    except (iss.QueryError, BatchVersionError) as exc:
        print("✗ 导出失败，未写入新版本: %s" % exc, file=sys.stderr)
        return 1

    _finalize_coverage(result)
    delivery_terminated = None
    if args.delivery_all:
        delivery_filters = {}
        if args.delivery_customer:
            delivery_filters["customer"] = args.delivery_customer
        if args.delivery_item:
            delivery_filters["item_code"] = args.delivery_item
        try:
            batches, delivery_terminated = _paginate_delivery(
                base_url, key, secret, args.page_size,
                include_inactive=args.delivery_include_inactive,
                base_filters=delivery_filters)
            first = batches[0]
            result["delivery_versions"] = {
                key: first.get(key) for key in
                ("schema_version", "rule_version", "implementation_version")}
            result["delivery_coverage"] = first.get("coverage") or {}
            result["delivery_pagination"] = first.get("pagination") or {}
            result["so_fulfillment"] = first.get("so_fulfillment") or []
            result["supply_ledger"] = first.get("supply_ledger") or []
            for batch in batches:
                for name in ("delivery_lines", "association_edges",
                             "coverage_exceptions", "physical_chain"):
                    result[name].extend(batch.get(name) or [])
            result["counts"].update({
                "delivery_lines": len(result["delivery_lines"]),
                "association_edges": len(result["association_edges"]),
                "coverage_exceptions": len(result["coverage_exceptions"]),
                "so_fulfillment": len(result["so_fulfillment"]),
                "supply_ledger": len(result["supply_ledger"]),
                "physical_chain": len(result["physical_chain"]),
            })
            if delivery_terminated == "page_limit":
                terminated_by = "page_limit"
            print("  delivery: %d 行（%s）"
                  % (len(result["delivery_lines"]), delivery_terminated))
        except (iss.QueryError, BatchVersionError) as exc:
            print("✗ 交付事实导出失败，未写入新版本: %s" % exc, file=sys.stderr)
            return 1
    receipt_links_source = "none"
    if receipt_payload is not None:
        links = [_shipment_link(row) for row in receipt_payload.get("links") or []]
        if links:
            receipt_links_source = "local_input"
        else:
            links = _auto_shipment_links(result)
            receipt_links_source = "server_auto" if links else "none"
        result["receipt_reconciliation"] = reconcile_signed_receipts(
            links,
            [_receipt_line(row) for row in receipt_payload.get("receipts") or []],
            coverage=receipt_payload.get("coverage"))
        result["_receipt_coverage"] = receipt_payload.get("coverage")
        result["_receipt_links_source"] = receipt_links_source
    if bridge_payload is not None:
        windows = [SnapshotWindow(row["start"], row["end"])
                   for row in bridge_payload["coverage"]["requested_windows"]]
        bridge = build_inventory_bridge(
            [_snapshot_row(row) for row in bridge_payload.get("snapshots") or []],
            [_receipt_event(row) for row in bridge_payload.get("receipts") or []],
            collection_windows=windows or None)
        result["inventory_bridge"] = bridge["rows"]
        result["_inventory_bridge_summary"] = bridge["summary"]
        result["_inventory_bridge_coverage"] = bridge_payload.get("coverage")
    query_seconds = time.perf_counter() - started
    run_id = hashlib.sha1(("%s|%s|%s|%s|%s|%s" % (
        result["as_of"], args.codes, args.so, args.all, len(result["lines"]),
        time.time_ns())).encode()).hexdigest()[:12]
    out = Path(args.out)
    run_dir = out / "runs" / ("run-%s-%s" % (time.strftime("%Y%m%d_%H%M%S"), run_id))
    tmp_dir = run_dir.with_suffix(".tmp")
    if tmp_dir.exists():
        shutil.rmtree(tmp_dir)
    tmp_dir.mkdir(parents=True)

    files = {}
    try:
        for name in DATASETS + DELIVERY_DATASETS:
            path = tmp_dir / ("%s.json" % name)
            _write_json(path, result[name])
            files["%s.json" % name] = {"path": "%s.json" % name, "sha256": _sha256(path),
                                       "rows": len(result[name])}
        generated_at = time.strftime("%Y-%m-%dT%H:%M:%S")
        manifest = {
            "schema_version": SCHEMA_VERSION,
            "supply_facts_schema_version": result["schema_version"] or None,
            "delivery_schema_version": result["delivery_versions"].get("schema_version"),
            "rule_version": result["rule_version"],
            "implementation_version": result["implementation_version"],
            "run_id": run_id,
            "as_of": result["as_of"],
            "as_of_kind": "current_snapshot",
            "observed_at": generated_at,
            "generated_at": generated_at,
            "scope": {"codes": args.codes, "so": args.so, "all": args.all,
                      "include_delivered": bool(args.include_delivered)},
            "coverage": dict(result["coverage"],
                             restricted=result["restricted"],
                             external_sources=result["external_sources"]),
            "sources": result["sources"],
            "source_timestamps": result["source_timestamps"],
            "restricted": result["restricted"],
            "external_sources": result["external_sources"],
            "delivery_coverage": result.get("delivery_coverage"),
            "delivery_pagination": result.get("delivery_pagination"),
            "counts": result["counts"],
            "partial": result["partial"],
            "notes": result["notes"],
            "truncated": terminated_by == "page_limit",
            "termination_reason": terminated_by,
            "performance": {"query_seconds": round(query_seconds, 3)},
        }
        if not args.no_excel:
            xlsx = tmp_dir / "supply_facts.xlsx"
            result["manifest"] = manifest
            _write_xlsx(xlsx, result)
            files["supply_facts.xlsx"] = {"path": "supply_facts.xlsx",
                                          "sha256": _sha256(xlsx), "rows": None}
        manifest["datasets"] = {
            "production_evidence": {
                "lines": len(result["production_evidence"]),
                "rows": len(_flatten_production_evidence(result["production_evidence"])),
            },
            "shipment_progress": {
                "rows": len(result["shipment_progress"]),
                "source": "server",
            },
            "receipt_reconciliation": {
                "rows": len(result["receipt_reconciliation"]),
                "source": ("local_receipt_state"
                           if receipt_payload is not None else "not_integrated"),
                "links_source": receipt_links_source,
                "coverage": result.get("_receipt_coverage"),
                "schema_version": (RECEIPT_STATE_SCHEMA
                                   if receipt_payload is not None else None),
            },
            "inventory_bridge": {
                "rows": len(result["inventory_bridge"]),
                "source": ("local_inventory_bridge"
                           if bridge_payload is not None else "not_integrated"),
                "coverage": result.get("_inventory_bridge_coverage"),
                "schema_version": (INVENTORY_BRIDGE_SCHEMA
                                   if bridge_payload is not None else None),
                "summary": result.get("_inventory_bridge_summary"),
                "eligible": False,
            },
            "delivery_lines": {
                "rows": len(result["delivery_lines"]),
                "source": "server" if args.delivery_all else "not_requested",
                "coverage": result.get("delivery_coverage"),
                "pagination": result.get("delivery_pagination"),
                "schema_version": result["delivery_versions"].get("schema_version"),
                "rule_version": result["delivery_versions"].get("rule_version"),
                "implementation_version": result["delivery_versions"].get(
                    "implementation_version"),
            },
            "association_edges": {
                "rows": len(result["association_edges"]),
                "source": "server" if args.delivery_all else "not_requested",
            },
            "coverage_exceptions": {
                "rows": len(result["coverage_exceptions"]),
                "source": "server" if args.delivery_all else "not_requested",
            },
            "so_fulfillment": {
                "rows": len(result["so_fulfillment"]),
                "source": "server" if args.delivery_all else "not_requested",
            },
            "supply_ledger": {
                "rows": len(result["supply_ledger"]),
                "source": "server" if args.delivery_all else "not_requested",
            },
            "physical_chain": {
                "rows": len(result["physical_chain"]),
                "source": "server" if args.delivery_all else "not_requested",
            },
        }
        manifest["files"] = files
        _write_json(tmp_dir / "manifest.json", manifest)
        tmp_dir.rename(run_dir)
    except Exception as exc:  # noqa: BLE001
        shutil.rmtree(tmp_dir, ignore_errors=True)
        print("✗ 写文件失败，未切换最新版本: %s" % exc, file=sys.stderr)
        return 1

    pointer = out / "latest.json"
    pointer_tmp = out / ".latest.json.tmp"
    _write_json(pointer_tmp, {
        "run_id": run_id, "run_dir": str(run_dir),
        "manifest": str(run_dir / "manifest.json"),
        "as_of": result["as_of"],
        "schema_version": result["schema_version"],
        "rule_version": result["rule_version"],
        "implementation_version": result["implementation_version"],
        "truncated": terminated_by == "page_limit",
        "termination_reason": terminated_by,
        "last_success_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "rows": len(result["lines"]),
        "delivery_rows": len(result["delivery_lines"]),
        "delivery_schema_version": result["delivery_versions"].get("schema_version"),
        "partial": result["partial"],
        "performance": {"query_seconds": round(query_seconds, 3),
                        "total_seconds": round(time.perf_counter() - started, 3)},
    })
    os.replace(pointer_tmp, pointer)

    print("\n✓ 完成: %s" % run_dir)
    print("  行: %d | 阶段余额: %d | 里程碑: %d | 可售片段: %d | 关联: %d" % (
        len(result["lines"]), len(result["stage_balances"]), len(result["milestones"]),
        len(result["availability_schedule"]), len(result["links"])))
    print("  partial: %s | 查询 %.2fs | 总计 %.2fs" % (
        result["partial"], query_seconds, time.perf_counter() - started))
    print("  最新指针: %s" % pointer)
    return 0


if __name__ == "__main__":
    sys.exit(main())
