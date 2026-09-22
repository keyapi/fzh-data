# -*- coding: utf-8 -*-
"""出货计划「同一菲号被多张单据重复占用」存量冲突清单 —— 只读。

背景：改造前扫菲号只在本单据内校验，导致多张出货计划（草稿/已提交）各占同一菲号的全额。
补丁上线后这些单据在保存时会被拦住，所以先出一份清单让用户人工处理。

用法:
  python EN_API/dp_tn_occupy_conflict.py --site test
  python EN_API/dp_tn_occupy_conflict.py --site prod

输出: EN_API/out/dp_tn_occupy_conflict_<site>_<ts>.xlsx
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

from bom_fabric_length_check import ENV_KEYS, ENV_URLS, ErpnextClient, load_env  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

SCRIPT_NAME = "zz_dp_tn_occupy_conflict"

SERVER_SCRIPT = '''mode = frappe.form_dict.get("mode")

if mode == "all":
    claims = frappe.db.sql("""
        SELECT dp.name AS plan, dp.docstatus AS ds, iq.tracking_number AS tn,
               iq.item_code AS item_code, iq.planned_delivery_qty AS qty
        FROM `tabDelivery Plan` dp
        INNER JOIN `tabDelivery Plan Item Qty` iq ON iq.parent = dp.name
        WHERE dp.docstatus IN (0, 1)
          AND iq.tracking_number IS NOT NULL AND iq.tracking_number != ''
    """, as_dict=True)
    sle = frappe.db.sql("""
        SELECT tracking_number AS tn, SUM(actual_qty) AS bal
        FROM `tabStock Ledger Entry`
        WHERE is_cancelled = 0 AND tracking_number IS NOT NULL AND tracking_number != ''
        GROUP BY tracking_number
    """, as_dict=True)
    frappe.response["data"] = {
        "claims": [{"plan": r.plan, "docstatus": r.ds, "tn": r.tn,
                    "item_code": r.item_code, "qty": frappe.utils.flt(r.qty)} for r in claims],
        "sle": {r.tn: frappe.utils.flt(r.bal) for r in sle},
    }
'''


def call(client: ErpnextClient, **data):
    r = client._request("POST", f"/api/method/{SCRIPT_NAME}", data=data, timeout=(30, 600))
    body = r.json()
    for k in ("message", "data"):
        if k in body:
            return body[k]
    return body


def ensure_script(client: ErpnextClient) -> None:
    try:
        client._request("DELETE", f"/api/resource/Server Script/{SCRIPT_NAME}", timeout=60, retries=0)
    except Exception:  # noqa: BLE001
        pass
    client._request("POST", "/api/resource/Server Script",
                    json={"name": SCRIPT_NAME, "script_type": "API", "api_method": SCRIPT_NAME,
                          "script": SERVER_SCRIPT, "allow_guest": 0, "disabled": 0}, timeout=60)


def drop_script(client: ErpnextClient) -> None:
    try:
        client._request("DELETE", f"/api/resource/Server Script/{SCRIPT_NAME}", timeout=60, retries=0)
    except Exception:  # noqa: BLE001
        pass
    try:
        client._request("GET", f"/api/resource/Server Script/{SCRIPT_NAME}", timeout=60, retries=0)
        print(f"  ⚠ {SCRIPT_NAME} 删除后仍能读到，请复查")
    except Exception:  # noqa: BLE001
        print(f"已删除临时 Server Script {SCRIPT_NAME}（残留 0）")


HEADERS = ["菲号", "物料", "占用单据数", "占用合计", "其中草稿", "其中已提交",
           "SLE裸余额(参考)", "占用-余额", "单据明细"]
WIDTHS = [26, 26, 11, 11, 10, 11, 15, 12, 70]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--site", choices=("prod", "test"), default="test")
    ap.add_argument("--out-dir", type=Path, default=Path(__file__).resolve().parent / "out")
    args = ap.parse_args()

    key, sec = load_env(args.site)
    if not key or not sec:
        print(f"✗ 缺 {ENV_KEYS[args.site]} 凭证")
        return 1
    client = ErpnextClient(ENV_URLS[args.site], key, sec)
    ensure_script(client)
    try:
        data = call(client, mode="all")
    finally:
        drop_script(client)

    claims = data["claims"]
    sle = data["sle"]

    by_tn: dict[str, list[dict]] = defaultdict(list)
    for c in claims:
        by_tn[c["tn"]].append(c)

    conflicts = []
    for tn, rows in by_tn.items():
        plans = {r["plan"] for r in rows}
        if len(plans) < 2:
            continue
        total = sum(r["qty"] for r in rows)
        draft = sum(r["qty"] for r in rows if r["docstatus"] == 0)
        submitted = sum(r["qty"] for r in rows if r["docstatus"] == 1)
        bal = sle.get(tn)
        conflicts.append({
            "tn": tn,
            "item": rows[0]["item_code"],
            "plan_count": len(plans),
            "total": total, "draft": draft, "submitted": submitted,
            "bal": bal, "gap": (total - bal) if bal is not None else None,
            "detail": "；".join(
                f"{p}({'草稿' if ds == 0 else '已提交'}) "
                + "+".join(str(r["qty"]) for r in rows if r["plan"] == p)
                for p, ds in sorted({(r["plan"], r["docstatus"]) for r in rows})
            ),
        })
    conflicts.sort(key=lambda x: (-x["plan_count"], -(x["total"] or 0)))

    args.out_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = args.out_dir / f"dp_tn_occupy_conflict_{args.site}_{ts}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "重复占用"
    ws.append(HEADERS)
    head = PatternFill("solid", fgColor="DDEBF7")
    red = PatternFill("solid", fgColor="FFC7CE")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head
        c.alignment = Alignment(horizontal="center", vertical="center")
    for x in conflicts:
        ws.append([x["tn"], x["item"], x["plan_count"], x["total"], x["draft"], x["submitted"],
                   x["bal"], x["gap"], x["detail"]])
        if x["gap"] is not None and x["gap"] > 1e-9:
            for ci in range(1, len(HEADERS) + 1):
                ws.cell(row=ws.max_row, column=ci).fill = red
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"
    wb.save(str(out))

    print("=" * 60)
    print(f"站点 {args.site}：{len(claims)} 条占用记录，{len(by_tn)} 个菲号，"
          f"其中 {len(conflicts)} 个被 >1 张单据占用")
    over = [x for x in conflicts if x["gap"] is not None and x["gap"] > 1e-9]
    print(f"  其中「占用合计 > SLE裸余额」的：{len(over)} 个（红底行）")
    for x in conflicts[:10]:
        print(f"  {x['tn']}: {x['plan_count']} 张单 合计 {x['total']} / 余额 {x['bal']}")
    print(f"\n✓ 清单: {out}")
    print("  注：SLE裸余额是按菲号直接求和的参考值，未按 物料+仓库 维度过滤；")
    print("      单据内扫码用的是严格口径值，可能与本列略有出入。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
