# -*- coding: utf-8 -*-
"""三角靠枕 皮壳 BOM 面料用量核对 —— 只读报告（不写任何数据）

需求：
  对每份「三角靠枕 皮壳 BOM」，按父物料尺寸算出应有面料长度，与 BOM 现用量比对：
      应有长度(米) = (尺寸 + 40) / 100
  例：父物料 PK#KS0001-DM-160-CHEEKRED 尺寸 160
      -> 应有 200cm = 2.00 米
      BOM 面料行 DMML1010-CHEEKRED-150-320 (涤麻-腮红色-150cm-320g/m2) 现用量 2.38 米
      2.00 < 2.38 -> 正常；若 应有 > 现用量 -> 高亮「需工厂确认」

口径：
  - 范围   : BOM.item_name 含「三角」且含「靠枕」，item 以 PK# 开头，is_default=1 且已提交且启用
  - 尺寸   : 父物料 PK#<型号>-<材质码>-<尺寸>-<颜色> 的倒数第 2 段；
             异形（140x50x24 / 75*58*55）取第一个数字
  - 主面料 : 该 BOM 里面料行（Item Group 祖先链含「面料」）中，颜色段与父物料颜色一致、用量最大的一条；
             无颜色命中则取第一条面料行
  - 内胆 / 底开内胆：本次不参与

只读：全程只有 GET /api/resource/...，不含任何 PUT/POST。

用法:
  python EN_API/bom_fabric_length_check.py                            # prod，联网拉取
  python EN_API/bom_fabric_length_check.py --env test
  python EN_API/bom_fabric_length_check.py --from-json out/xxx.json   # 用本地快照跑，只联网取物料 meta

输出（EN_API/out/）:
  bom_fabric_length_check_<ts>.xlsx       报告
  bom_fabric_length_check_raw_<ts>.json   本次拉到的皮壳 BOM 原文快照（留档核对）
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from urllib.parse import quote

import requests
from requests.adapters import HTTPAdapter
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

_DIR = Path(__file__).resolve().parent
OUT_DIR = _DIR / "out"

ENV_URLS = {"prod": "https://erpnext.vilavi.cn", "test": "https://ensh.vilavi.cn"}
ENV_KEYS = {
    "prod": ("PROD_ERP_API_KEY", "PROD_ERP_API_SECRET"),
    "test": ("TEST_ERP_API_KEY", "TEST_ERP_API_SECRET"),
}

# ── 业务口径 ─────────────────────────────────────────
FABRIC_ROOT = "面料"            # Item Group 树里「面料」这一层，祖先链含它即为面料
SIZE_BONUS_CM = 40              # 应有长度 = 尺寸 + 40cm
VERDICT_BAD = "需工厂确认"
VERDICT_OK = "正常"

HEADERS = [
    "BOM", "父物料", "父物料名称", "型号", "材质码", "尺寸", "颜色",
    "主面料物料", "主面料名称", "现用量(米)", "需求(米)", "差额", "判定", "备注",
]
WIDTHS = [40, 30, 34, 10, 10, 7, 14, 34, 32, 11, 10, 9, 14, 34]
NUM_RE = re.compile(r"\d+")


def load_env(env: str) -> tuple[str, str]:
    """读 EN_API/.env（必要时回退上层 .env）。"""
    vals: dict[str, str] = {}
    for p in (_DIR / ".env", _DIR.parent / ".env", _DIR.parent.parent / ".env"):
        if not p.is_file():
            continue
        for line in p.read_text(encoding="utf-8", errors="ignore").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, _, v = line.partition("=")
            v = v.strip()
            if len(v) >= 2 and v[0] == v[-1] and v[0] in ('"', "'"):
                v = v[1:-1]
            vals.setdefault(k.strip(), v)
    key_name, sec_name = ENV_KEYS[env]
    return vals.get(key_name, ""), vals.get(sec_name, "")


class _NoExpectAdapter(HTTPAdapter):
    def send(self, request, **kwargs):  # noqa: ANN001
        request.headers.pop("Expect", None)
        return super().send(request, **kwargs)


class ErpnextClient:
    def __init__(self, base_url: str, api_key: str, api_secret: str) -> None:
        self.base_url = base_url.rstrip("/")
        self.session = requests.Session()
        self.session.headers["Authorization"] = f"token {api_key}:{api_secret}"
        self.session.mount("https://", _NoExpectAdapter())
        self.session.mount("http://", _NoExpectAdapter())

    def _request(self, method: str, path: str, *, retries: int = 3,
                 retry_delay: float = 3.0, **kwargs: Any) -> requests.Response:
        timeout = kwargs.pop("timeout", (30, 300))
        url = f"{self.base_url}{path}"
        last: Exception | None = None
        for attempt in range(retries + 1):
            try:
                r = self.session.request(method, url, timeout=timeout, **kwargs)
                r.raise_for_status()
                return r
            except Exception as exc:  # noqa: BLE001
                last = exc
                if attempt < retries:
                    time.sleep(retry_delay * (attempt + 1))
        assert last is not None
        raise last

    def get_list(self, doctype: str, filters: list | None = None,
                 fields: list[str] | None = None, limit_page_length: int = 0) -> list[dict]:
        params: dict[str, str] = {"limit_page_length": str(limit_page_length)}
        if filters is not None:
            params["filters"] = json.dumps(filters)
        if fields is not None:
            params["fields"] = json.dumps(fields)
        return self._request("GET", f"/api/resource/{quote(doctype, safe='')}",
                             params=params).json()["data"]

    def get_doc(self, doctype: str, name: str) -> dict:
        return self._request(
            "GET", f"/api/resource/{quote(doctype, safe='')}/{quote(name, safe='')}"
        ).json()["data"]

    def get_items_meta(self, codes: list[str], batch: int = 60) -> dict[str, dict]:
        out: dict[str, dict] = {}
        for s in range(0, len(codes), batch):
            rows = self.get_list(
                "Item", filters=[["name", "in", codes[s:s + batch]]],
                fields=["name", "item_code", "item_name", "item_group", "stock_uom"],
            )
            for r in rows:
                out[r["name"]] = r
        return out


class ItemGroupTree:
    """Item Group 树：判断某物料组是否属于「面料」（祖先链含 FABRIC_ROOT）。"""

    def __init__(self, groups: list[dict]) -> None:
        self.parent = {g["name"]: (g["parent_item_group"] or "") for g in groups}
        self._cache: dict[str, bool] = {}

    def chain(self, name: str) -> list[str]:
        out: list[str] = []
        cur = name
        while cur and cur not in out:
            out.append(cur)
            cur = self.parent.get(cur, "")
        return out

    def is_fabric(self, name: str) -> bool:
        if not name:
            return False
        if name not in self._cache:
            self._cache[name] = FABRIC_ROOT in self.chain(name)
        return self._cache[name]


def is_triangle_backrest(name: str) -> bool:
    return bool(name) and "三角" in name and "靠枕" in name


# ── 父物料解析：PK#<型号>-<材质码>-<尺寸>-<颜色> ──────────────────

def split_parent(item_code: str) -> list[str]:
    body = item_code.split("#", 1)[1] if "#" in item_code else item_code
    return body.split("-")


def parse_size(item_code: str) -> tuple[int | None, str]:
    """返回 (尺寸数值, 尺寸原文)。140x50x24 -> (140, '140x50x24')；解析不出 -> (None, token)。"""
    parts = split_parent(item_code)
    if len(parts) < 4:
        return None, ""
    token = parts[-2]
    if token.isdigit():
        return int(token), token
    m = NUM_RE.match(token)
    if m:
        return int(m.group()), token
    return None, token


def parent_attrs(item_code: str) -> tuple[str, str, str]:
    """返回 (型号, 材质码, 颜色)。"""
    parts = split_parent(item_code)
    model = parts[0] if parts else ""
    material = parts[-3] if len(parts) >= 3 else ""
    color = parts[-1] if len(parts) >= 4 else ""
    return model, material, color


def _color_matches(code: str, color: str) -> bool:
    """面料物料里是否有与父物料颜色「完全相同」的段（避免 BLUEGREY 误命中 GREY）。"""
    return bool(color) and color in code.split("-")


def pick_main_fabric(fabric_rows: list[dict], color: str) -> dict | None:
    """面料行中优先取颜色段与父物料一致、用量最大的一条；无命中取第一条。"""
    if not fabric_rows:
        return None
    matched = [r for r in fabric_rows if _color_matches(r["code"], color)]
    if matched:
        return max(matched, key=lambda r: r["qty"])
    return fabric_rows[0]


def need_meter(size_cm: int) -> Decimal:
    return (Decimal(size_cm) + Decimal(SIZE_BONUS_CM)) / Decimal(100)


def collect(env: str, client: ErpnextClient, workers: int,
            from_json: Path | None) -> tuple[list[dict], dict[str, dict], ItemGroupTree]:
    # 1) Item Group 树（面料判定口径）
    print("[1/5] 取 Item Group 树（面料判定口径）…", flush=True)
    tree = ItemGroupTree(client.get_list(
        "Item Group",
        fields=["name", "item_group_name", "parent_item_group", "is_group"],
    ))
    print(f"      物料组 {len(tree.parent)} 个；「{FABRIC_ROOT}」祖先链: "
          f"{' < '.join(tree.chain(FABRIC_ROOT))}", flush=True)

    # 2) 三角靠枕 皮壳 BOM 清单 / 或直接读本地快照
    if from_json is not None:
        print(f"\n[2/5] 读本地快照 {from_json} …", flush=True)
        payload = json.loads(from_json.read_text(encoding="utf-8"))
        raw = payload.get("boms", payload if isinstance(payload, list) else [])
        docs = [d for d in raw if (d.get("item") or "").startswith("PK#")]
        print(f"      快照 {len(raw)} 份，其中皮壳 {len(docs)} 份", flush=True)
    else:
        print("\n[2/5] 拉取三角靠枕 皮壳 BOM 清单 …", flush=True)
        boms = client.get_list(
            "BOM",
            filters=[["item_name", "like", "%三角靠枕%"],
                     ["is_default", "=", 1],
                     ["docstatus", "=", 1],
                     ["is_active", "=", 1]],
            fields=["name", "item", "item_name", "is_default", "is_active", "docstatus"],
        )
        scoped = [b for b in boms
                  if (b["item"] or "").startswith("PK#") and is_triangle_backrest(b["item_name"])]
        print(f"      命中 {len(boms)} 份，其中皮壳 {len(scoped)} 份", flush=True)

        print(f"\n[3/5] 逐份取 BOM 单据（{workers} 线程，子表无法 list）…", flush=True)
        docs = []
        names = [b["name"] for b in scoped]
        with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
            futures = {pool.submit(client.get_doc, "BOM", n): n for n in names}
            for i, fut in enumerate(as_completed(futures), start=1):
                nm = futures[fut]
                try:
                    docs.append(fut.result())
                except Exception as exc:  # noqa: BLE001
                    print(f"      [WARN] 取 {nm} 失败: {exc}", flush=True)
                if i % 200 == 0 or i == len(names):
                    print(f"      已取 {i}/{len(names)}", flush=True)

    # 3) 全部组件物料的 meta
    print("\n[4/5] 查询组件物料的 item_group …", flush=True)
    all_codes = sorted({i["item_code"] for d in docs for i in (d.get("items") or [])})
    meta = client.get_items_meta(all_codes)
    missing = [c for c in all_codes if c not in meta]
    if missing:
        print(f"      [WARN] {len(missing)} 个组件查不到 meta，样例: {missing[:5]}", flush=True)
    print(f"      组件物料 {len(all_codes)} 个，已取 meta {len(meta)} 个", flush=True)
    return docs, meta, tree


def evaluate(docs: list[dict], meta: dict[str, dict], tree: ItemGroupTree) -> tuple[list[dict], list[dict]]:
    """返回 (可判定记录, 无法判定记录)。"""
    print("\n[5/5] 逐份解析尺寸 / 取主面料 / 判定 …", flush=True)
    records: list[dict] = []
    undecided: list[dict] = []

    for d in docs:
        item, item_name = d["item"], d.get("item_name") or ""
        model, material, color = parent_attrs(item)
        size, size_token = parse_size(item)
        irregular = not size_token.isdigit()

        fabric_rows = []
        for idx, row in enumerate(d.get("items") or [], start=1):
            cm = meta.get(row["item_code"], {})
            grp = cm.get("item_group", "")
            if not tree.is_fabric(grp):
                continue
            fabric_rows.append({
                "idx": idx, "code": row["item_code"], "row_name": row.get("name"),
                "name": cm.get("item_name") or row.get("item_name") or "",
                "group": grp, "uom": row.get("uom"),
                "qty": float(row.get("qty") or 0), "rate": float(row.get("rate") or 0),
                "amount": float(row.get("amount") or 0),
            })

        base = {"bom": d["name"], "item": item, "item_name": item_name,
                "model": model, "material": material, "color": color}

        if size is None or not fabric_rows:
            reason = "解析不出尺寸" if size is None else "BOM 里没有面料行"
            undecided.append({**base, "reason": reason})
            continue

        meter = [r for r in fabric_rows if r["uom"] == "米"]
        main = pick_main_fabric(meter or fabric_rows, color)
        assert main is not None

        need = need_meter(size)
        qty = Decimal(str(main["qty"]))
        diff = (qty - need).quantize(Decimal("0.01"))

        notes = []
        if irregular:
            notes.append(f"拼色/异形款，尺寸取 {size_token} 中的 {size}")
        if len(fabric_rows) > 1:
            notes.append(f"该 BOM 有 {len(fabric_rows)} 条面料行，取主面料（颜色段={color}）")
        if main["uom"] != "米":
            notes.append(f"主面料单位为「{main['uom']}」，非米")

        records.append({
            **base,
            "size": size, "size_token": size_token,
            "irregular": irregular,
            "code": main["code"], "name": main["name"], "uom": main["uom"],
            "qty": qty, "need": need, "diff": diff,
            "verdict": VERDICT_BAD if diff < 0 else VERDICT_OK,
            "note": "；".join(notes),
            "fabric_rows": fabric_rows,
        })

    records.sort(key=lambda r: (r["model"], r["material"], r["size"], r["color"]))
    undecided.sort(key=lambda r: r["bom"])
    print(f"      可判定 {len(records)} 份，无法判定 {len(undecided)} 份", flush=True)
    return records, undecided


def write_report(out: Path, meta_info: dict, records: list[dict], undecided: list[dict]) -> None:
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    gray_fill = PatternFill("solid", fgColor="EDEDED")

    # ── 汇总 ──
    ws = wb.active
    ws.title = "汇总"
    ws.append(["三角靠枕 皮壳 BOM 面料用量核对（只读，未写库）"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["环境", meta_info["env"], "生成时间", meta_info["generated_at"]])
    ws.append(["需求口径", f"需求(米) = (尺寸 + {SIZE_BONUS_CM}) / 100",
               "", "例：尺寸 160 → 200cm → 2.00 米"])
    ws.append(["面料判定", f"Item Group 祖先链含「{FABRIC_ROOT}」", "", "拉链/扣胚/流苏/蕾丝等辅料自动排除"])
    ws.append(["主面料选取", "面料行中颜色段与父物料一致、用量最大的一条；无命中取第一条"])
    ws.append(["范围", "三角靠枕 皮壳 BOM（is_default=1 且已提交且启用）", "", "内胆/底开内胆不参与"])
    ws.append([])

    n_ok = sum(1 for r in records if r["verdict"] == VERDICT_OK)
    n_bad = sum(1 for r in records if r["verdict"] == VERDICT_BAD)
    n_irr = sum(1 for r in records if r["irregular"])
    ws.append(["统计项", "数量", "说明"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = head_fill
    ws.append(["皮壳 BOM 可判定", len(records), "一行一个 BOM"])
    ws.append([VERDICT_OK, n_ok, "需求 ≤ 现用量"])
    ws.append([VERDICT_BAD, n_bad, f"需求 > 现用量，整行标红，交工厂确定"])
    ws.append(["其中 拼色/异形款", n_irr, "按主尺寸 + 主面料判定"])
    ws.append(["无法判定", len(undecided), "解析不出尺寸 / 没有面料行，见「无法判定」表"])
    ws.append([])

    ws.append(["按尺寸分布", "BOM 数", VERDICT_OK, VERDICT_BAD])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = head_fill
    by_size: dict[Any, list[dict]] = defaultdict(list)
    for r in records:
        by_size[r["size"]].append(r)
    for size in sorted(by_size, key=lambda s: (s is None, s)):
        rows = by_size[size]
        ws.append([size, len(rows),
                   sum(1 for r in rows if r["verdict"] == VERDICT_OK),
                   sum(1 for r in rows if r["verdict"] == VERDICT_BAD)])
    for col, w in {"A": 30, "B": 14, "C": 16, "D": 46}.items():
        ws.column_dimensions[col].width = w

    # ── 明细表 ──
    def write_rows(title: str, recs: list[dict]) -> None:
        sheet = wb.create_sheet(title)
        sheet.append(HEADERS)
        for c in sheet[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
        if not recs:
            sheet.append(["（无）"])
        for r in recs:
            sheet.append([
                r["bom"], r["item"], r["item_name"], r["model"], r["material"],
                r["size"], r["color"], r["code"], r["name"],
                float(r["qty"]), float(r["need"]), float(r["diff"]),
                r["verdict"], r["note"],
            ])
            for ci in (10, 11, 12):                       # 现用量 / 需求 / 差额
                sheet.cell(row=sheet.max_row, column=ci).number_format = "0.00"
            if r["verdict"] == VERDICT_BAD:
                for ci in range(1, len(HEADERS) + 1):
                    sheet.cell(row=sheet.max_row, column=ci).fill = red_fill
        for i, w in enumerate(WIDTHS, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = w
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{sheet.max_row}"

    write_rows("皮壳BOM核对", records)
    write_rows("需工厂确认", [r for r in records if r["verdict"] == VERDICT_BAD])
    write_rows("拼色异形款", [r for r in records if r["irregular"]])

    # ── 无法判定 ──
    sheet = wb.create_sheet("无法判定")
    sheet.append(["BOM", "父物料", "父物料名称", "原因"])
    for c in sheet[1]:
        c.font = Font(bold=True)
        c.fill = head_fill
    for r in undecided:
        sheet.append([r["bom"], r["item"], r["item_name"], r["reason"]])
        for ci in range(1, 5):
            sheet.cell(row=sheet.max_row, column=ci).fill = gray_fill
    for col, w in {"A": 40, "B": 30, "C": 34, "D": 26}.items():
        sheet.column_dimensions[col].width = w

    wb.save(str(out))


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--env", choices=("prod", "test"), default="prod")
    ap.add_argument("--workers", type=int, default=16)
    ap.add_argument("--out-dir", type=Path, default=OUT_DIR)
    ap.add_argument("--from-json", type=Path, default=None,
                    help="用本地 BOM 快照跑（只联网取 Item Group 树与组件 meta）")
    args = ap.parse_args()

    key, sec = load_env(args.env)
    if not key or not sec:
        print(f"✗ 缺少 {ENV_KEYS[args.env]} 凭证")
        return 1
    base = ENV_URLS[args.env]
    client = ErpnextClient(base, key, sec)
    print(f"环境: {args.env} ({base})   只读模式\n", flush=True)

    t0 = time.time()
    docs, meta, tree = collect(args.env, client, args.workers, args.from_json)
    if not docs:
        print("✗ 没有取到任何皮壳 BOM")
        return 1
    records, undecided = evaluate(docs, meta, tree)

    args.out_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx = args.out_dir / f"bom_fabric_length_check_{ts}.xlsx"
    raw_json = args.out_dir / f"bom_fabric_length_check_raw_{ts}.json"

    meta_info = {
        "env": args.env, "base_url": base, "generated_at": datetime.now().isoformat(),
        "scope": "三角靠枕 皮壳 BOM (item 以 PK# 开头, is_default=1)",
        "formula": f"需求(米) = (尺寸 + {SIZE_BONUS_CM}) / 100",
        "fabric_rule": f"Item Group 祖先链含「{FABRIC_ROOT}」",
    }
    write_report(xlsx, meta_info, records, undecided)
    raw_json.write_text(json.dumps(
        {"meta": meta_info, "boms": docs}, ensure_ascii=False, indent=1), encoding="utf-8")

    n_ok = sum(1 for r in records if r["verdict"] == VERDICT_OK)
    n_bad = len(records) - n_ok
    print("\n" + "=" * 64)
    print("汇总")
    print("=" * 64)
    print(f"  皮壳 BOM 可判定 : {len(records)} 份")
    print(f"    {VERDICT_OK:<8}: {n_ok}")
    print(f"    {VERDICT_BAD:<8}: {n_bad}   ← 整行标红")
    print(f"    其中 拼色/异形款: {sum(1 for r in records if r['irregular'])}")
    print(f"  无法判定        : {len(undecided)} 份")
    if n_bad:
        cnt = Counter((r["size"], float(r["qty"])) for r in records if r["verdict"] == VERDICT_BAD)
        print(f"\n  需工厂确认明细（尺寸, 现用量）→ 份数:")
        for (size, qty), n in sorted(cnt.items(), key=lambda t: -t[1]):
            print(f"    尺寸 {size:>4}  现用量 {qty:<6}  → {n} 份")
    print(f"\n  耗时 {time.time() - t0:.0f}s")
    print(f"✓ 报告: {xlsx}")
    print(f"✓ 快照: {raw_json}  ({raw_json.stat().st_size / 1024 / 1024:.1f} MB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
