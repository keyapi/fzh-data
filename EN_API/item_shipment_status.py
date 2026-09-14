#!/usr/bin/env python3
"""客户物料号 / EN 物料号 → 销售订单发货状态 + 生产计划/工单/工序卡

回答的问题：某个物料的销售订单，哪些已发货、哪些没发，以及没发的原因。

数据链:
  Sales Order Item.customer_item_code ─┐
  Sales Order Item.item_code          ─┴→ Sales Order (子表过滤, 按父单去重)
       ├─ get_single → items[]            → SO 行 (name = Production Plan Item.sales_order_item)
       ├─ Delivery Note Item.against_sales_order → Delivery Note    (已发/草稿)
       ├─ Work Order.production_item             → Work Order       (工单)
       │     └─ Job Card.work_order              → Job Card         (工序卡 = 真实进度)
       └─ Production Plan Item.sales_order       → Production Plan (生产计划)
  Item.customer_items[].ref_code ─────────────→ 客户码 ↔ 物料 交叉校验

── 五条 API 硬约束 (实测, 别再踩) ─────────────────────────────────────────
1. 子表不能直接 list: GET /api/resource/Sales Order Item → 403 PermissionError
2. 父单直查子字段 → 417 "Field not permitted in query"
   必须用子表过滤语法 [["Sales Order Item","customer_item_code","like","%码%"]]
3. 子表过滤查询"一行一子行", 同一父单会重复 → 先按父单 name 去重, 再 get_single 扇出
4. in 列表不能太长: 308 个 work_order 塞进过滤 → 400 Request Line is too large (6737 > 4094)
   所有 in 查询必须分块 (IN_CHUNK = 40)
5. fields 只能是父单字段, 传未知字段 417 (实测 Item.has_bom 被拒);
   子表值必须 get_single 后读 doc["items"] / doc["po_items"]

── 两条容易搞错的业务口径 ─────────────────────────────────────────────────
* Work Order.status / produced_qty 不可信 —— 实测 WO-26-02609 头部 Not Started、
  produced_qty=0, 但 14 张工序卡显示 44 件已过 裁剪/皮壳整件/锁扣眼/拷边。
  判断进度看 Job Card, 不要只看 WO.operation 状态。
* ERPNext 的 Closed ≠ 已发完 —— 实测 10 张 SO 里 3 张 Closed 却仍有未发量。

示例:
  uv run python EN_API/item_shipment_status.py --customer-code CENKZ1325-Yellow-138
  uv run python EN_API/item_shipment_status.py --item "PK#KS0001-DM-140-YELLOW"
  uv run python EN_API/item_shipment_status.py --customer-code CENKZ1325-Yellow-138 --no-excel
  uv run python EN_API/item_shipment_status.py --customer-code X --test -o EN_API/out/custom.xlsx
"""
import argparse, json, os, re, sys, time, unicodedata
from calendar import monthrange
from datetime import date, datetime
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
import requests
import urllib.parse as _urlparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── 路径 (脚本在 EN_API/ 下 → PROJECT_ROOT = 仓库根) ──
PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "EN_API" / "out"

ENV_URLS = {"test": "https://ensh.vilavi.cn", "prod": "https://erpnext.vilavi.cn"}

# .env 候选: worktree 优先, 再退回主仓库 (worktree 里没有 .env)
DEFAULT_ENV_FILES = (
    PROJECT_ROOT / "EN_API" / ".env",
    PROJECT_ROOT / ".env",
    PROJECT_ROOT.parent / "EN_API" / ".env",
    PROJECT_ROOT.parents[2] / "EN_API" / ".env",
)

IN_CHUNK = 40             # in 过滤分块大小 (约束 4)
MAX_RETRIES = 3           # 服务器偶发 ConnectTimeoutError
DUP_WINDOW_DAYS = 3       # 疑似重复下单: 交货日相差天数上限

SO_FIELDS = ["name", "status", "docstatus", "customer", "customer_name", "company",
             "delivery_date", "transaction_date", "per_delivered", "amended_from"]
WO_FIELDS = ["name", "status", "docstatus", "sales_order", "production_item", "qty",
             "produced_qty", "planned_start_date", "planned_end_date", "company", "bom_no"]
JC_FIELDS = ["name", "work_order", "operation", "status", "for_quantity",
             "total_completed_qty", "actual_start_date", "actual_end_date", "docstatus"]


# ── 小工具 ────────────────────────────────────────────
def chunks(seq, n=IN_CHUNK):
    """把 list 切块 —— in 过滤太长会被 nginx 拒 (400 Request Line is too large)。"""
    seq = list(seq)
    for i in range(0, len(seq), n):
        yield seq[i:i + n]


def close(a, b, tol=1e-6) -> bool:
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return False


def num(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def fmt_qty(v) -> str:
    f = num(v)
    return f"{f:g}"


def add_months(d: date, n: int) -> date:
    """日期加 n 个月, 月末自动收敛 (1/31 + 1 月 → 2/28)。"""
    m = d.month - 1 + n
    y = d.year + m // 12
    m = m % 12 + 1
    return date(y, m, min(d.day, monthrange(y, m)[1]))


def parse_date(s):
    if not s:
        return None
    try:
        return date.fromisoformat(str(s)[:10])
    except ValueError:
        return None


def slugify(s: str) -> str:
    s = str(s).replace("#", "PK_").replace("+", "plus_")
    return re.sub(r"[^A-Za-z0-9._-]+", "_", s).strip("_") or "report"


def code_in(haystack, needles, ci=True) -> bool:
    """大小写不敏感匹配 (实测客户码 -60CM 与 -60cm 并存)。"""
    if not haystack:
        return False
    h = haystack.lower() if ci else haystack
    for n in needles:
        if not n:
            continue
        if (n.lower() if ci else n) in h:
            return True
    return False


def dedupe_by(rows, key="name"):
    """子表过滤查询一行一子行 → 按父单折叠, 保持首次出现顺序 (约束 3)。"""
    seen, out = set(), []
    for r in rows:
        k = r.get(key)
        if not k or k in seen:
            continue
        seen.add(k)
        out.append(r)
    return out


def child_filter(child_doctype, fieldname, operator, value):
    """构造父单查询用的子表过滤 (约束 1+2)。

    子表不能直接 list, 父单直查子字段报 417, 只有这个四元组形式能被 frappe
    展开成 tab<child_doctype> 的子查询。
    注意: 返回结果一行一匹配子行, 必须按父单 name 去重。
    """
    return [[child_doctype, fieldname, operator, value]]


# ── 凭证 ──────────────────────────────────────────────
def resolve_env_files(explicit=None):
    if explicit:
        return [Path(explicit)]
    return [p for p in DEFAULT_ENV_FILES if p.is_file()]


def load_credentials(env_name="prod", env_files=None):
    """优先环境变量, 再读 .env 候选列表。"""
    if env_name == "test":
        key, secret = os.environ.get("TEST_ERP_API_KEY", ""), os.environ.get("TEST_ERP_API_SECRET", "")
    else:
        key, secret = os.environ.get("PROD_ERP_API_KEY", ""), os.environ.get("PROD_ERP_API_SECRET", "")
    if not key:
        key = os.environ.get("ERP_API_KEY", "")
    if not secret:
        secret = os.environ.get("ERP_API_SECRET", "")
    if key and secret:
        return key, secret

    wanted = ("TEST_ERP_API_KEY", "TEST_ERP_API_SECRET") if env_name == "test" \
        else ("PROD_ERP_API_KEY", "PROD_ERP_API_SECRET")
    for path in (env_files or resolve_env_files()):
        with open(path, encoding="utf-8-sig") as f:
            for line in f:
                line = line.strip()
                if not line or line.startswith("#") or "=" not in line:
                    continue
                k, v = line.split("=", 1)
                k, v = k.strip(), v.strip().strip("'\"")
                if not key and k in wanted:
                    key = v
                elif not secret and k in wanted:
                    secret = v
                elif not key and k == "ERP_API_KEY":
                    key = v
                elif not secret and k == "ERP_API_SECRET":
                    secret = v
        if key and secret:
            break
    return key, secret


# ── API helpers (简单 requests, 不用 session/keep-alive —— 避免 nginx 417) ──
def _build_url(resource, filters, fields, limit_start=0, limit_page_length=100, order_by=None):
    q = {
        "filters": json.dumps(filters, ensure_ascii=False),
        "fields": json.dumps(fields, ensure_ascii=False),
        "limit_start": str(limit_start),
        "limit_page_length": str(limit_page_length),
    }
    if order_by:
        q["order_by"] = order_by
    return f"/api/resource/{_urlparse.quote(resource, safe='')}?{_urlparse.urlencode(q)}"


def _auth_headers(key, secret):
    return {"Authorization": f"token {key}:{secret}"}


def api_get_status(path, base_url, key, secret, timeout=60):
    """返回 (status_code, body)。连接异常重试 —— 实测服务器偶发 ConnectTimeout。"""
    url = f"{base_url}{path}"
    last = None
    for attempt in range(MAX_RETRIES):
        try:
            resp = requests.get(url, headers=_auth_headers(key, secret), timeout=timeout)
            try:
                return resp.status_code, resp.json()
            except ValueError:
                return resp.status_code, {}
        except requests.RequestException as e:
            last = e
            if attempt < MAX_RETRIES - 1:
                time.sleep(2 * (attempt + 1))
    print(f"  ✗ 请求失败 ({type(last).__name__}): {path[:120]}", file=sys.stderr)
    return 0, {}


def api_get(path, base_url, key, secret):
    return api_get_status(path, base_url, key, secret)[1]


def paginated_get(resource, base_url, key, secret, filters, fields,
                  page_size=100, order_by="name asc", fallback_fields=None,
                  label=""):
    """翻页拉取。

    与 EN_API/dn_trace_report.py 的两处差异 (有意为之, 不回灌旧脚本):
      * order_by 默认 "name asc" —— 无稳定排序时跨页 LIMIT/OFFSET 会漏行或重复
      * 按实收行数前移 (start += len(data)), 不再 page * page_size
        —— 服务端返回不足一页时会跳行
    """
    all_rows, start, page = [], 0, 0
    while page < 200:
        status, body = api_get_status(
            _build_url(resource, filters, fields, start, page_size, order_by),
            base_url, key, secret)
        if status == 417 and fallback_fields and fields != fallback_fields:
            print(f"  ⚠ {label or resource}: 字段被拒 (417), 回退字段集 {fallback_fields}")
            return paginated_get(resource, base_url, key, secret, filters,
                                 fallback_fields, page_size, order_by, None, label)
        if status != 200:
            print(f"  ✗ {label or resource}: HTTP {status} {str(body)[:140]}", file=sys.stderr)
            break
        data = body.get("data", []) or []
        if not data:
            break
        all_rows.extend(data)
        if len(data) < page_size:
            break
        start += len(data)
        page += 1
    return all_rows


def get_single(resource, docname, base_url, key, secret):
    path = "/api/resource/{}/{}".format(_urlparse.quote(resource, safe=""),
                                        _urlparse.quote(docname, safe=""))
    return api_get(path, base_url, key, secret).get("data", {})


# ── Excel 样式 ────────────────────────────────────────
HDR_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HDR_FONT = Font(bold=True, size=11, color="FFFFFF")
ALT_FILL = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
ERR_FONT = Font(bold=True, color="C00000")
WARN_FONT = Font(color="BF8F00")


def border():
    return Border(left=Side(style="thin"), right=Side(style="thin"),
                  top=Side(style="thin"), bottom=Side(style="thin"))


def _write_sheet(wb, title, headers, rows, group_col=0, first=False, number_cols=()):
    """写一个 sheet。

    与旧脚本的 _write_sheet 的差别: 分组判断放在行处理"开始处",
    否则每组首行会继承上一组的底色 (旧版先算底色后判分组)。
    """
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for c, (h, w) in enumerate(headers, 1):
        cl = ws.cell(row=1, column=c, value=h)
        cl.font, cl.fill = HDR_FONT, HDR_FILL
        cl.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cl.border = border()
        ws.column_dimensions[get_column_letter(c)].width = w

    prev_group, alt = None, 0
    for r, vals in enumerate(rows, 2):
        g = vals[group_col] if group_col < len(vals) else None
        if prev_group is not None and g != prev_group:
            alt += 1
        prev_group = g
        for c in range(1, len(headers) + 1):
            v = vals[c - 1] if c <= len(vals) else ""
            cl = ws.cell(row=r, column=c, value=v)
            cl.border = border()
            cl.alignment = Alignment(vertical="center", wrap_text=True)
            if c in number_cols and isinstance(v, (int, float)):
                cl.number_format = "#,##0.##"
            if alt % 2 == 1:
                cl.fill = ALT_FILL
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


# ── 控制台表格 ────────────────────────────────────────
def disp_width(s) -> int:
    """CJK 双宽字符的显示宽度。"""
    return sum(2 if unicodedata.east_asian_width(c) in ("W", "F") else 1 for c in str(s))


def pad(s, w, align="left") -> str:
    s = str(s)
    gap = max(0, w - disp_width(s))
    return (" " * gap + s) if align == "right" else (s + " " * gap)


def print_table(headers, rows, widths=None, aligns=None):
    if widths is None:
        widths = [max([disp_width(h)] + [disp_width(r[i]) for r in rows]) + 1
                  for i, h in enumerate(headers)]
    aligns = aligns or ["left"] * len(headers)
    join = " "
    print("  " + join.join(pad(h, widths[i]) for i, h in enumerate(headers)).rstrip())
    print("  " + "-" * (sum(widths) + len(headers) - 1))
    for r in rows:
        cells = [pad(r[i] if i < len(r) else "", widths[i], aligns[i]) for i in range(len(headers))]
        print("  " + join.join(cells).rstrip())


def rule(title=""):
    head = f"── {title} " if title else ""
    print("\n" + head + "─" * max(4, 96 - disp_width(head)))


# ── 查询 ──────────────────────────────────────────────
def find_sales_orders_by_customer_code(code, base_url, key, secret):
    rows = paginated_get("Sales Order", base_url, key, secret,
                         child_filter("Sales Order Item", "customer_item_code", "like", f"%{code}%"),
                         SO_FIELDS, fallback_fields=["name", "status", "docstatus", "customer",
                                                     "delivery_date", "transaction_date"],
                         label="SO(客户码)")
    print(f"  客户码过滤: {len(rows)} 行")
    return rows


def find_sales_orders_by_items(item_codes, base_url, key, secret, customers=None):
    out = []
    for chunk in chunks(item_codes):
        filters = child_filter("Sales Order Item", "item_code", "in", chunk)
        if customers:
            filters.append(["customer", "in", list(customers)])
        out.extend(paginated_get("Sales Order", base_url, key, secret, filters, SO_FIELDS,
                                 fallback_fields=["name", "status", "docstatus", "customer",
                                                  "delivery_date", "transaction_date"],
                                 label="SO(物料码)"))
    print(f"  物料码过滤: {len(out)} 行")
    return out


def line_matches(line, item_codes, customer_code) -> bool:
    if code_in(line.get("item_code"), item_codes):
        return True
    return bool(customer_code) and code_in(line.get("customer_item_code"), [customer_code])


def fetch_so_lines(so_names, item_codes, customer_code, base_url, key, secret, sleep_s):
    """Q3: 逐单取明细, 返回命中的订单行。"""
    out = []
    for i, name in enumerate(so_names):
        d = get_single("Sales Order", name, base_url, key, secret)
        if i + 1 < len(so_names):
            time.sleep(sleep_s)
        if not d:
            continue
        for line in d.get("items", []):
            if line_matches(line, item_codes, customer_code):
                out.append({
                    "so": name,
                    "so_status": d.get("status"),
                    "so_docstatus": d.get("docstatus"),
                    "customer": d.get("customer_name") or d.get("customer"),
                    "customer_id": d.get("customer"),
                    "company": d.get("company"),
                    "so_delivery_date": d.get("delivery_date"),
                    "transaction_date": d.get("transaction_date"),
                    "amended_from": d.get("amended_from"),
                    "line_id": line.get("name"),
                    "item_code": line.get("item_code"),
                    "item_name": line.get("item_name"),
                    "customer_item_code": line.get("customer_item_code"),
                    "warehouse": line.get("warehouse") or d.get("set_warehouse"),
                    "uom": line.get("uom"),
                    "qty": num(line.get("qty")),
                    "delivered_qty": num(line.get("delivered_qty")),
                    "delivery_date": line.get("delivery_date") or d.get("delivery_date"),
                })
    return out


def fetch_delivery_notes(so_names, item_codes, customer_code, base_url, key, secret, sleep_s):
    """Q4: 出库单。必须按 item_code 匹配 —— 实测旧出库行 customer_item_code 为 None。"""
    raw = []
    for chunk in chunks(so_names):
        raw.extend(paginated_get(
            "Delivery Note", base_url, key, secret,
            [["Delivery Note Item", "against_sales_order", "in", chunk],
             ["Delivery Note Item", "item_code", "in", list(item_codes)]],
            ["name", "status", "docstatus", "posting_date", "customer", "company", "is_return"],
            fallback_fields=["name", "status", "docstatus", "posting_date"], label="DN"))
    dns = dedupe_by(raw, "name")
    print(f"  去重: {len(raw)} 行 → {len(dns)} 张 DN")

    so_set, item_set = set(so_names), set(item_codes)
    out = []
    for i, dn in enumerate(dns):
        d = get_single("Delivery Note", dn["name"], base_url, key, secret)
        if i + 1 < len(dns):
            time.sleep(sleep_s)
        for line in d.get("items", []):
            so = line.get("against_sales_order")
            if so not in so_set:
                continue
            by_item = line.get("item_code") in item_set
            by_code = code_in(line.get("customer_item_code"), [customer_code]) if customer_code else False
            if not (by_item or by_code):
                continue
            out.append({
                "so": so, "so_line_id": line.get("so_detail"), "dn": dn["name"],
                "dn_status": dn.get("status"), "dn_docstatus": dn.get("docstatus"),
                "dn_date": dn.get("posting_date"), "customer": dn.get("customer"),
                "company": dn.get("company"), "is_return": dn.get("is_return"),
                "warehouse": line.get("warehouse"), "item_code": line.get("item_code"),
                "item_name": line.get("item_name"),
                "customer_item_code": line.get("customer_item_code"),
                "qty": num(line.get("qty")), "uom": line.get("uom"),
                "rate": num(line.get("rate")), "amount": num(line.get("amount")),
                "match_by": "物料码" if by_item else "客户码",
            })
    return out


def fetch_work_orders(item_codes, base_url, key, secret):
    """Q5: 按 production_item 查本物料工单。

    不做 sales_order 侧并集 —— production_item 已覆盖全部相关工单
    (含无 SO 的 WO-26-00072 与指向集外 SO 的 WO-26-01242→SO-26-00045),
    而按 sales_order 并集会拉出 SO-26-00101 名下 308 张工单 (噪音 + 超 Request Line 限制)。
    """
    raw = []
    for chunk in chunks(item_codes):
        raw.extend(paginated_get("Work Order", base_url, key, secret,
                                 [["production_item", "in", chunk]], WO_FIELDS,
                                 fallback_fields=["name", "status", "docstatus", "sales_order",
                                                  "production_item", "qty", "produced_qty"],
                                 label="WO"))
    wos = dedupe_by(raw, "name")
    print(f"  去重: {len(raw)} 行 → {len(wos)} 张 WO")
    return wos


def fetch_work_order_details(wos, base_url, key, secret, sleep_s):
    """补 WO.operations 路由 (列表查询不含子表)。"""
    out = []
    for i, wo in enumerate(wos):
        d = get_single("Work Order", wo["name"], base_url, key, secret)
        if i + 1 < len(wos):
            time.sleep(sleep_s)
        merged = dict(wo)
        merged["operations"] = d.get("operations") or []
        merged["status"] = d.get("status") or wo.get("status")
        merged["produced_qty"] = num(d.get("produced_qty"))
        out.append(merged)
    return out


def fetch_job_cards(wo_names, base_url, key, secret):
    """Q8: 工序卡 = 真实进度。只查本物料的工单, 且分块 (约束 4)。"""
    out = []
    for chunk in chunks(wo_names):
        out.extend(paginated_get("Job Card", base_url, key, secret,
                                 [["work_order", "in", chunk]], JC_FIELDS,
                                 fallback_fields=["name", "work_order", "operation", "status"],
                                 label="JobCard"))
    return out


def fetch_production_plans(so_names, item_codes, base_url, key, secret, sleep_s):
    """Q6: 生产计划。必须同时按物料码过滤 po_items —— 光按 sales_order 会把
    同一张 PP 里属于其它物料的计划行也带出来。"""
    raw = []
    for chunk in chunks(so_names):
        raw.extend(paginated_get("Production Plan", base_url, key, secret,
                                 child_filter("Production Plan Item", "sales_order", "in", chunk),
                                 ["name", "status", "docstatus", "posting_date", "company"],
                                 fallback_fields=["name", "status", "docstatus", "posting_date"],
                                 label="PP"))
    pps = dedupe_by(raw, "name")
    print(f"  去重: {len(raw)} 行 → {len(pps)} 张 PP")

    so_set, item_set = set(so_names), set(item_codes)
    out = []
    for i, pp in enumerate(pps):
        d = get_single("Production Plan", pp["name"], base_url, key, secret)
        if i + 1 < len(pps):
            time.sleep(sleep_s)
        for row in d.get("po_items", []):
            so = row.get("sales_order")
            if row.get("item_code") not in item_set:
                continue
            if so and so not in so_set:
                continue
            out.append({
                "pp": pp["name"], "pp_status": d.get("status"), "pp_date": d.get("posting_date"),
                "item_code": row.get("item_code"), "item_name": row.get("item_name"),
                "planned_qty": num(row.get("planned_qty")), "produced_qty": num(row.get("produced_qty")),
                "pending_qty": num(row.get("pending_qty")), "so": so,
                "so_line_id": row.get("sales_order_item"), "work_order": row.get("work_order"),
                "warehouse": row.get("warehouse"),
            })
    return out


def resolve_customer_code(code, base_url, key, secret):
    """客户码 → 注册在哪些物料上 (走 Item 子表过滤, 实测可行)。

    注意: 客户码一般注册在 KS 成品上, 而订单行卖的是 PK# 皮壳 —— 两者不是同一个
    item_code。所以订单行物料的 customer_items 往往是空的, 这不是数据缺失。
    实测 CENKZ1325-Yellow-138 挂在 KS0001-DM-140-YELLOW, 订单行卖 PK#KS0001-DM-140-YELLOW。
    """
    rows = paginated_get("Item", base_url, key, secret,
                         child_filter("Item Customer Detail", "ref_code", "like", f"%{code}%"),
                         ["name", "item_name", "item_group", "variant_of"],
                         fallback_fields=["name", "item_name"], label="Item(客户码)")
    return dedupe_by(rows, "name")


def fetch_item_master(item_codes, base_url, key, secret):
    """Q7: 物料主档客户码 (这条路径不 403/417)。"""
    out = {}
    for code in item_codes:
        d = get_single("Item", code, base_url, key, secret)
        if not d:
            continue
        out[code] = {
            "item_name": d.get("item_name"), "item_group": d.get("item_group"),
            "stock_uom": d.get("stock_uom"), "variant_of": d.get("variant_of"),
            "ref_codes": [c.get("ref_code") for c in (d.get("customer_items") or []) if c.get("ref_code")],
        }
    return out


# ── 工序进度聚合 ──────────────────────────────────────
def jobcards_by_wo(job_cards):
    m = {}
    for j in job_cards:
        m.setdefault(j.get("work_order"), []).append(j)
    return m


def op_progress(wo, jcs):
    """单个工单的工序进度。

    done_qty 取"已完成件数"而非各工序完成量之和 —— 每道工序都有自己的工序卡,
    求和会把同一批件数按工序数重复累加 (实测 WO-26-02609 会算成 4×44=176)。
    首选第一道工序 (裁剪/开料) 的完成量, 退化时取各工序最大值。
    """
    routing = wo.get("operations") or []
    done_ops, open_ops = set(), set()
    last_end = ""

    def op_done_qty(opname):
        return sum(num(j.get("for_quantity")) for j in jcs
                   if j.get("operation") == opname and j.get("status") == "Completed")

    for j in jcs:
        op = j.get("operation")
        if j.get("status") == "Completed":
            done_ops.add(op)
            last_end = max(last_end, str(j.get("actual_end_date") or ""))
        else:
            open_ops.add(op)
    open_ops -= done_ops

    first_op = min(routing, key=lambda o: (o.get("sequence_id") or 0)) if routing else None
    done_qty = op_done_qty(first_op.get("operation")) if first_op else 0.0
    if done_qty <= 0:
        done_qty = max([op_done_qty(o.get("operation")) for o in routing] or [0.0])

    return {
        "total_ops": len(routing), "routing_ops": {o.get("operation") for o in routing},
        "done_ops": done_ops, "open_ops": open_ops,
        "done_jc": sum(1 for j in jcs if j.get("status") == "Completed"),
        "open_jc": sum(1 for j in jcs if j.get("status") != "Completed"),
        "done_qty": done_qty, "last_end": last_end,
        "jc_count": len(jcs),
        "any_pending_op": any((o.get("status") or "") not in ("Completed",) for o in routing),
    }


def progress_text(wos, jc_map) -> str:
    if not wos:
        return "-"
    parts = []
    for wo in wos:
        p = op_progress(wo, jc_map.get(wo["name"], []))
        if p["jc_count"] == 0:
            parts.append(f"{wo['name']} 未开工(0工序卡)")
        else:
            parts.append(f"{wo['name']} {len(p['done_ops'])}/{p['total_ops']}工序完({fmt_qty(p['done_qty'])}件)")
    return " | ".join(parts)


# ── 组装 ──────────────────────────────────────────────
def build_context(args, base_url, key, secret):
    customer_code = args.customer_code
    print(f"\n查询: {'--customer-code ' + customer_code if customer_code else '--item ' + args.item}")

    # ── 找 SO ──
    raw = []
    if customer_code:
        raw.extend(find_sales_orders_by_customer_code(customer_code, base_url, key, secret))
    else:
        item_codes = [c.strip() for c in args.item.split(",") if c.strip()]
        raw.extend(find_sales_orders_by_items(item_codes, base_url, key, secret))
    so_list = dedupe_by(raw, "name")
    if not so_list:
        return {"empty": True}

    # ── 订单行明细 ──
    # --item 模式: 物料码已知; --customer-code 模式: 由订单行反推
    item_codes = [] if customer_code else item_codes
    so_rows = fetch_so_lines([s["name"] for s in so_list], item_codes, customer_code,
                             base_url, key, secret, args.sleep)
    item_codes = sorted({r["item_code"] for r in so_rows if r["item_code"]})

    # 客户码模式: 同客户内按物料码补一轮, 抓订单行客户码为空的单
    customers = {s["customer"] for s in so_list if s.get("customer")}
    if customer_code and item_codes:
        extra = dedupe_by(find_sales_orders_by_items(item_codes, base_url, key, secret, customers), "name")
        have = {s["name"] for s in so_list}
        extra = [e for e in extra if e["name"] not in have]
        if extra:
            print(f"  物料码并集补充: {len(extra)} 张 SO ({', '.join(e['name'] for e in extra)})")
            so_list.extend(extra)
            so_rows.extend(fetch_so_lines([e["name"] for e in extra], item_codes, customer_code,
                                          base_url, key, secret, args.sleep))
        else:
            print("  物料码并集: 无新增 SO")

    item_codes = sorted({r["item_code"] for r in so_rows if r["item_code"]})
    so_names = sorted({r["so"] for r in so_rows})
    print(f"  SO 命中: {len(so_names)} 张, 订单行 {len(so_rows)} 行, 物料 {item_codes}")

    print("拉取物料主档…")
    item_master = fetch_item_master(item_codes, base_url, key, secret)
    master_items = (resolve_customer_code(customer_code, base_url, key, secret)
                    if customer_code else []).copy()
    if not master_items:
        master_items = [{"name": c, "item_name": item_master.get(c, {}).get("item_name"),
                         "item_group": item_master.get(c, {}).get("item_group")}
                        for c in item_codes]
    master_text = ", ".join(m["name"] for m in master_items)
    if customer_code:
        print(f"  客户码解析: {customer_code} → {master_text or '(未注册在任何物料上)'}")
    print("拉取出库单…")
    dn_rows = fetch_delivery_notes(so_names, item_codes, customer_code, base_url, key, secret, args.sleep)
    print("拉取生产工单…")
    wos = fetch_work_order_details(fetch_work_orders(item_codes, base_url, key, secret),
                                  base_url, key, secret, args.sleep)
    print("拉取工序卡…")
    jc_map = jobcards_by_wo(fetch_job_cards([w["name"] for w in wos], base_url, key, secret))
    print(f"  工序卡 {sum(len(v) for v in jc_map.values())} 张, 覆盖 {len(jc_map)} 张工单")
    print("拉取生产计划…")
    pp_rows = fetch_production_plans(so_names, item_codes, base_url, key, secret, args.sleep)

    # ── 每行补: 出库 / 工单 / 生产计划 ──
    wo_by_so = {}
    for w in wos:
        wo_by_so.setdefault(w.get("sales_order") or "", []).append(w)
    pp_by_line, pp_by_so = {}, {}
    for p in pp_rows:
        if p.get("so_line_id"):
            pp_by_line.setdefault(p["so_line_id"], []).append(p)
        if p.get("so"):
            pp_by_so.setdefault(p["so"], []).append(p)

    for r in so_rows:
        line_dns = [d for d in dn_rows
                    if d["so"] == r["so"] and (d["so_line_id"] == r["line_id"]
                                               or (not d["so_line_id"] and d["item_code"] == r["item_code"]))]
        r["dns"] = line_dns
        r["shipped_dn"] = sum(d["qty"] for d in line_dns if d["dn_docstatus"] == 1 and not d["is_return"])
        r["draft_dn"] = sum(d["qty"] for d in line_dns if d["dn_docstatus"] == 0)
        r["open_qty"] = r["qty"] - r["delivered_qty"]
        r["wos"] = wo_by_so.get(r["so"], [])
        r["pps"] = pp_by_line.get(r["line_id"]) or pp_by_so.get(r["so"], [])
        r["progress"] = progress_text(r["wos"], jc_map)
        r["master_item"] = master_text

    ctx = {
        "empty": False, "so_rows": so_rows, "dn_rows": dn_rows, "wos": wos, "jc_map": jc_map,
        "pp_rows": pp_rows, "item_master": item_master, "item_codes": item_codes,
        "master_items": master_items, "master_text": master_text,
        "customer_code": customer_code, "lead_months": args.lead_months,
        "dup_window": DUP_WINDOW_DAYS,
    }
    classify_anomalies(ctx)
    return ctx


def classify_anomalies(ctx):
    """异常判定 (顺序即优先级, 规则间有互斥关系)。"""
    today = date.today()
    lead = ctx["lead_months"]
    so_rows = ctx["so_rows"]
    anoms = ctx["anomalies"] = []

    def add(level, kind, key, msg):
        anoms.append({"level": level, "kind": kind, "key": key, "message": msg})

    amended_by = {r["amended_from"] for r in so_rows if r.get("amended_from")}

    totals = {"order_qty": 0.0, "valid_qty": 0.0, "cancelled_qty": 0.0, "voided_qty": 0.0,
              "shipped_qty": 0.0, "open_qty": 0.0, "dead_qty": 0.0,
              "in_progress_qty": 0.0, "draft_qty": 0.0}
    for r in so_rows:
        r["anomalies"] = []
        r["lead_date"] = None
        td = parse_date(r.get("transaction_date"))
        if td:
            r["lead_date"] = add_months(td, lead).isoformat()

        r["voided"] = (r["so_docstatus"] == 2 and r["so"] in amended_by)
        r["cancelled"] = (r["so_docstatus"] == 2 and not r["voided"])
        totals["order_qty"] += r["qty"]

        # 规则 2: 已改单 (作废, 不算未发)
        if r["voided"]:
            r["anomalies"].append("已改单(作废)")
            totals["voided_qty"] += r["qty"]
            add("info", "已改单", r["so"],
                f"{r['so']} 已取消并被后继单取代, 未发 {fmt_qty(r['open_qty'])} 不计入")
            continue
        # 规则 1': 单纯取消
        if r["cancelled"]:
            r["anomalies"].append("已取消")
            totals["cancelled_qty"] += r["qty"]
            if r["open_qty"] > 0:
                add("warn", "已取消未发", r["so"], f"{r['so']} 已取消, 未发 {fmt_qty(r['open_qty'])}")
            continue

        totals["valid_qty"] += r["qty"]
        totals["shipped_qty"] += r["delivered_qty"]
        totals["open_qty"] += r["open_qty"]
        totals["draft_qty"] += r["draft_dn"]

        if r["open_qty"] <= 0:
            continue

        # 规则 1: Closed 但未发完
        if r["so_status"] == "Closed":
            r["anomalies"].append("死单(Closed未发)")
            totals["dead_qty"] += r["open_qty"]
            add("error", "死单", r["so"],
                f"[死单] {r['so']} Closed 但未发 {fmt_qty(r['open_qty'])} "
                f"(交货日 {r['delivery_date']}, {progress_text(r['wos'], ctx['jc_map'])})")
            continue

        pros = [op_progress(w, ctx["jc_map"].get(w["name"], [])) for w in r["wos"]]
        started = any(p["jc_count"] > 0 for p in pros)
        not_started = bool(r["wos"]) and not started
        no_wo = not r["wos"]
        no_pp = not r["pps"]

        # 规则 6: 无工单无生产计划
        if no_wo and no_pp:
            r["anomalies"].append("无生产计划")
            add("error", "无生产计划", r["so"],
                f"[无生产计划] {r['so']} 未发 {fmt_qty(r['open_qty'])}, 无工单无生产计划 "
                f"(交货日 {r['delivery_date']})")
        # 规则 3/4: 工单未开工
        elif not_started:
            late = (parse_date(r["delivery_date"]) or today) < today
            r["anomalies"].append("逾期未开工" if late else "工单未开工")
            add("error" if late else "warn", "工单未开工", r["so"],
                f"[{'逾期未开工' if late else '工单未开工'}] "
                f"{', '.join(w['name'] for w in r['wos'])} 0 张工序卡, 工序全 Pending "
                f"(交货日 {r['delivery_date']})")

        # 规则 5: WO.status 未回写
        for w, p in zip(r["wos"], pros):
            if (w.get("status") in ("Not Started", "Draft", "Pending")
                    and p["jc_count"] > 0 and p["done_jc"] > 0):
                r["anomalies"].append("工单状态未回写")
                add("warn", "工单状态未回写", w["name"],
                    f"[工单状态未回写] {w['name']} 头部 {w.get('status')}/produced={fmt_qty(w.get('produced_qty'))}, "
                    f"但 {fmt_qty(p['done_qty'])} 件已完成 {'/'.join(sorted(p['done_ops']))}")

        # 规则 7: 超出交付约定
        ld = parse_date(r["lead_date"])
        if ld and today > ld:
            r["anomalies"].append("超出交付约定")
            add("warn", "超出交付约定", r["so"],
                f"[超出交付约定] {r['so']} 下单 {r['transaction_date']} + {lead}月 = {r['lead_date']}, 已逾期")

        if r["wos"]:
            totals["in_progress_qty"] += r["open_qty"]

    # 规则 8: 疑似重复下单 (同物料 + 同量 + 交货日近)
    for i, a in enumerate(so_rows):
        for b in so_rows[i + 1:]:
            if a["item_code"] != b["item_code"] or a["open_qty"] <= 0 or b["open_qty"] <= 0:
                continue
            if not close(a["qty"], b["qty"]):
                continue
            da, db = parse_date(a["delivery_date"]), parse_date(b["delivery_date"])
            if da and db and abs((da - db).days) <= ctx["dup_window"]:
                a["anomalies"].append("疑似重复下单")
                b["anomalies"].append("疑似重复下单")
                add("warn", "疑似重复下单", f"{a['so']}|{b['so']}",
                    f"[疑似重复] {a['so']} 与 {b['so']} 同物料同量 {fmt_qty(a['qty'])}, "
                    f"交货日差 {abs((da - db).days)} 天")

    ctx["totals"] = totals


# ── 输出 ──────────────────────────────────────────────
def render_console(ctx, args, base_url):
    so_rows, jc_map = ctx["so_rows"], ctx["jc_map"]
    rule("物料解析")
    if ctx["customer_code"]:
        print(f"  客户物料号: {ctx['customer_code']}")
        if ctx["master_items"]:
            for m in ctx["master_items"]:
                print(f"    → 注册物料: {m['name']}  {m.get('item_name') or ''}  [{m.get('item_group') or ''}]")
        else:
            print("    → ⚠ 未注册在任何物料上")
    print(f"  订单行实际物料: {', '.join(ctx['item_codes'])}")
    for code, m in ctx["item_master"].items():
        refs = ", ".join(m.get("ref_codes") or [])
        if not refs:
            refs = "（无 —— 客户码挂在成品 KS 上是常态，订单行卖的是 PK# 皮壳）"
        print(f"    {code}  {m.get('item_name') or ''}  物料自身客户码: {refs}")

    rule(f"订单发货状态 ({len(so_rows)} 行 / {len({r['so'] for r in so_rows})} 张 SO)")
    rows = []
    for r in sorted(so_rows, key=lambda x: (x["delivery_date"] or "", x["so"])):
        dn_txt = ", ".join(sorted({d["dn"] for d in r["dns"] if d["dn_docstatus"] == 1})) or "-"
        wo_txt = ", ".join(sorted({w["name"] for w in r["wos"]})) or "-"
        rows.append([r["so"], r["so_status"], r["delivery_date"] or "-", fmt_qty(r["qty"]),
                     fmt_qty(r["delivered_qty"]), fmt_qty(r["open_qty"]), dn_txt, wo_txt,
                     r["progress"]])
    print_table(["SO", "状态", "交货日期", "订单量", "已发", "未发", "出库单", "工单", "工序进度"],
                rows, aligns=["left", "left", "left", "right", "right", "right", "left", "left", "left"])

    t = ctx["totals"]
    print(f"\n  订单量 {fmt_qty(t['order_qty'])} = 有效 {fmt_qty(t['valid_qty'])}"
          + (f" + 已改单(作废) {fmt_qty(t['voided_qty'])}" if t["voided_qty"] else "")
          + (f" + 已取消 {fmt_qty(t['cancelled_qty'])}" if t["cancelled_qty"] else ""))
    print(f"  已发 {fmt_qty(t['shipped_qty'])} / 未发 {fmt_qty(t['open_qty'])}")
    dead = [r for r in so_rows if "死单(Closed未发)" in r["anomalies"]]
    if dead:
        print(f"    ├ 死单(Closed未发)  {fmt_qty(t['dead_qty'])}   "
              + " + ".join(f"{r['so']} {fmt_qty(r['open_qty'])}" for r in dead))
    prog = [r for r in so_rows if r["open_qty"] > 0 and "死单(Closed未发)" not in r["anomalies"]
            and not r["voided"] and not r["cancelled"]]
    if prog:
        print(f"    ├ 在产/待发        {fmt_qty(t['in_progress_qty'])}   "
              + " + ".join(f"{r['so']} {fmt_qty(r['open_qty'])}" for r in prog))
    if t["voided_qty"]:
        print(f"    ├ 已改单(作废)     {fmt_qty(t['voided_qty'])}")
    if t["cancelled_qty"]:
        print(f"    └ 已取消           {fmt_qty(t['cancelled_qty'])}")

    dn_open = sum(r["qty"] - r["shipped_dn"] for r in so_rows
                  if not r["voided"] and not r["cancelled"])
    mark = "✓" if close(dn_open, t["open_qty"]) else "✗"
    print(f"  DN 实算未发 {fmt_qty(dn_open)} vs ERP delivered_qty 口径 {fmt_qty(t['open_qty'])}  {mark}")

    lead = ctx["lead_months"]
    rule(f"交付预估 (下单 + {lead} 个月)")
    rows = []
    for r in so_rows:
        if r["open_qty"] <= 0 or r["voided"] or r["cancelled"]:
            continue
        note = r["progress"]
        if "无生产计划" in r["anomalies"]:
            note = "⚠ 无生产计划, 实际不会发"
        elif "死单(Closed未发)" in r["anomalies"]:
            note = "⚠ ERP 已 Closed, 不会发"
        rows.append([r["so"], r["transaction_date"] or "-", r["lead_date"] or "-",
                     f"{fmt_qty(r['open_qty'])} 件", note])
    if rows:
        print_table(["SO", "下单日期", f"预估可发(+{lead}月)", "未发", "工序/备注"], rows,
                    aligns=["left", "left", "left", "right", "left"])

    rule("异常")
    order = {"error": 0, "warn": 1, "info": 2}
    if ctx["anomalies"]:
        for a in sorted(ctx["anomalies"], key=lambda x: order.get(x["level"], 9)):
            print(f"  {a['message']}")
    else:
        print("  无")


def build_sheets(ctx):
    so_rows, jc_map = ctx["so_rows"], ctx["jc_map"]
    lead = ctx["lead_months"]

    s1 = [("SO单据号", 15), ("SO状态", 21), ("docstatus", 9), ("客户", 24), ("公司", 8),
          ("交货日期", 12), ("下单日期", 12), (f"预估可发(+{lead}月)", 15), ("订单行ID", 13),
          ("物料编码", 27), ("物料名称", 30), ("客户物料号(订单行)", 24), ("客户码注册物料", 26),
          ("发货仓", 18), ("订单量", 9), ("ERP已发", 9), ("ERP未发", 9), ("出库单号", 30),
          ("出库日期", 30), ("出库已发", 10), ("口径差异", 9), ("工单号", 22), ("工单状态", 18),
          ("工序进度", 44), ("生产计划", 22), ("计划待产", 10), ("异常标记", 30)]
    rows1 = []
    for r in sorted(so_rows, key=lambda x: (x["delivery_date"] or "", x["so"])):
        dns = [d for d in r["dns"] if d["dn_docstatus"] == 1]
        rows1.append([
            r["so"], r["so_status"], r["so_docstatus"], r["customer"], r["company"],
            r["delivery_date"], r["transaction_date"], r["lead_date"], r["line_id"],
            r["item_code"], r["item_name"], r["customer_item_code"], r["master_item"],
            r["warehouse"], r["qty"], r["delivered_qty"], r["open_qty"],
            ", ".join(sorted({d["dn"] for d in dns})), ", ".join(sorted({d["dn_date"] for d in dns})),
            r["shipped_dn"], "是" if not close(r["open_qty"], r["qty"] - r["shipped_dn"]) else "否",
            ", ".join(sorted({w["name"] for w in r["wos"]})),
            ", ".join(sorted({w.get("status") or "" for w in r["wos"]})),
            r["progress"],
            ", ".join(sorted({p["pp"] for p in r["pps"]})),
            sum(p["pending_qty"] for p in r["pps"]),
            " / ".join(r["anomalies"]),
        ])

    s2 = [("SO单据号", 15), ("SO行ID", 13), ("DN单据号", 15), ("DN状态", 12), ("docstatus", 9),
          ("DN日期", 12), ("客户", 20), ("公司", 8), ("出库仓", 18), ("物料编码", 27),
          ("物料名称", 30), ("客户物料号(DN行)", 24), ("数量", 9), ("单位", 7), ("单价", 9),
          ("金额", 11), ("计入已发", 9), ("匹配方式", 10)]
    rows2 = []
    for d in sorted(ctx["dn_rows"], key=lambda x: (x["dn_date"] or "", x["dn"], x["so"])):
        rows2.append([d["so"], d["so_line_id"], d["dn"], d["dn_status"], d["dn_docstatus"],
                      d["dn_date"], d["customer"], d["company"], d["warehouse"], d["item_code"],
                      d["item_name"], d["customer_item_code"], d["qty"], d["uom"], d["rate"],
                      d["amount"], "是" if d["dn_docstatus"] == 1 and not d["is_return"] else "否",
                      d["match_by"]])
    if rows2:
        rows2.append(["合计", "", "", "", "", "", "", "", "", "", "", "",
                      sum(d["qty"] for d in ctx["dn_rows"] if d["dn_docstatus"] == 1 and not d["is_return"]),
                      "", "", sum(d["amount"] for d in ctx["dn_rows"]), "", ""])

    s3 = [("WO单据号", 15), ("WO状态", 12), ("WO数量", 9), ("WO已产", 9), ("计划开工", 12),
          ("关联SO", 15), ("物料编码", 27), ("生产计划", 15), ("工序名", 20), ("工序序号", 9),
          ("工序状态", 12), ("工序卡号", 34), ("批次量", 9), ("工序完工时间", 20), ("备注", 24)]
    # PP ↔ WO 的关联: 这些生产计划行的 work_order 字段是空的, 所以还要按
    # (sales_order, item_code) 兜一层, 否则工序表里看不到所属生产计划。
    pp_by_wo, pp_by_so_item = {}, {}
    for p in ctx["pp_rows"]:
        if p.get("work_order"):
            pp_by_wo.setdefault(p["work_order"], []).append(p["pp"])
        if p.get("so") and p.get("item_code"):
            pp_by_so_item.setdefault((p["so"], p["item_code"]), []).append(p["pp"])

    def pp_of(wo):
        names = set(pp_by_wo.get(wo["name"], []))
        names |= set(pp_by_so_item.get((wo.get("sales_order"), wo.get("production_item")), []))
        return ", ".join(sorted(names))

    rows3 = []
    for w in sorted(ctx["wos"], key=lambda x: x["name"]):
        jcs = jc_map.get(w["name"], [])
        routing = w.get("operations") or []
        if not routing:
            rows3.append([w["name"], w.get("status"), w["qty"], w.get("produced_qty"),
                          str(w.get("planned_start_date") or "")[:10], w.get("sales_order") or "",
                          w.get("production_item"), pp_of(w),
                          "-", "", "", "", "", "", "无工序路由"])
            continue
        for op in sorted(routing, key=lambda o: (o.get("sequence_id") or 0, o.get("operation") or "")):
            opname = op.get("operation")
            op_jcs = [j for j in jcs if j.get("operation") == opname]
            notes = []
            if not jcs:
                notes.append("未开工(0工序卡)")
            elif not op_jcs:
                notes.append("无工序卡")
            if not w.get("sales_order"):
                notes.append("工单无SO")
            if w.get("status") == "Cancelled":
                notes.append("工单已取消")
            if num(w.get("produced_qty")) > num(w.get("qty")):
                notes.append("已产>计划(超产)")
            rows3.append([
                w["name"], w.get("status"), w["qty"], w.get("produced_qty"),
                str(w.get("planned_start_date") or "")[:10], w.get("sales_order") or "",
                w.get("production_item"), pp_of(w),
                opname, op.get("sequence_id"), op.get("status"),
                ", ".join(sorted(j["name"] for j in op_jcs)),
                sum(num(j.get("for_quantity")) for j in op_jcs),
                max([str(j.get("actual_end_date") or "")[:19] for j in op_jcs] or [""]),
                " / ".join(notes),
            ])
    return (s1, rows1), (s2, rows2), (s3, rows3)


# ── fixture 断言 ──────────────────────────────────────
FIXTURE = {
    "customer_code": "CENKZ1325-Yellow-138",
    "item": "PK#KS0001-DM-140-YELLOW",
    # ── --customer-code CENKZ1325-Yellow-138 (10 张 SO / 330 件) ──
    "so_count": 10,
    "order_qty": 330, "valid_qty": 310, "voided_qty": 20, "cancelled_qty": 0,
    "shipped_qty": 126, "open_qty": 184, "dead_qty": 64, "in_progress_qty": 120,
    "dns": {"DN-25-00043": 20, "DN-25-00138": 60, "DN-26-00003": 26, "DN-26-00039": 20},
    "wo_count": 18,
    # 待产的生产计划 (其余为历史计划, 同一 PP 会覆盖多个 SO)
    "pp_names_pending": {"PP-26-00033", "PP-26-00036"},
    "dead_sos": {"SO-25-00198", "SO-26-00003", "SO-26-00099"},
    "voided_sos": {"SO-26-00053"},
    # 客户码注册在成品 KS 上, 订单行卖的是 PK# 皮壳 —— 两者不是同一 item_code
    "master_item": "KS0001-DM-140-YELLOW",
    "lead": {"SO-26-00101": "2026-11-17", "SO-26-00110": "2026-12-02",
             "SO-26-00099": "2026-11-13", "SO-26-00003": "2026-04-09"},
    # ── 工序卡 (真实进度, WO.status 不可信) ──
    # 生产在进行中, 所以"已完成工序"只做下界断言 —— 实测 2026-09-14 当天
    # 翻面 就从 Pending 变 Completed。
    "jc_done_ops_min": {"裁剪", "皮壳整件", "锁扣眼", "拷边"},
    "jc_not_done": {"质检"},
    "jc_done_qty": 44,
    "jc_not_started_wo": "WO-26-03264",
}


def assert_fixture(ctx):
    f = FIXTURE
    t = ctx["totals"]
    checks = []

    def ck(name, got, want):
        ok = close(got, want) if isinstance(want, (int, float)) else got == want
        checks.append((ok, name, got, want))

    ck("SO 张数", len({r["so"] for r in ctx["so_rows"]}), f["so_count"])
    for k in ("order_qty", "valid_qty", "voided_qty", "cancelled_qty", "shipped_qty",
              "open_qty", "dead_qty", "in_progress_qty"):
        ck(f"合计 {k}", t[k], f[k])
    for dn, q in f["dns"].items():
        ck(f"{dn} 已发量", sum(d["qty"] for d in ctx["dn_rows"]
                              if d["dn"] == dn and d["dn_docstatus"] == 1), q)
    ck("工单张数", len(ctx["wos"]), f["wo_count"])
    ck("生产计划 ⊇ 待产两张", f["pp_names_pending"] <= {p["pp"] for p in ctx["pp_rows"]}, True)
    ck("死单集合", {r["so"] for r in ctx["so_rows"] if "死单(Closed未发)" in r["anomalies"]}, f["dead_sos"])
    ck("作废集合", {r["so"] for r in ctx["so_rows"] if r.get("voided")}, f["voided_sos"])
    ck("客户码注册物料", {m["name"] for m in ctx["master_items"]}, {f["master_item"]})
    for so, want in f["lead"].items():
        got = next((r["lead_date"] for r in ctx["so_rows"] if r["so"] == so), None)
        ck(f"{so} 预估可发", got, want)
    jc = ctx["jc_map"].get("WO-26-02609", [])
    _done = {j["operation"] for j in jc if j["status"] == "Completed"}
    ck("WO-26-02609 已完成工序 ⊇ 前4道", f["jc_done_ops_min"] <= _done, True)
    ck("WO-26-02609 质检仍未完成", not (f["jc_not_done"] & _done), True)
    _wo = next((w for w in ctx["wos"] if w["name"] == "WO-26-02609"), None)
    ck("WO-26-02609 完成件数(非各工序求和)", op_progress(_wo, jc)["done_qty"] if _wo else 0, f["jc_done_qty"])
    ck(f"{f['jc_not_started_wo']} 工序卡数", len(ctx["jc_map"].get(f["jc_not_started_wo"], [])), 0)

    print("\n── fixture 断言 ──")
    bad = 0
    for ok, name, got, want in checks:
        if not ok:
            bad += 1
        print(f"  [{'OK' if ok else 'FAIL'}] {name}: {got!r} (期望 {want!r})")
    print(f"  {len(checks) - bad}/{len(checks)} 通过")
    return bad == 0


# ── main ──────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(
        description="客户物料号 / EN 物料号 → 销售订单发货状态 + 生产/工序进度",
        epilog="示例:\n"
               "  %(prog)s --customer-code CENKZ1325-Yellow-138\n"
               '  %(prog)s --item "PK#KS0001-DM-140-YELLOW" --no-excel\n'
               "  %(prog)s --customer-code X --test\n"
               "  %(prog)s --customer-code X -o EN_API/out/custom.xlsx",
        formatter_class=argparse.RawDescriptionHelpFormatter)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--customer-code", help="客户物料号 (模糊匹配, 大小写不敏感)")
    group.add_argument("--item", help="EN 物料编码, 逗号分隔多个")
    parser.add_argument("--test", action="store_true", help="使用测试系统 (ensh.vilavi.cn)")
    parser.add_argument("--output", "-o", help="输出路径 (默认 EN_API/out/ 下)")
    parser.add_argument("--no-excel", action="store_true", help="只打控制台, 不生成 Excel")
    parser.add_argument("--env-file", help="指定 .env 路径 (默认自动探测)")
    parser.add_argument("--lead-months", type=int, default=3,
                        help="交付约定: 下单日期 + N 个月 (默认 3, 业务口径非系统字段)")
    parser.add_argument("--sleep", type=float, default=0.15, help="单文档请求间隔秒数 (默认 0.15)")
    parser.add_argument("--assert-fixture", action="store_true", help="跑已知样例断言 (回归自检)")
    args = parser.parse_args()

    env = "test" if args.test else "prod"
    key, secret = load_credentials(env, resolve_env_files(args.env_file))
    if not key or not secret or "your_" in key:
        print("✗ API 凭证未配置。请先检查 EN_API/.env 文件", file=sys.stderr)
        sys.exit(1)

    base_url = ENV_URLS[env]
    print(f"系统: {base_url} (env: {env})")

    ctx = build_context(args, base_url, key, secret)
    if ctx.get("empty"):
        print("\n无匹配的销售订单。")
        sys.exit(0)

    render_console(ctx, args, base_url)

    if args.assert_fixture:
        if not (args.customer_code == FIXTURE["customer_code"] or args.item == FIXTURE["item"]):
            print("\n⚠ --assert-fixture 只对 FIXTURE 中的物料有效, 跳过。", file=sys.stderr)
            sys.exit(2)
        sys.exit(0 if assert_fixture(ctx) else 1)

    if args.no_excel:
        return

    print("\n生成 Excel…")
    wb = Workbook()
    (h1, r1), (h2, r2), (h3, r3) = build_sheets(ctx)
    _write_sheet(wb, "订单发货状态", h1, r1, group_col=0, first=True,
                 number_cols=(15, 16, 17, 20, 26))
    _write_sheet(wb, "出库明细", h2, r2, group_col=0, number_cols=(13, 15, 16))
    _write_sheet(wb, "工单与工序", h3, r3, group_col=0, number_cols=(3, 4, 13))

    if args.output:
        output = Path(args.output)
    else:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        src = args.customer_code or args.item
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        output = DATA_DIR / f"{slugify(src)}_发货状态_{ts}.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    print(f"OK: {output}")
    print(f"  Sheet1 订单发货状态 {len(r1)} 行, Sheet2 出库明细 {len(r2)} 行, Sheet3 工单与工序 {len(r3)} 行")


if __name__ == "__main__":
    main()
