"""EN 生产 BOM 成本 → item-group-browser「BOM成本概览」Excel（三角/平条/弧形靠枕）

数据源 = ERPNext 生产报表「BOM Cost List V2」(work_order_task; prepared 报表需 ignore_prepared_report 同步执行)。
展示结构复刻 vilavi_pim item-group-browser 成本信息页签的 BOM成本概览表（不含产品图片列）：
  - 两行分组表头：产品编号/名称/面料/尺寸/发货方式/成品实际重量/皮壳单重/体积/FBM出厂价 / [美中USTX·美东USNJ·波兰PL] 各(头程运费+海外生产费用) / FBA出厂价 / USFBA·DEFBA·UKFBA·JPFBA 头程运费
  - 面料、尺寸从中文品名提取（面料=尺寸前一段，尺寸=倒数第二段）；皮壳单重=报表 cover_weight(g)
  - FBM出厂价=按绍兴发货方式取 皮壳cost_fg | 半成品cost_fg+半成品包装成本 | 成品sx_cost_all
  - 头程运费=按发货方式取 皮壳cover_freight / 半成品sfg_freight / 成品freight（对应仓）
  - 海外生产费用=对应仓估值(valuation_rate)；FBA出厂价=绍兴总成本(sx_cost_all)
仅保留 三角靠枕 / 平条靠枕 / 弧形靠枕 三类成品，并按产品编号(item_fg)去重取首行。

用法:
    uv run python EN_API/bom_cost_list_excel.py                # 默认 prod 联网拉最新
    uv run python EN_API/bom_cost_list_excel.py --use-local 数据源/xxx.json.gz   # 读本地快照不联网
凭证: EN_API/.env 的 PROD_ERP_API_KEY/SECRET (同 customs_export.py)。输出 → EN_API/out/
"""

from __future__ import annotations

import argparse
import datetime
import gzip
import json
import re
import sys
from collections import Counter
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

_DIR = Path(__file__).resolve().parent
DATA_DIR = _DIR / "数据源"
OUT_DIR = _DIR / "out"

ENV_URLS = {"test": "https://ensh.vilavi.cn", "prod": "https://erpnext.vilavi.cn"}
REPORT_NAME = "BOM Cost List V2"

# 与 item-group-browser(item_group_cost_service) 同源；多一个 simplified_column_view=0 以取全列
REPORT_FILTERS = {
    "item_group": "产品",
    "show_disabled": 1,
    "show_ref_code": 1,
    "sum_columns_at_end": 1,
    "pllc_sfg_missing_use_cover": 0,
    "simplified_column_view": 0,
}
REQUIRED_COLS = ["item_fg", "item_fg_name", "shipping_method", "cost_fg", "sx_cost_all"]

GROUP_DESTS = ["USTX", "USNJ", "PL"]                      # 分组：头程运费 + 海外生产费用
FBA_FREIGHT_DESTS = ["USFBA", "DEFBA", "UKFBA", "JPFBA"]  # 单列头程运费
DEST_LABELS = {"USTX": "美中 USTX", "USNJ": "美东 USNJ", "PL": "波兰 PL"}
VALUATION_FIELD = {"USTX": "ustx_valuation_rate", "USNJ": "usnj_valuation_rate", "PL": "pl_valuation_rate"}

# 概览叶子列规划：单列(single) 或 分组(组内两列：g_freight/g_val)
# 「面料/尺寸」从产品名称解析，紧随产品名称之后。
BASE_LEAVES = [
    ("item_fg", "产品编号"),
    ("item_fg_name", "产品名称"),
    ("fabric", "面料"),
    ("size", "尺寸"),
    ("shipping_method", "绍兴发货方式"),
    ("fg_actual_weight_kg", "成品实际重量\n(kg)"),
    ("cover_weight", "皮壳单重\n(g)"),
    ("package_volume", "单个包装体积\n(cm³)"),
    ("fbm_factory_price", "FBM出厂价"),
]


def build_plan() -> list[dict]:
    """叶子列顺序（不含图片列）。entry: kind in single/g_freight/g_val, 附带 label/fn。"""
    plan: list[dict] = []
    for fn, label in BASE_LEAVES:
        plan.append({"kind": "single", "field": fn, "label": label})
    for dest in GROUP_DESTS:
        plan.append({"kind": "g_freight", "dest": dest})
        plan.append({"kind": "g_val", "dest": dest})
    plan.append({"kind": "single", "field": "fba_factory_price", "label": "FBA出厂价"})
    for dest in FBA_FREIGHT_DESTS:
        plan.append({"kind": "single", "field": f"freight_{dest}", "label": f"头程运费\n{dest}", "dest": dest})
    return plan


def load_credentials(env: str) -> tuple[str, str]:
    vals: dict[str, str] = {}
    env_file = _DIR / ".env"
    if env_file.is_file():
        for line in env_file.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, _, v = line.partition("=")
            vals[k.strip()] = v.strip()
    key = vals.get("PROD_ERP_API_KEY" if env == "prod" else "TEST_ERP_API_KEY", "")
    sec = vals.get("PROD_ERP_API_SECRET" if env == "prod" else "TEST_ERP_API_SECRET", "")
    return key or vals.get("ERP_API_KEY", ""), sec or vals.get("ERP_API_SECRET", "")


def fetch_report(env: str, key: str, sec: str) -> dict:
    base = ENV_URLS[env]
    url = f"{base}/api/method/frappe.desk.query_report.run"
    params = {
        "report_name": REPORT_NAME,
        "filters": json.dumps(REPORT_FILTERS),
        "ignore_prepared_report": 1,   # prepared 报表改同步执行，直接返回 result/columns
    }
    print(f"GET {url}  (env={env}, report={REPORT_NAME}, {len(REPORT_FILTERS)} filters)")
    r = requests.get(url, params=params, headers={"Authorization": f"token {key}:{sec}"}, timeout=300)
    if r.status_code != 200:
        raise SystemExit(f"✗ 报表请求失败 HTTP {r.status_code}: {r.text[:500]}")
    body = r.json()
    payload = body.get("message") if isinstance(body, dict) and isinstance(body.get("message"), dict) else body
    if not isinstance(payload, dict) or "result" not in payload or "columns" not in payload:
        raise SystemExit(f"✗ 响应结构异常，缺少 columns/result: {list(body.keys()) if isinstance(body, dict) else type(body)}")
    return payload


def load_local(path: Path) -> dict:
    with gzip.open(path, "rt", encoding="utf-8") as f:
        payload = json.load(f)
    print(f"读本地快照: {path}")
    return payload


def validate(payload: dict) -> None:
    cols = payload["columns"]
    fieldnames = [c.get("fieldname") for c in cols]
    if not all(fieldnames):
        raise SystemExit(f"✗ columns 缺 fieldname: {fieldnames[:10]}")
    rows = payload["result"]
    missing = [c for c in REQUIRED_COLS if c not in fieldnames]
    if missing or not rows:
        raise SystemExit(f"✗ 缺少关键列或空结果: missing={missing}, rows={len(rows)}")
    print(f"✓ columns={len(cols)}  rows={len(rows)}")


def classify(name: str) -> str | None:
    if not name:
        return None
    if "三角" in name and "靠枕" in name:
        return "三角靠枕"
    if "平条" in name and "靠枕" in name:
        return "平条靠枕"
    if ("弧形" in name or "月牙" in name or "圆弧" in name) and "靠枕" in name:
        return "弧形靠枕"
    return None


_SIZE_RE = re.compile(r"^\d[\d.]*([xX*][\d.]+)*\s*(cm)?$")


def extract_fabric_size(name: str) -> tuple[str, str]:
    """从中文品名提取 (面料, 尺寸)：形如 三角靠枕-荷兰绒-153x50x24cm-蓝色。
    从右向左找第一个尺寸段（纯数字 / a x b x c 带 cm），其左侧一段即面料。"""
    if not name:
        return "", ""
    segs = [s for s in name.split("-") if s]
    idx = None
    for j in range(len(segs) - 1, -1, -1):
        if _SIZE_RE.fullmatch(segs[j]):
            idx = j
            break
    if idx is None or idx == 0:
        return "", ""
    return segs[idx - 1], segs[idx]


def audit_near_misses(rows: list[dict], matched: set[int]) -> list[dict]:
    """未命中但品名含相关词的行：返回审计清单（防静默丢弃）。"""
    key_words = ("靠枕", "弧形", "月牙", "圆弧", "三角", "平条")
    near: list[dict] = []
    seen = set()
    for i, r in enumerate(rows):
        if i in matched:
            continue
        name = (r.get("item_fg_name") or "").strip()
        code = (r.get("item_fg") or "").strip()
        if not name or not any(w in name for w in key_words):
            continue
        if code not in seen:
            seen.add(code)
            near.append({"item_fg": code, "item_fg_name": name})
    print(f"⚠ 审计: 含相关词但未命中 {len(near)} 行（样例）")
    for x in near[:12]:
        print(f"    {x['item_fg']}  {x['item_fg_name'][:40]}")
    return near


# ──────────────────────────────────────────────
# 概览表取数（复刻 item_group_cost_service.py 逻辑）
# ──────────────────────────────────────────────
def build_dest_freight_map(payload: dict) -> dict[str, str]:
    """按 label 含目的地 → 识别各仓 freight 字段。返回 {dest: freight_fieldname}"""
    dest_field: dict[str, str] = {}
    for dest in [*GROUP_DESTS, *FBA_FREIGHT_DESTS]:
        for c in payload["columns"]:
            fn = c.get("fieldname") or ""
            label = c.get("label") or ""
            if not fn.startswith("freight_") or not _label_matches(label, dest):
                continue
            dest_field[dest] = fn
            break
    return dest_field


def _label_matches(label: str, dest: str) -> bool:
    if dest == "USTX":
        return ("USTX" in label or "美中" in label) and "USNJ" not in label
    if dest == "USNJ":
        return "USNJ" in label
    if dest == "PL":
        return "PL" in label or "波兰" in label
    return dest in label  # USFBA/DEFBA/UKFBA/JPFBA


def freight_value(row: dict, dest: str, dest_freight: dict[str, str], method: str) -> float:
    """按绍兴发货方式取该仓头程运费：皮壳→cover_freight_*，半成品→sfg_freight_*，成品→freight_*。"""
    base_fn = dest_freight.get(dest)
    if not base_fn or not method:
        return 0.0
    suffix = base_fn[len("freight_"):]
    if "皮壳" in method:
        field = f"cover_freight_{suffix}"
    elif "半成品" in method:
        field = f"sfg_freight_{suffix}"
    elif "成品" in method:
        field = base_fn
    else:
        return 0.0
    return float(row.get(field) or 0)


def fbm_factory_price(row: dict) -> float:
    """FBM出厂价：皮壳→cost_fg；半成品→cost_fg+绍兴包装半成品成本；成品→sx_cost_all。"""
    method = (row.get("shipping_method") or "").strip()
    if not method:
        return 0.0
    cost_fg = float(row.get("cost_fg") or 0)
    if "皮壳" in method:
        return cost_fg
    if "半成品" in method:
        return cost_fg + float(row.get("cost_sxbzbcp") or 0)
    if "成品" in method:
        return float(row.get("sx_cost_all") or 0)
    return 0.0


def project_row(row: dict, dest_freight: dict[str, str]) -> list:
    """按概览叶子列规划取值（顺序见 build_plan）。"""
    method = (row.get("shipping_method") or "").strip()
    cells: list = []
    for entry in build_plan():
        kind = entry["kind"]
        if kind == "single":
            field = entry["field"]
            if "dest" in entry:                       # FBA 各仓单列头程运费
                cells.append(freight_value(row, entry["dest"], dest_freight, method))
            elif field == "fbm_factory_price":
                cells.append(fbm_factory_price(row))
            elif field == "fba_factory_price":
                cells.append(float(row.get("sx_cost_all") or 0))
            elif field in ("fabric", "size"):
                fabric, size = extract_fabric_size(row.get("item_fg_name") or "")
                cells.append(fabric if field == "fabric" else size)
            elif field in ("fg_actual_weight_kg", "package_volume", "cover_weight"):
                cells.append(row.get(field))
            else:
                cells.append(row.get(field) or "")
        elif kind == "g_freight":
            cells.append(freight_value(row, entry["dest"], dest_freight, method))
        else:  # g_val
            cells.append(float(row.get(VALUATION_FIELD[entry["dest"]]) or 0))
    return cells


def export_overview(rows: list[list], fam_of: list[str], near_miss: list[dict], out: Path) -> None:
    plan = build_plan()
    n_cols = len(plan)

    wb = Workbook()
    ws = wb.active
    ws.title = "BOM成本概览-三种靠枕"

    head_fill = PatternFill("solid", fgColor="DDEBF7")
    group_fill = PatternFill("solid", fgColor="BDD7EE")

    # 分组起始列（1-based）与单列纵向合并
    group_start: dict[str, int] = {}
    single_cols: list[int] = []
    for i, e in enumerate(plan):
        col = i + 1
        if e["kind"] == "single":
            ws.cell(row=1, column=col, value=e["label"])
            single_cols.append(col)
        elif e["kind"] == "g_freight":
            group_start[e["dest"]] = col
    for dest in GROUP_DESTS:
        c0 = group_start[dest]
        ws.cell(row=1, column=c0, value=DEST_LABELS[dest])
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + 1)
        for cc in (c0, c0 + 1):
            ws.cell(row=1, column=cc).fill = group_fill
        ws.cell(row=2, column=c0, value="头程运费")
        ws.cell(row=2, column=c0 + 1, value="海外生产费用")
    # 单列：纵向合并两行表头
    for col in single_cols:
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
        ws.cell(row=1, column=col).fill = head_fill

    # 表头样式
    for col in range(1, n_cols + 1):
        for row in (1, 2):
            c = ws.cell(row=row, column=col)
            c.font = Font(bold=True, size=9)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "A3"

    # 数据
    start_row = 3
    for cells in rows:
        for ci, v in enumerate(cells, start=1):
            ws.cell(row=start_row, column=ci, value=v)
        start_row += 1

    # 数字格式：金额列带 ¥ 千分位；重量/单重/体积保持默认(General)，避免整数值出现尾随小数点
    for r in range(3, start_row):
        for ci in range(1, n_cols + 1):
            e = plan[ci - 1]
            c = ws.cell(row=r, column=ci)
            if not isinstance(c.value, (int, float)):
                continue
            is_money = e["kind"] in ("g_freight", "g_val") or (
                e["kind"] == "single" and (e["field"] in ("fbm_factory_price", "fba_factory_price") or e["field"].startswith("freight_"))
            )
            if is_money:
                c.number_format = '"¥"#,##0.00'

    # 列宽
    widths = {1: 24, 2: 32, 3: 12, 4: 16, 5: 11, 6: 10, 7: 11, 8: 12, 9: 11,
              10: 10, 11: 12, 12: 10, 13: 12, 14: 10, 15: 12, 16: 11,
              17: 12, 18: 12, 19: 12, 20: 12}
    for ci, w in widths.items():
        ws.column_dimensions[get_column_letter(ci)].width = w

    # 统计 sheet（含审计，不丢记录）
    ws2 = wb.create_sheet("筛选统计")
    counts = Counter(fam_of)
    ws2.append(["家族", "去重后产品数"])
    for k, v in counts.items():
        ws2.append([k, v])
    ws2.append(["总计(去重后)", sum(counts.values())])
    ws2.append([])
    ws2.append(["未命中但含相关词（审计，未纳入导出）", ""])
    for x in near_miss:
        ws2.append([x["item_fg"], x["item_fg_name"]])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    print(f"\n✓ 导出完成: {out}")
    print(f"  去重后产品数={len(rows)}  家族分布: " + "  ".join(f"{k}={v}" for k, v in counts.items()))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--env", choices=("test", "prod"), default="prod")
    ap.add_argument("--use-local", metavar="JSONGZ", default=None, help="读本地 gzip JSON 快照，不联网")
    ap.add_argument("--out", default=None, help="输出 xlsx 路径（默认 EN_API/out/ 自动命名）")
    args = ap.parse_args()

    if args.use_local:
        payload = load_local(Path(args.use_local))
    else:
        key, sec = load_credentials(args.env)
        if not key or not sec:
            raise SystemExit(f"✗ {_DIR / '.env'} 缺少 {args.env.upper()} 凭证")
        payload = fetch_report(args.env, key, sec)
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        ts = datetime.datetime.now().strftime("%Y-%m-%d-%H-%M")
        raw = DATA_DIR / f"bom_cost_list_v2_{ts}.json.gz"
        with gzip.open(raw, "wt", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False)
        print(f"原始报表快照已存: {raw}")

    validate(payload)

    rows = payload["result"]
    fams, matched = [], set()
    for i, r in enumerate(rows):
        fam = classify((r.get("item_fg_name") or "").strip())
        fams.append(fam)
        if fam:
            matched.add(i)

    near_miss = audit_near_misses(rows, matched)

    # 去重：同一产品编号只保留首次出现的行
    kept, fam_of = [], []
    seen = set()
    for r, f in zip(rows, fams):
        if not f:
            continue
        code = (r.get("item_fg") or "").strip()
        if not code or code in seen:
            continue
        seen.add(code)
        kept.append(r)
        fam_of.append(f)

    if not kept:
        raise SystemExit("✗ 三种靠枕未命中任何行，请核对品名关键词口径。")

    dest_freight = build_dest_freight_map(payload)
    projected = [project_row(r, dest_freight) for r in kept]

    out = Path(args.out) if args.out else OUT_DIR / f"BOM成本概览_三角平条弧形靠枕_{datetime.datetime.now():%Y%m%d_%H%M}.xlsx"
    export_overview(projected, fam_of, near_miss, out)


if __name__ == "__main__":
    main()
