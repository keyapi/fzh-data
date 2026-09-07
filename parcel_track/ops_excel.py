"""运营异常 Excel 写出（FedEx 单承运商报表与混合报表共用）。"""

from __future__ import annotations

from parcel_track.classify import (
    CLASS,
    HANDLING_DAYS,
    MISSING_AFTER_DAYS,
    STUCK_DAYS,
    TRANSIT_SLOW_DAYS,
)

COLOR = {
    "green": "C6EFCE", "yellow": "FFEB9C", "orange": "FCD5B4", "red": "FFC7CE",
    "gray": "D9D9D9", "blue": "DDEBF7", "dark": "404040", "header": "2F5597",
}
LEVEL_COLOR = {
    "漏发/未交接": "red", "卡件": "red", "严重延误": "red", "重度迟发": "red",
    "中度迟发": "orange", "FedEx延误": "orange", "承运延误": "orange",
    "轻度迟发": "yellow", "建标未收件": "yellow", "在途": "yellow",
    "已取消": "gray", "数据异常/查无": "gray", "复用旧票(缺建标)": "blue",
    "未支持/停放": "gray",
    "正常交付": "green", "准时": "green",
}
SLOW_KEYS = ("fedex_slow", "carrier_slow")
ANOMALY_CATS = [
    "missing_not_handed", "fresh_no_pickup", "late_handover",
    "fedex_slow", "carrier_slow", "stuck", "cancelled", "not_found",
]
DETAIL_COLS = [
    "跟踪号", "承运商", "分类(EN)", "分类(中文)", "等级", "订单号", "包裹号", "渠道账号", "渠道",
    "平台SKU", "通途SKU", "产品名称", "发货仓库", "执行发货人", "是否补发货", "销售站点",
    "买家姓名", "国家", "城市", "省/州", "发货日期", "建标时间", "站点收件时间", "交付时间",
    "日历延迟(天)", "营业日延迟(天)", "承运延迟(营业日)", "Amazon是否判迟", "所属Sheet", "建议动作", "处理状态",
]


def write_ops_workbook(df, out_xlsx: str, *, title: str, slow_label: str, notes_title: str, parked=None):
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if "_key" not in df.columns:
        raise ValueError("df 缺少 _key")
    parked = parked if parked is not None else pd.DataFrame()

    def pick(*keys):
        return df[df["_key"].isin(keys)]

    wb = Workbook()
    thin = Border(*[Side(style="thin", color="BFBFBF")] * 4)
    ws = wb.active
    ws.title = "总览"
    ws.cell(row=1, column=1, value=title).font = Font(size=16, bold=True, color="FFFFFF")
    ws.cell(row=1, column=1).fill = PatternFill("solid", fgColor="404040")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    n_total = len(df) + len(parked)
    n_anom = len(pick(*ANOMALY_CATS))
    kpi = [
        ("总单(票)", n_total), ("异常(票)", n_anom),
        ("漏发/未交接", int(sum(df["_key"] == "missing_not_handed"))),
        ("迟发", int(sum(df["_key"] == "late_handover"))),
        (slow_label, int(sum(df["_key"].isin(SLOW_KEYS)))),
        ("卡件", int(sum(df["_key"] == "stuck"))),
        ("已取消", int(sum(df["_key"] == "cancelled"))),
        ("停放", len(parked)),
    ]
    for i, (k, v) in enumerate(kpi):
        col = 1 + i * 2
        ws.cell(row=3, column=col, value=k).font = Font(bold=True)
        c = ws.cell(row=4, column=col, value=v)
        c.font = Font(size=14, bold=True)
        c.alignment = Alignment(horizontal="center")
        if k in ("漏发/未交接", "卡件"):
            c.fill = PatternFill("solid", fgColor="FFC7CE")
    r = 6
    ws.cell(row=r, column=1, value="分类(中文)").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=r, column=2, value="数量").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=r, column=3, value="所属Sheet").font = Font(bold=True, color="FFFFFF")
    for c in (1, 2, 3):
        ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="2F5597")
    order = [
        "missing_not_handed", "fresh_no_pickup", "late_handover", "carrier_slow", "fedex_slow",
        "stuck", "cancelled", "not_found", "reused_no_label", "in_transit", "delivered_ok",
    ]
    seen = set()
    for k in order:
        if k in seen:
            continue
        cnt = int(sum(df["_key"] == k))
        if cnt == 0 and k == "carrier_slow" and int(sum(df["_key"] == "fedex_slow")):
            continue
        if cnt == 0 and k == "fedex_slow" and int(sum(df["_key"] == "carrier_slow")):
            continue
        seen.add(k)
        r += 1
        zh, _en, col = CLASS[k]
        ws.cell(row=r, column=1, value=zh)
        ws.cell(row=r, column=2, value=cnt).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=3, value="承运异常" if k in SLOW_KEYS or k == "stuck" else "")
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=COLOR[col])
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=COLOR[col])
    cols = [c for c in DETAIL_COLS if c in df.columns]
    widths = [16] * len(cols)
    detail_plan = [
        ("异常处理", pick(*ANOMALY_CATS), True, "异常子集，含建议动作"),
        ("漏发未交接", pick("missing_not_handed", "fresh_no_pickup"), True, "通知仓库/货代核查"),
        ("迟发", pick("late_handover"), True, "含营业日延迟与 Amazon是否判迟"),
        ("承运异常", pick("fedex_slow", "carrier_slow", "stuck"), True, f"{slow_label}/卡件"),
        ("取消·其他", pick("cancelled", "not_found"), True, "已取消 / 查无"),
        ("全部明细", df, False, "全量审计"),
    ]
    for name, data, colorize, note in detail_plan:
        ws2 = wb.create_sheet(name)
        ws2.cell(row=1, column=1, value=note).font = Font(italic=True, color="808080")
        ws2.freeze_panes = "A3"
        start = 3
        for c, h in enumerate(cols, 1):
            cell = ws2.cell(row=start, column=c, value=h)
            cell.fill = PatternFill("solid", fgColor="2F5597")
            cell.font = Font(color="FFFFFF", bold=True)
        for i, (_, rd) in enumerate(data.iterrows()):
            rr = start + 1 + i
            for c, h in enumerate(cols, 1):
                cell = ws2.cell(row=rr, column=c, value=rd.get(h, ""))
                cell.border = thin
            if colorize:
                col = LEVEL_COLOR.get(rd.get("等级", ""), "green")
                ws2.cell(row=rr, column=1).fill = PatternFill("solid", fgColor=COLOR[col])
        for c, w in enumerate(widths, 1):
            ws2.column_dimensions[get_column_letter(c)].width = min(w, 24)
    if len(parked):
        ws_p = wb.create_sheet("未支持停放")
        pcols = [c for c in ("跟踪号", "跳过原因", "邮寄方式", "渠道", "订单号", "包裹号") if c in parked.columns or True]
        pcols = ["跟踪号", "跳过原因", "邮寄方式", "渠道", "订单号", "包裹号"]
        for c, h in enumerate(pcols, 1):
            cell = ws_p.cell(row=1, column=c, value=h)
            cell.fill = PatternFill("solid", fgColor="2F5597")
            cell.font = Font(color="FFFFFF", bold=True)
        for i, (_, rd) in enumerate(parked.iterrows()):
            for c, h in enumerate(pcols, 1):
                ws_p.cell(row=2 + i, column=c, value=rd.get(h, ""))
    ws3 = wb.create_sheet("口径说明")
    notes = [
        (notes_title, ""),
        ("1. 起点与确认", "起点=建标时间；确认发货=站点收件/首次取件扫描（UPS 用实际发货时间）。"),
        ("2. 迟发", "营业日延迟=建标→收件营业日−处理时间。周末与美国联邦假日不计。"),
        ("3. 处理时间", f"默认 {HANDLING_DAYS} 个营业日（UPS/FedEx 同日历；UPS 在途阈值未单独校准）。"),
        ("4. 承运延误", f"收件→交付营业日 > {TRANSIT_SLOW_DAYS}。一单既迟发又延误时主分类为承运延误。"),
        ("5. 卡件", f"未交付且最近扫描超过 {STUCK_DAYS} 天。漏发/未交接：有发货日期超过 {MISSING_AFTER_DAYS} 天无收件。"),
        ("6. 停放", "GLS/GOFO/无法识别承运商的行不查询，计入停放，不丢弃。"),
    ]
    for i, (k, v) in enumerate(notes, 1):
        ws3.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws3.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True)
        ws3.column_dimensions["A"].width = 22
        ws3.column_dimensions["B"].width = 118
    wb.save(out_xlsx)
    return df
