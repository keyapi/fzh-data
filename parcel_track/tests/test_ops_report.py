from __future__ import annotations

import datetime as _dt
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from parcel_track.normalize import classified_row
from parcel_track.ops_excel import write_ops_workbook


def test_mixed_sheets(tmp_path: Path):
    now = pd.Timestamp(_dt.datetime(2026, 9, 20, 12, 0))
    late = classified_row(
        {
            "跟踪号": "1ZAAA",
            "建标时间": "2026-09-01 10:00:00",
            "站点收件时间": "2026-09-04 10:00:00",
            "交付时间": "2026-09-05 10:00:00",
            "当前状态": "Delivered",
        },
        {"渠道": "Amazon US", "订单号": "O1"},
        now,
        "ups",
    )
    slow = classified_row(
        {
            "跟踪号": "382954490594",
            "建标时间": "2026-09-01 10:00:00",
            "站点收件时间": "2026-09-01 16:00:00",
            "交付时间": "2026-09-11 16:00:00",
            "当前状态": "Delivered",
        },
        {"渠道": "Amazon US", "订单号": "O2"},
        now,
        "fedex",
    )
    df = pd.DataFrame([late, slow])
    parked = pd.DataFrame([{"跟踪号": "GFUS1", "跳过原因": "unsupported:gofo", "邮寄方式": "GOFO", "渠道": "", "订单号": "O3", "包裹号": ""}])
    out = tmp_path / "ops.xlsx"
    write_ops_workbook(df, str(out), title="t", slow_label="承运延误", notes_title="n", parked=parked)
    wb = load_workbook(out)
    assert late["_key"] == "late_handover"
    assert slow["_key"] == "carrier_slow"
    names = set(wb.sheetnames)
    assert "迟发" in names
    assert "承运异常" in names
    assert "未支持停放" in names
