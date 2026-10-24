# -*- coding: utf-8 -*-
"""U7：对账闸门 + 产出（JSON / Excel / 管理摘要）—— 纯本地，不联网。"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

import pytest

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from tongtool_order_cost import profitability_audit as P  # noqa: E402
from tongtool_order_cost import profitability_report as R  # noqa: E402


def _finding(cat, owner, severity="high", amount=None):
    return P.Finding(cat, "order", "k", severity, {}, amount, P.PROVEN, owner)


def _passing_report():
    gates = R.reconcile_gates(
        row_conservation={"input": 3, "matched": 3, "unmatched": 0, "failed": 0,
                          "excluded": 0, "closes": True},
        revenue_ordinary=1000.0, cost_ordinary=400.0, profit_ordinary=600.0,
        reported_profit_ordinary=600.0,
    )
    findings = [
        _finding("cost_missing_anywhere", P.OWNER_MASTER),
        _finding("profit_inversion", P.OWNER_FINANCE),
        _finding("freight_repeated_across_rows", P.OWNER_TECH, amount=50.0),
        _finding("cost_present", P.OWNER_TECH, severity="info"),   # 常态，不进决策
    ]
    metrics = P.metric_matrix(
        P.build_source_rows(
            ["订单号", "通途SKU", "包裹号", "发货数量", "1件产品成本", "售价*汇率", "订单总成本", "订单利润"],
            [["A-1", "S", "P1", "1", "10", "100", "40", "60"]],
            "wb",
        ),
        [P.MetricSpec("产品成本_行均值", "1件产品成本", P.DEN_ROWS)],
    )
    return R.build_report(
        run_id="202607", metrics=metrics,
        row_conservation={"conservation": {"input": 3, "matched": 3, "closes": True}},
        findings=findings, gates=gates,
        cohorts={"by_account": [{"key": "A", "revenue": 100.0, "profit": 60.0, "margin": 0.6}]},
        questions={"product_cost": "x", "freight": "y", "profit_inversion": "z"},
        now="2026-10-14T00:00:00",
    )


def test_passing_gates_allow_conclusions(tmp_path):
    report = _passing_report()
    assert report.conclusions_blocked is False
    out = R.new_run_dir(tmp_path, "analyze", now=datetime(2026, 10, 14, 1, 2, 3))
    written = R.write_all(report, out)
    assert Path(written["json"]).exists()
    assert Path(written["xlsx"]).exists()
    assert Path(written["memo"]).exists()
    memo = Path(written["memo"]).read_text(encoding="utf-8")
    assert "可给结论" in memo
    assert "三个问题" in memo


def test_all_required_sheets_present_and_counts_match_json(tmp_path):
    import openpyxl

    report = _passing_report()
    out = R.new_run_dir(tmp_path, "analyze", now=datetime(2026, 10, 14, 1, 2, 3))
    written = R.write_all(report, out)
    wb = openpyxl.load_workbook(written["xlsx"], read_only=True)
    assert wb.sheetnames == R.REQUIRED_SHEETS
    data = json.loads(Path(written["json"]).read_text(encoding="utf-8"))
    # 决策清单行数 = JSON 决策清单总条数 + 表头
    assert wb["16_决策清单"].max_row == (
        sum(len(v) for v in data["decision_lists"].values()) + 1
    )
    # 指标口径行数 = JSON 指标数 + 表头
    assert wb["01_指标口径"].max_row == len(data["metrics"]) + 1


def test_row_conservation_failure_blocks_conclusions(tmp_path):
    gates = R.reconcile_gates(
        row_conservation={"input": 5, "matched": 3, "unmatched": 0, "failed": 0,
                          "excluded": 0, "closes": False},
        revenue_ordinary=100.0, cost_ordinary=40.0, profit_ordinary=60.0,
        reported_profit_ordinary=60.0,
    )
    report = R.build_report(run_id="x", gates=gates,
                            questions={"product_cost": "p", "freight": "f", "profit_inversion": "i"})
    assert report.conclusions_blocked is True
    memo = R.render_memo(report)
    assert "已阻断" in memo


def test_profit_identity_mismatch_blocks_even_if_rows_pass():
    gates = R.reconcile_gates(
        row_conservation={"input": 1, "matched": 1, "closes": True},
        revenue_ordinary=1000.0, cost_ordinary=400.0, profit_ordinary=500.0,  # 应为 600
        reported_profit_ordinary=500.0,
    )
    byname = {g.name: g for g in gates}
    assert byname["row_conservation"].passed is True
    assert byname["profit_identity"].passed is False


def test_no_gates_means_blocked():
    report = R.build_report(run_id="x")   # 没给闸门
    assert report.conclusions_blocked is True


def test_new_run_dir_never_overwrites(tmp_path):
    now = datetime(2026, 10, 14, 1, 2, 3)
    R.new_run_dir(tmp_path, "analyze", now=now)
    with pytest.raises(FileExistsError):
        R.new_run_dir(tmp_path, "analyze", now=now)


def test_decision_lists_separated_and_exclude_info():
    report = _passing_report()
    assert len(report.decision_lists[P.OWNER_MASTER]) == 1
    assert len(report.decision_lists[P.OWNER_FINANCE]) == 1
    assert len(report.decision_lists[P.OWNER_TECH]) == 1          # info 项被排除
    assert all(d["category"] != "cost_present"
               for items in report.decision_lists.values() for d in items)
