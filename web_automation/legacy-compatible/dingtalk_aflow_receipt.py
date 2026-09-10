#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""钉钉 aflow「销售收款确认单」单据导出为 Excel + 离职发起人附件补齐（浏览器，持久化登录）。

用法:
  # 1) 导出单据表
  uv run python web_automation/legacy-compatible/dingtalk_aflow_receipt.py \
      --from 2026-07-04 --to 2026-09-09
  # 2) 按导出表补离职发起人的附件（API userNotExist 的那批）
  uv run python web_automation/legacy-compatible/dingtalk_aflow_receipt.py \
      --mode attachments

产出:
  excel        <out>/Amazon&新平台成本 {from}-{to} 销售收款确认单-{导出时间戳}.xlsx
  attachments  <out>/dingtalk_oa_approval_data/aflow_attachments/<数据id>/<原名>
               + aflow_attachments_manifest.jsonl（按 数据id 一行，断点续传）

流程（选择器与踩坑详见 web_automation/docs/reference/aflow-receipt-export.md）:
  excel:       登录 → 表单名称级联(已启用→销售收款确认单) → 发起时间 → 查询
               → 导出全部(异步任务) → 操作记录/导出记录 轮询 → 下载 → 落盘
  attachments: 读导出表 → 筛离职发起人 → 逐单开 plainapproval 详情页 → 点文件名下载

约定:
  - 默认 headed：cookie 失效时用户可在窗口里现场登录（脚本会自动尝试一键头像授权）。
  - 文件名自构造，不用 download.suggested_filename。
  - 只写 <out>（仓库外）与 gitignore 的 profile 目录，不写仓库内任何目录。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from playwright.sync_api import Page, TimeoutError as PWTimeout, sync_playwright

SCRIPT_DIR = Path(__file__).resolve().parent
WEB_ROOT = SCRIPT_DIR.parent
DEFAULT_PROFILE = WEB_ROOT / "dingtalk-profile"
DEFAULT_OUT = Path(r"D:\Work\王忠于\成本核算")

# 直接开 aflow 未登录会落到一个没有任何登录控件的 error.vm；
# 必须先走 oa.dingtalk.com 才能触发钉钉统一身份认证 + 选组织。
AFLOW_URL = (
    "https://aflow.dingtalk.com/dingtalk/web/query/dashboard"
    "?dinghash=aflowSetting#/aflowSetting/dataManage?tabKey=default"
)
RECORD_URL = (
    "https://aflow.dingtalk.com/dingtalk/web/query/dashboard"
    "?dinghash=aflowSetting#/aflowSetting/dataManage?tab=export&tabKey=record"
)
LOGIN_URL = "https://oa.dingtalk.com/index.htm"
DEFAULT_FORM = "销售收款确认单"
FILE_PREFIX = "Amazon&新平台成本"

LOGIN_HINT = (
    "[信息] 未检测到登录态。已在浏览器里打开钉钉登录页。\n"
    "       登录优先级：账号密码（.env 配了就优先）→ 一键头像 → 人工扫码。\n"
    "       若需要你介入：\n"
    "       1) 账号密码登录失败时会自动回退，可手动扫码\n"
    "       2) 钉钉对**陌生设备**可能再要一次短信验证码 —— 脚本会提示，请在弹出的\n"
    "          浏览器窗口里输入（同 profile 只需一次）\n"
    "       3) Chrome 左上角若弹「访问此设备上的其它应用和服务」→ 点【允许】\n"
    "          （原生气泡，脚本点不到；按 profile 只问一次）\n"
    f"       4) 登录态持久化到 {DEFAULT_PROFILE}，之后不再需要\n"
)


def emit_failure(code: str, detail: str = "") -> None:
    if detail:
        print(detail)
    print(f"FAILURE_CODE={code}")


def ymd(s: str) -> str:
    """2026-07-04 -> 20260704"""
    return datetime.strptime(s, "%Y-%m-%d").strftime("%Y%m%d")


def is_logged_in(page: Page) -> bool:
    """登录后才会出现的锚点：筛选表单里的「表单名称」级联。"""
    try:
        return page.locator("text=表单名称").count() > 0 and "error.vm" not in page.url
    except Exception:
        return False


def dismiss_modals(page: Page) -> None:
    """首屏两个弹窗会挡住点击。"""
    for sel in (
        "button.upgrade-guide-button:not(.button-primary)",  # 新增搜索功能 → 稍后升级
        "button.dtd-modal-close",                            # 智能OA试用已过期 → ×
    ):
        try:
            loc = page.locator(sel)
            if loc.count():
                loc.first.click(timeout=3000)
                page.wait_for_timeout(500)
        except Exception:
            pass


def try_avatar_login(page: Page) -> bool:
    """钉钉登录页的「一键头像授权」：勾选自动登录 + 点头像。

    依赖本机钉钉客户端在 127.0.0.1:8441-8443 可达（客户端版本不匹配时不可用，
    页面只剩二维码）→ 返回 False，交由用户扫码。
    另需用户在 Chrome 的「访问此设备上的其它应用和服务」气泡点允许（原生 UI，脚本够不到）。
    """
    try:
        if not page.locator(".app-qr-login-page").count():
            return False
    except Exception:
        return False
    try:
        box = page.locator(".app-qr-login-page .base-comp-check-box-rememberme-box").first
        if box.count() and "checkbox-done" not in (box.get_attribute("class") or ""):
            box.click(timeout=3000)
            page.wait_for_timeout(300)
    except Exception:
        pass
    try:
        av = page.locator(".app-qr-login-page .module-qrcode-user-avatar").first
        if av.count():
            av.click(timeout=5000)
            page.wait_for_timeout(1500)
            return True
    except Exception:
        pass
    return False


def pick_org(page: Page, org: str) -> bool:
    """登录后「选择你管理的组织」页：点目标组织。"""
    try:
        if not page.locator(".app-page-curr").count():
            return False
        target = page.locator(f'.app-page-curr :text-is("{org}")')
        if target.count():
            target.first.click(timeout=5000)
            page.wait_for_timeout(2000)
            print(f"[信息] 已选择组织：{org}")
            return True
    except Exception:
        pass
    return False


def load_env() -> dict:
    """读 web_automation/.env（已被 .gitignore 排除）。凭据只从这里来，不走命令行。"""
    vals: dict[str, str] = {}
    env_path = WEB_ROOT / ".env"
    if not env_path.is_file():
        return vals
    for raw in env_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        vals[k.strip()] = v.strip().strip('"').strip("'")
    return vals


def _click_in(page: Page, scope: str, text: str, attempt_each_ms: int = 1500, rounds: int = 10) -> bool:
    """在 scope 内点「文本相等且真的可点」的按钮。

    **必须限定 scope**：整页有 3 个文本为「登录」的按钮（另两个是扫码登录的
    `module-qrscan-login-btn`/`module-localscan-login-btn`），点到它们会触发
    钉钉客户端跳转并把页面搞崩（实测 page crashed）。
    `is_visible()` 也不够——被遮住的元素它照样返回 True，只有 Playwright 的
    click 可点击性检查（visible+stable+receives events）才靠谱，所以逐个试。
    """
    for _ in range(rounds):
        for el in page.locator(f"{scope} .base-comp-button").all():
            try:
                if (el.inner_text() or "").strip() != text:
                    continue
                el.click(timeout=attempt_each_ms)
                return True
            except Exception:
                continue
        page.wait_for_timeout(800)
    return False


def _fill_in(page: Page, scope: str, selector: str, value: str,
             attempt_each_ms: int = 1500, rounds: int = 10) -> bool:
    """在 scope 内填「真的能填」的输入框。"""
    for _ in range(rounds):
        for el in page.locator(f"{scope} {selector}").all():
            try:
                el.fill(value, timeout=attempt_each_ms)
                if el.input_value():
                    return True
            except Exception:
                continue
        page.wait_for_timeout(800)
    return False


def _check_remember_me(page: Page, scope: str = ".module-pass-login") -> bool:
    """勾上「自动登录」。

    **必须勾**：不勾的话钉钉只发短期会话，profile 里留不住登录态，
    下次运行又要重新登录（还会再触发一次短信验证码）。
    """
    try:
        for box in page.locator(f"{scope} .base-comp-check-box-rememberme-box").all():
            try:
                parent_txt = box.evaluate(
                    "el => (el.parentElement && el.parentElement.innerText) || ''"
                ) or ""
            except Exception:
                continue
            if "自动登录" not in parent_txt:
                continue
            if "checkbox-done" in (box.get_attribute("class") or ""):
                return True
            box.click(timeout=3000)
            page.wait_for_timeout(400)
            return True
    except Exception:
        pass
    return False


def login_with_password(page: Page, user: str, password: str) -> bool:
    """钉钉「账号登录」：手机号 → 下一步 → 密码 → 登录。

    实测（2026-09-10）这一步**没有验证码、没有短信**，可直接脚本化。
    所有控件都限定在 `.module-pass-login` 作用域内 —— 该容器恰好只含本流程的
    手机号/密码框与「下一步/登录」两个按钮，能避开整页几十个同名控件。
    """
    scope = ".module-pass-login"
    try:
        for el in page.locator("[role=tab]").all():
            try:
                if "账号登录" in (el.inner_text() or "") and el.is_visible():
                    el.click(timeout=3000)
                    page.wait_for_timeout(1500)
                    break
            except Exception:
                continue

        if not _fill_in(page, scope, 'input[type=tel][placeholder="请输入手机号码"]', user):
            print("[警告] 没填上手机号，回退人工登录")
            return False
        page.wait_for_timeout(800)

        if not _click_in(page, scope, "下一步"):
            print("[警告] 点不动「下一步」")
            return False
        page.wait_for_timeout(2500)

        if not _fill_in(page, scope, 'input[type=password][placeholder="请输入密码"]', password):
            print("[警告] 没填上密码（可能被要求验证码/短信验证）")
            return False
        page.wait_for_timeout(800)

        if _check_remember_me(page, scope):
            print("[信息] 已勾选「自动登录」（保证登录态能持久化到 profile）")
        else:
            print("[警告] 没能勾上「自动登录」——登录态可能不持久，下次还要重登")

        if not _click_in(page, scope, "登录"):
            print("[警告] 点不动「登录」")
            return False
        page.wait_for_timeout(4000)
        print("[信息] 已提交手机号+密码；钉钉可能再要求短信验证码（见下方提示）")
        return True
    except Exception as e:
        print(f"[警告] 账号密码登录异常（{type(e).__name__}）：{str(e)[:150]}")
        return False


def _sms_step_present(page: Page) -> bool:
    """是否停在「短信验证码」这一步。

    钉钉对陌生设备/风控会插一道短信验证码（`.module-verify-code-input`）。
    这是**预期内的**，不是脚本 bug；同 profile 登录成功一次后就不再需要。
    """
    try:
        for el in page.locator(".module-verify-code-input").all():
            if el.is_visible():
                return True
    except Exception:
        pass
    return False


def wait_for_login(page: Page, timeout_s: int, org: str, user: str = "", password: str = "") -> bool:
    """开 oa.dingtalk.com 触发统一身份认证。

    登录方式优先级：账号密码（.env 有就优先）→ 一键头像 → 等用户扫码。
    账号密码**只尝试一次**：错了就交给人工，避免连续失败触发风控/锁定。
    中途若出现短信验证码，会**明确提示**并等用户输入（这一步无法自动化）。
    """
    if is_logged_in(page):
        return True
    print(LOGIN_HINT)
    try:
        page.goto(LOGIN_URL, wait_until="domcontentloaded", timeout=60000)
    except Exception:
        pass
    page.wait_for_timeout(3000)
    waited = 0
    pwd_tried = False
    avatar_done = False
    sms_reported = False
    while waited < timeout_s:
        if page.url.startswith("https://oa.dingtalk.com/index.htm") and "welcome" in page.url:
            break
        if "login.dingtalk.com" in page.url:
            if user and password and not pwd_tried:
                pwd_tried = True
                login_with_password(page, user, password)
            elif not avatar_done and try_avatar_login(page):
                avatar_done = True
                print("[信息] 已勾选自动登录并点击头像，等待授权…")
            if not sms_reported and _sms_step_present(page):
                sms_reported = True
                print(
                    "\n" + "=" * 66 + "\n"
                    "[需要你介入] 钉钉要求**短信验证码**（陌生设备/风控触发，正常现象）。\n"
                    "             请在浏览器窗口里填入手机收到的验证码，输完脚本会自动继续。\n"
                    "[说明] 这一步无法自动化（短信只有你手机上有）。\n"
                    "       同一个 profile 登录成功一次后不再需要——登录态已持久化。\n"
                    + "=" * 66 + "\n"
                )
            pick_org(page, org)
        page.wait_for_timeout(2000)
        waited += 2
    try:
        page.goto(AFLOW_URL, wait_until="domcontentloaded", timeout=60000)
    except Exception:
        pass
    page.wait_for_timeout(4000)
    return is_logged_in(page)


def current_form(page: Page) -> str:
    """当前已选的表单名（aflow 按账号记住上次选择，可能是「已启用 / 销售收款确认单」）。"""
    try:
        el = page.locator(".ant-select-selection-item").first
        if not el.count():
            return ""
        return (el.get_attribute("title") or el.inner_text() or "").strip()
    except Exception:
        return ""


def select_form(page: Page, form: str) -> bool:
    """表单名称是二级级联：先状态，再表单。必须精确匹配（有多个近似名）。

    幂等：已经选好就直接跳过——否则显示已选值的 span 会挡住输入框。
    """
    want = f"已启用 / {form}"
    if current_form(page) == want:
        print(f"[信息] 表单已是「{want}」，跳过选择")
        return True
    try:
        # 点容器本身，不要点内层 search input（会被 selection-item 拦截）
        page.locator(".ant-select.ant-cascader").first.click(timeout=10000)
        page.wait_for_timeout(800)
        page.locator('.ant-cascader-menu-item:has-text("已启用")').first.click()
        page.wait_for_timeout(800)
        page.locator(f'.ant-cascader-menu :text-is("{form}")').first.click()
        page.wait_for_timeout(800)
        got = current_form(page)
        if got != want:
            print(f"[错误] 表单选择未生效: 期望「{want}」，实际「{got}」")
            return False
        return True
    except Exception as e:
        print(f"[错误] 选表单失败: {e}")
        return False


def set_date_range(page: Page, d_from: str, d_to: str) -> bool:
    """发起时间用的是 dtd RangePicker，input 是 readOnly，只能走日历面板。"""
    try:
        cur_from = page.locator('input[placeholder="开始日期"]').first.input_value()
        cur_to = page.locator('input[placeholder="结束日期"]').first.input_value()
        if cur_from == d_from and cur_to == d_to:
            print(f"[信息] 发起时间已是 {d_from}~{d_to}，跳过设置")
            return True
        page.locator('input[placeholder="开始日期"]').first.click()
        page.wait_for_timeout(1000)
        target = datetime.strptime(d_from, "%Y-%m-%d")
        # 面板一次显示两个月，翻到包含起始月为止
        for _ in range(60):
            headers = page.locator(".dtd-picker-header-view:visible").all_inner_texts()
            if not headers:
                break
            m = re.search(r"(\d{4})年(\d{1,2})月", headers[0])
            if m and (int(m.group(1)), int(m.group(2))) >= (target.year, target.month):
                break
            page.locator(".dtd-picker-header-next-btn:visible").first.click()
            page.wait_for_timeout(300)
        page.locator(f'[title="{d_from}"]').first.click()
        page.wait_for_timeout(800)
        page.locator(f'[title="{d_to}"]').first.click()
        page.wait_for_timeout(800)
        got_from = page.locator('input[placeholder="开始日期"]').first.input_value()
        got_to = page.locator('input[placeholder="结束日期"]').first.input_value()
        if got_from != d_from or got_to != d_to:
            print(f"[错误] 日期未生效: 期望 {d_from}~{d_to}，实际 {got_from}~{got_to}")
            return False
        return True
    except Exception as e:
        print(f"[错误] 设置发起时间失败: {e}")
        return False


def list_export_rows(page: Page) -> list[str]:
    """操作记录里每行的导出文件名（第 1 列）。"""
    out: list[str] = []
    for tr in page.locator("tr:visible").all():
        try:
            first = tr.locator("td").first
            if first.count():
                t = (first.inner_text() or "").strip()
                if t and t != "导出文件名称":
                    out.append(t)
        except Exception:
            continue
    return out


def goto_records(page: Page) -> None:
    # 注意：只改 hash 的 goto 在 SPA 里不是重载，必须 reload 才能拿到最新进度
    page.goto(RECORD_URL, wait_until="domcontentloaded", timeout=60000)
    page.wait_for_timeout(1500)
    page.reload(wait_until="domcontentloaded")
    page.wait_for_timeout(4000)


# ---------------------------------------------------------------------------
# 附件模式：按单据走 plainapproval 详情页取附件
#
# 背景：批量附件导出（「导出全部」hover 的"仅导出审批单附件"）产物进钉盘【云盘-团队文件】，
# aflow 侧取不回本地；而导出表里的 #/previewAttachments 深链**需要钉钉客户端**才能打开。
# 浏览器里能走通的是「查看」对应的详情页：
#   .../pchomepage.htm?from=oflow&op=true&corpid=<corp>#/plainapproval?procInstId=<数据id>
# 页面上附件卡片虽然带 `file-list disabled`（"预览"动作不可见），但点**文件名**仍会触发真实下载。
# 这条正是离职发起人（API `userNotExist`）唯一的取件路径。
# ---------------------------------------------------------------------------

CORP_DEFAULT = "dingb0f80e79aafefc8635c2f4657eb6378f"
DETAIL_URL = (
    "https://aflow.dingtalk.com/dingtalk/web/query/pchomepage.htm"
    "?from=oflow&op=true&corpid={corp}#/plainapproval?procInstId={pid}"
)


def safe_name(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "_", (name or "file").strip()) or "file"


def read_export_records(xlsx: Path) -> tuple[list[dict], str]:
    """从导出表读 (数据id, 发起人姓名, 审批编号)，并从超链接里取 corpid。

    导出表是 2 行表头、数据自第 3 行；同一单据会因明细表重复成多行 → 按 数据id 去重。
    """
    import openpyxl

    wb = openpyxl.load_workbook(xlsx)
    recs: list[dict] = []
    seen: set[str] = set()
    corp = ""
    for sn in wb.sheetnames:
        ws = wb[sn]
        h1 = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        ix = {h: i for i, h in enumerate(h1) if h}
        if "数据id" not in ix:
            continue
        c_id = ix["数据id"] + 1
        c_name = ix.get("发起人姓名", ix["数据id"]) + 1
        c_bid = ix.get("审批编号", ix["数据id"]) + 1
        for r in range(3, ws.max_row + 1):
            pid = ws.cell(row=r, column=c_id).value
            if not pid:
                continue
            pid = str(pid)
            if pid in seen:
                continue
            seen.add(pid)
            recs.append({
                "pid": pid,
                "name": str(ws.cell(row=r, column=c_name).value or ""),
                "bid": ws.cell(row=r, column=c_bid).value,
            })
            if not corp:
                hl = ws.cell(row=r, column=102).hyperlink
                if hl and hl.target:
                    m = re.search(r"corpid=([A-Za-z0-9]+)", hl.target)
                    if m:
                        corp = m.group(1)
    return recs, (corp or CORP_DEFAULT)


def list_record_files(page: Page, corp: str, pid: str, fields: list[str]) -> list[dict]:
    """打开某单据详情页，列出要下载的附件（只取指定字段标签，天然跳过图片控件）。"""
    page.goto(DETAIL_URL.format(corp=corp, pid=pid), wait_until="domcontentloaded", timeout=60000)
    page.reload(wait_until="domcontentloaded")  # 只改 hash 的 goto 不重渲染
    page.wait_for_timeout(5000)
    try:
        return page.evaluate(
            """(fields) => {
                const out = [];
                document.querySelectorAll('.m-field-view').forEach(v => {
                    const lab = v.querySelector('.m-field-view-label');
                    const label = lab ? lab.innerText.trim() : '';
                    if (!fields.includes(label)) return;
                    v.querySelectorAll('.file-list-item').forEach((it, i) => {
                        const n = it.querySelector('.item-name');
                        const s = it.querySelector('.item-size');
                        out.push({
                            label,
                            idx: i,
                            name: n ? n.innerText.trim() : '',
                            size: s ? s.innerText.trim() : '',
                        });
                    });
                });
                return out;
            }""",
            fields,
        )
    except Exception as e:
        print(f"    [警告] 读取附件列表失败: {str(e)[:120]}")
        return []


def fetch_record(page: Page, corp: str, pid: str, fields: list[str], dest_dir: Path) -> dict:
    """取某单据的附件。已存在且非空的跳过（断点续传）。"""
    files = list_record_files(page, corp, pid, fields)
    saved: list[dict] = []
    errors: list[dict] = []
    for f in files:
        dest = dest_dir / safe_name(f["name"])
        try:
            if dest.exists() and dest.stat().st_size > 0:
                saved.append({**f, "path": dest.name, "skipped": True})
                continue
            scope = page.locator(f'.m-field-view:has(label:text-is("{f["label"]}"))').first
            name_el = scope.locator(".file-list-item").nth(f["idx"]).locator(".item-name")
            with page.expect_download(timeout=60000) as dl_info:
                name_el.click(force=True, timeout=15000)
            dl = dl_info.value
            dl.save_as(str(dest))
            saved.append({**f, "path": dest.name, "bytes": dest.stat().st_size})
            page.wait_for_timeout(400)
        except Exception as e:
            errors.append({**f, "reason": str(e)[:200]})
    return {"files": saved, "errors": errors}


def run_attachments(page: Page, args, base_out: Path, att_out: Path) -> int:
    xlsx = Path(args.from_xlsx) if args.from_xlsx else _newest_export(base_out, args.form)
    if not xlsx or not xlsx.is_file():
        emit_failure("INVALID_ARGUMENT",
                     f"[错误] 找不到导出表，请用 --from-xlsx 指定（默认在 {base_out} 找最新的）")
        return 1
    recs, corp = read_export_records(xlsx)
    if args.only_departed:
        targets = [r for r in recs if "离职" in r["name"]]
    else:
        targets = recs
    fields = [s.strip() for s in args.fields.split(",") if s.strip()]
    print(f"[信息] 导出表 {xlsx.name}：{len(recs)} 单 → 目标 {len(targets)} 单"
          f"（{'仅离职发起人' if args.only_departed else '全部'}），字段={fields}，corpId={corp}")
    if not targets:
        print("[信息] 没有符合条件的单据")
        return 0

    manifest = att_out / "aflow_attachments_manifest.jsonl"
    n_ok = n_err = n_skip = n_files = 0
    with manifest.open("a", encoding="utf-8") as fh:
        for i, rec in enumerate(targets, 1):
            dest_dir = att_out / rec["pid"]
            try:
                res = fetch_record(page, corp, rec["pid"], fields, dest_dir)
            except Exception as e:
                res = {"files": [], "errors": [{"reason": f"{type(e).__name__}: {str(e)[:200]}"}]}
            got = [f for f in res["files"] if not f.get("skipped")]
            skipped = [f for f in res["files"] if f.get("skipped")]
            n_files += len(res["files"])
            n_skip += len(skipped)
            if res["errors"]:
                n_err += 1
            else:
                n_ok += 1
            rec_out = {**rec, "corp": corp, **res, "ok": not res["errors"]}
            fh.write(json.dumps(rec_out, ensure_ascii=False) + "\n")
            fh.flush()
            print(f"[{i}/{len(targets)}] {rec['name']} {rec['bid']} "
                  f"取到={len(got)} 已存在={len(skipped)} 失败={len(res['errors'])}")
            if res["errors"]:
                for e in res["errors"]:
                    print(f"    [错误] {e.get('name')}: {e.get('reason')}")
    print(f"[汇总] 目标单据 {len(targets)}：成功 {n_ok} / 有失败 {n_err}；"
          f"文件 {n_files}（新下载 {n_files - n_skip}，已存在跳过 {n_skip}）")
    print(f"[汇总] 落盘 {att_out}；manifest {manifest.name}")
    if n_err:
        emit_failure("PARTIAL_FAILURE", f"[警告] {n_err} 个单据有附件未取到，详见 manifest")
        return 1
    return 0


def _newest_export(out_dir: Path, form: str) -> Path | None:
    cands = sorted(
        out_dir.glob(f"{FILE_PREFIX} *{form}-*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    return cands[0] if cands else None


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="钉钉 aflow 销售收款确认单导出")
    ap.add_argument("--mode", default="excel", choices=["excel", "attachments"])
    ap.add_argument("--from", dest="d_from", default="2026-07-04")
    ap.add_argument("--to", dest="d_to", default="2026-09-09")
    ap.add_argument("--form", default=DEFAULT_FORM)
    ap.add_argument("--org", default="方州汇国际", help="登录后要选择的组织名")
    ap.add_argument("--out", default=str(DEFAULT_OUT))
    ap.add_argument("--profile", default=str(DEFAULT_PROFILE))
    ap.add_argument("--headless", action="store_true")
    ap.add_argument("--channel", default="chrome",
                    help="浏览器通道；默认本机 Chrome（钉钉登录页在自带 Chromium 上会崩）。传空串则强制自带 Chromium")
    ap.add_argument("--login-timeout", type=int, default=300)
    ap.add_argument("--export-timeout", type=int, default=600)
    # 附件模式（--mode attachments）
    ap.add_argument("--from-xlsx", default="",
                    help="附件模式：从哪张导出表读 数据id/发起人；默认取 --out 下最新的")
    ap.add_argument("--fields", default="账期明细",
                    help="附件模式：要取哪些字段标签（逗号分隔）；默认只取账期明细，图片控件不取")
    ap.add_argument("--only-departed", dest="only_departed", action="store_true", default=True,
                    help="附件模式：只取离职发起人的单据（默认开）")
    ap.add_argument("--all-originators", dest="only_departed", action="store_false",
                    help="附件模式：不限离职，取导出表里全部单据")
    ap.add_argument("--auto-login", action="store_true",
                    help="仅为 dispatcher/CLI 对齐保留；钉钉网页是扫码/SSO，非图形验证码，本参数是 no-op")
    args = ap.parse_args()

    if args.mode not in ("excel", "attachments"):
        emit_failure("INVALID_ARGUMENT", f"[错误] 未知 mode: {args.mode}")
        return 1
    if args.auto_login:
        print("[信息] --auto-login 对钉钉无意义（非验证码登录），已忽略")

    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    profile_dir = Path(args.profile)
    profile_dir.mkdir(parents=True, exist_ok=True)

    # 凭据只从 web_automation/.env 读（gitignore），不进命令行、不进日志
    env = load_env()
    dt_user = env.get("DINGTALK_USER", "").strip()
    dt_pwd = env.get("DINGTALK_PASSWORD", "").strip()
    if dt_user and dt_pwd:
        print(f"[信息] 已从 .env 读到钉钉账号（{dt_user[:3]}****{dt_user[-4:]}），登录时优先用账号密码")
    else:
        print("[信息] .env 未配置 DINGTALK_USER/DINGTALK_PASSWORD，将用手动登录（一键头像/扫码）")

    with sync_playwright() as p:
        ctx = None
        # 优先用本机安装的 Chrome：钉钉登录页在 Playwright 自带 Chromium 上会崩（实测 page crashed），
        # 而 MCP 用的就是本机 Chrome，登录页在那边正常。
        launch_kw = dict(
            user_data_dir=str(profile_dir),
            headless=args.headless,
            accept_downloads=True,
            viewport={"width": 1500, "height": 950},
            args=["--disable-blink-features=AutomationControlled"],
        )
        if args.channel:
            try:
                ctx = p.chromium.launch_persistent_context(channel=args.channel, **launch_kw)
                print(f"[信息] 浏览器通道：{args.channel}")
            except Exception as e:
                print(f"[警告] 打不开 {args.channel} 通道（{str(e)[:80]}），回退自带 Chromium")
        if ctx is None:
            ctx = p.chromium.launch_persistent_context(**launch_kw)
            print("[信息] 浏览器通道：自带 Chromium")
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        try:
            page.goto(AFLOW_URL, wait_until="domcontentloaded", timeout=60000)
            page.wait_for_timeout(3000)

            if not is_logged_in(page) and not wait_for_login(
                page, args.login_timeout, args.org, dt_user, dt_pwd
            ):
                emit_failure("LOGIN_TIMEOUT", "[错误] 登录未完成，请手动登录后重试")
                return 1

            if args.mode == "attachments":
                att_out = out_dir / "dingtalk_oa_approval_data" / "aflow_attachments"
                att_out.mkdir(parents=True, exist_ok=True)
                return run_attachments(page, args, out_dir, att_out)

            dismiss_modals(page)

            if not select_form(page, args.form):
                emit_failure("FORM_NOT_FOUND", f"[错误] 未找到表单：{args.form}")
                return 1
            if not set_date_range(page, args.d_from, args.d_to):
                emit_failure("INVALID_ARGUMENT", "[错误] 发起时间未设置成功，中止以免导出默认区间")
                return 1

            page.locator('button.dtd-btn-primary:has-text("查询")').first.click()
            page.wait_for_timeout(4000)
            rows_now = list_export_rows(page)
            if page.locator("text=暂无数据").count():
                print("[信息] 查询结果为空，仍继续尝试导出（导出全部按当前筛选）")

            # 先记下已有导出记录，用于锚定本次新任务
            goto_records(page)
            before = set(list_export_rows(page))
            page.goto(AFLOW_URL, wait_until="domcontentloaded", timeout=60000)
            page.wait_for_timeout(4000)
            dismiss_modals(page)
            select_form(page, args.form)
            set_date_range(page, args.d_from, args.d_to)
            page.locator('button.dtd-btn-primary:has-text("查询")').first.click()
            page.wait_for_timeout(3000)

            # 导出全部：点按钮本体直接触发（没有下拉菜单）
            page.locator('button:has-text("导出全部")').first.click()
            page.wait_for_timeout(2500)
            # 弹窗「数据正在导出…」
            if page.locator('button:has-text("查看进度")').count():
                page.locator('button:has-text("查看进度")').first.click()
                page.wait_for_timeout(3000)

            # 轮询：本次新出现的「销售收款确认单-*.xlsx」且 已完成
            bench_ts = datetime.now().strftime("%Y%m%d%H%M%S")
            print(f"[信息] 已提交导出，本次基准时间 {bench_ts}，开始轮询…")
            target_name = ""
            row_prefix = f"{args.form}-"
            waited = 0
            while waited < args.export_timeout:
                goto_records(page)
                names = list_export_rows(page)
                new = [n for n in names if n not in before and n.startswith(row_prefix)]
                if new:
                    target_name = new[0]
                    tr = page.locator("tr:visible").filter(has_text=target_name).first
                    txt = (tr.inner_text() or "").replace("\n", " ")
                    if "已完成" in txt:
                        print(f"[信息] 任务完成: {target_name}")
                        break
                    print(f"[信息] {waited}s 进度: {txt.strip()[:80]}")
                waited += 10
                page.wait_for_timeout(10000)
            else:
                emit_failure("EXPORT_TIMEOUT", f"[错误] 等 {args.export_timeout}s 仍未完成")
                return 1

            m = re.search(r"-(\d{14})\.xlsx$", target_name)
            stamp = m.group(1) if m else bench_ts
            final = out_dir / f"{FILE_PREFIX} {ymd(args.d_from)}-{ymd(args.d_to)} {args.form}-{stamp}.xlsx"

            tr = page.locator("tr:visible").filter(has_text=target_name).first
            with page.expect_download(timeout=120000) as dl_info:
                tr.locator("a.export-file-download").first.click()
            download = dl_info.value
            download.save_as(str(final))

            size_kb = final.stat().st_size / 1024
            print(f"[完成] {final}")
            print(f"[汇总] 单据导出 1 个文件，{size_kb:.0f} KB，来源任务 {target_name}")
            return 0
        except PWTimeout as e:
            emit_failure("EXPORT_TIMEOUT", f"[错误] 超时: {e}")
            return 1
        except Exception as e:
            emit_failure("UNCLASSIFIED_FAILURE", f"[错误] {type(e).__name__}: {e}")
            return 1
        finally:
            try:
                ctx.close()
            except Exception:
                pass


if __name__ == "__main__":
    raise SystemExit(main())
