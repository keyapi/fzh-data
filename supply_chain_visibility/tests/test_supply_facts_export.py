"""Conservative merge tests for the supply facts batch exporter."""
import importlib.util
import json
import hashlib
import tempfile
from pathlib import Path
from unittest import mock
import unittest


def _load_exporter():
    path = Path(__file__).resolve().parents[2] / "EN_API" / "supply_facts_export.py"
    spec = importlib.util.spec_from_file_location("supply_facts_export_test", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _empty_target():
    return {"lines": [], "stage_balances": [], "milestones": [],
            "availability_schedule": [], "links": [], "production_evidence": [],
            "shipment_progress": [], "receipt_reconciliation": [],
            "inventory_bridge": [], "partial": False,
            "counts": {}, "coverage": {}, "sources": {}, "source_timestamps": {},
            "restricted": {}, "external_sources": {},
            "notes": [], "schema_version": "", "rule_version": "",
            "implementation_version": "", "as_of": ""}


def _coverage(**overrides):
    value = {
        "query_range_complete": True,
        "requested_windows": [{"start": "2026-10-01T00:00:00+08:00",
                               "end": "2026-10-31T00:00:00+08:00"}],
        "truncated": False,
        "not_collected": [],
        "source_as_of": "2026-10-31T00:00:00+08:00",
        "covered_po_codes": ["PO-1"],
        "coverage_mode": "time_range",
    }
    value.update(overrides)
    return value


def _state_payload():
    return {
        "schema_version": "supply-facts/receipt-state/v1",
        "coverage": _coverage(),
        "links": [{"shipment": "SHIP-1", "dn": "DN-1",
                   "po_codes": ["PO-1"],
                   "signed_at": "2026-10-20T10:00:00+08:00",
                   "qty": 10, "uom": "个", "sku": "SKU-1",
                   "warehouse": "WH-1", "identity_confidence": "proven",
                   "status_semantics": "unproven"}],
        "receipts": [{"receipt_id": "R1", "line_id": "L1",
                      "po_code": "PO-1", "sku": "SKU-1",
                      "warehouse": "WH-1", "uom": "个", "qty": 10,
                      "effective_at": "2026-10-25T10:00:00+08:00",
                      "quantity_semantics": "increment",
                      "uom_proven": True, "identity_proven": True}],
    }


def _bridge_payload():
    return {
        "schema_version": "supply-facts/inventory-bridge/v1",
        "coverage": _coverage(),
        "snapshots": [{"inventory_snapshot_id": "SNAP-1", "sku": "SKU-1",
                       "warehouse": "WH-1", "uom": "个", "available": 100,
                       "as_of": "2026-10-05T00:00:00+00:00",
                       "window_start": "2026-10-01T00:00:00+00:00",
                       "window_end": "2026-10-05T23:59:59+00:00"}],
        "receipts": [{"receipt_id": "R1", "line_id": "L1",
                      "po_code": "PO-1", "sku": "SKU-1",
                      "warehouse": "WH-1", "uom": "个", "qty": 10,
                      "receipt_effective_at": "2026-10-10T10:00:00+00:00",
                      "uom_proven": True,
                      "quantity_semantics_proven": True}],
    }


def _batch(**coverage):
    return {"lines": [], "stage_balances": [], "milestones": [],
            "availability_schedule": [], "links": [], "production_evidence": [],
            "shipment_progress": [], "receipt_reconciliation": [],
            "inventory_bridge": [],
            "manifest": {"schema_version": "supply-facts/v1",
                         "implementation_version": "impl-v1",
                         "partial": True, "coverage": coverage,
                         "notes": ["note"], "counts": {"input": 1},
                         "rule_version": "v1", "as_of": "2026-10-27"}}


class MergeSemanticsTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.exporter = _load_exporter()

    def test_restricted_is_or_and_sources_is_and(self):
        target = _empty_target()
        self.exporter._merge(target, _batch(
            sources={"delivery_note": True, "work_order": True},
            restricted={"delivery_note": True, "work_order": False}))
        self.exporter._merge(target, _batch(
            sources={"delivery_note": False, "work_order": True},
            restricted={"delivery_note": False, "work_order": False}))
        self.assertTrue(target["restricted"]["delivery_note"])
        self.assertFalse(target["restricted"]["work_order"])
        self.assertFalse(target["sources"]["delivery_note"])
        self.assertTrue(target["sources"]["work_order"])
        self.assertTrue(target["partial"])

    def test_external_sources_and_timestamps_are_conservative(self):
        target = _empty_target()
        self.exporter._merge(target, _batch(
            external_sources={"receipt": "not_integrated"},
            source_timestamps={"delivery_note": "2026-10-27T10:00:00"}))
        self.exporter._merge(target, _batch(
            external_sources={"receipt": "integrated"},
            source_timestamps={"delivery_note": "2026-10-27T09:00:00"}))
        self.assertEqual(target["external_sources"]["receipt"], "not_integrated")
        self.assertEqual(target["source_timestamps"]["delivery_note"],
                         "2026-10-27T09:00:00")

    def test_empty_reason_is_preserved(self):
        target = _empty_target()
        self.exporter._merge(target, _batch(empty_reason="no_matching_lines"))
        self.assertEqual(target["coverage"]["empty_reason"], "no_matching_lines")

    def test_source_status_is_conservative_across_pages(self):
        target = _empty_target()
        self.exporter._merge(target, _batch(
            source_status={"stock_entry": {"available": False,
                                           "reason": "permission"}}))
        self.exporter._merge(target, _batch(
            source_status={"stock_entry": {"available": True}}))
        status = target["coverage"]["source_status"]["stock_entry"]
        self.assertFalse(status["available"])
        self.assertEqual(status["reason"], "permission")

    def test_implementation_version_is_carried(self):
        target = _empty_target()
        batch = _batch()
        batch["manifest"]["implementation_version"] = "key_test.supply_facts@2026-10-27.3"
        self.exporter._merge(target, batch)
        self.assertEqual(target["implementation_version"],
                         "key_test.supply_facts@2026-10-27.3")

    def test_mixed_implementation_version_is_rejected(self):
        target = _empty_target()
        self.exporter._merge(target, _batch())
        batch = _batch()
        batch["manifest"]["implementation_version"] = "impl-v2"
        with self.assertRaises(self.exporter.BatchVersionError):
            self.exporter._merge(target, batch)
        self.assertEqual(target["implementation_version"], "impl-v1")
        self.assertEqual(target["lines"], [])

    def test_old_and_new_implementation_versions_are_rejected(self):
        target = _empty_target()
        target["implementation_version"] = "key_test.supply_facts@2026-10-27.7"
        batch = _batch()
        batch["manifest"]["implementation_version"] = "key_test.supply_facts@2026-10-29.1"
        with self.assertRaises(self.exporter.BatchVersionError):
            self.exporter._merge(target, batch)

    def test_previous_and_current_implementation_versions_are_rejected(self):
        target = _empty_target()
        target["implementation_version"] = "key_test.supply_facts@2026-10-29.3"
        batch = _batch()
        batch["manifest"]["implementation_version"] = \
            "key_test.supply_facts@2026-10-29.4"
        with self.assertRaises(self.exporter.BatchVersionError):
            self.exporter._merge(target, batch)

    def test_mixed_rule_schema_or_as_of_is_rejected(self):
        target = _empty_target()
        self.exporter._merge(target, _batch())
        for key, value in (("rule_version", "v2"), ("schema_version", "supply-facts/v2"),
                           ("as_of", "2026-10-28T00:00:00")):
            bad = _batch()
            bad["manifest"][key] = value
            with self.assertRaises(self.exporter.BatchVersionError):
                self.exporter._merge(target, bad)

    def test_batch_missing_version_is_rejected(self):
        target = _empty_target()
        for key in ("schema_version", "rule_version", "implementation_version", "as_of"):
            bad = _batch()
            del bad["manifest"][key]
            with self.assertRaises(self.exporter.BatchVersionError):
                self.exporter._merge(target, bad)


class ExportPaginationTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.exporter = _load_exporter()

    def _batch_page(self, total, offset, limit):
        end = min(offset + limit, total)
        lines = [{"so": "SO-1", "so_line": "SOI-%05d" % i} for i in range(offset, end)]
        return {
            "lines": lines, "stage_balances": [], "milestones": [],
            "availability_schedule": [], "links": [],
            "manifest": {
                "schema_version": "supply-facts/v1",
                "rule_version": "v1",
                "implementation_version": "impl-v1",
                "as_of": "2026-10-29",
                "partial": True,
                "counts": {"input": len(lines)},
                "coverage": {},
                "notes": [],
            },
            "pagination": {
                "limit": limit, "offset": offset, "returned": len(lines), "total": total,
                "next_offset": None if end >= total else end,
                "next_cursor": lines[-1]["so_line"] if lines else None,
                "exhausted": end >= total,
                "partial": end < total,
                "termination_reason": "exhausted" if end >= total else "page_limit",
            },
        }

    def _run_export(self, argv, total=501):
        calls = []

        def fake_call(base, key, secret, filters):
            calls.append(dict(filters))
            return self._batch_page(total, int(filters.get("offset") or 0),
                                    int(filters["limit"]))

        with mock.patch.object(self.exporter.iss, "load_credentials",
                               return_value=("k", "s")), \
             mock.patch.object(self.exporter.iss, "resolve_env_files",
                               return_value=None), \
             mock.patch.object(self.exporter.iss, "ENV_URLS",
                               {"test": "http://test"}), \
             mock.patch.object(self.exporter, "_call", side_effect=fake_call), \
             mock.patch.object(self.exporter, "_resolve_sos",
                               return_value=["SO-1"]):
            with tempfile.TemporaryDirectory() as td:
                rc = self.exporter.main(argv + ["--out", td])
                self.assertEqual(rc, 0)
                runs = list((Path(td) / "runs").glob("run-*"))
                self.assertEqual(len(runs), 1)
                lines = json.loads((runs[0] / "lines.json").read_text(encoding="utf-8"))
                manifest = json.loads((runs[0] / "manifest.json").read_text(encoding="utf-8"))
        return calls, lines, manifest

    def test_single_so_501_rows_paginates_to_exhaustion(self):
        calls, lines, manifest = self._run_export(
            ["--so", "SO-1", "--test", "--page-size", "500", "--no-excel"])
        self.assertEqual(len(lines), 501)
        self.assertEqual([c["offset"] for c in calls], [0, 500])
        self.assertFalse(manifest["truncated"])
        self.assertEqual(manifest["termination_reason"], "exhausted")
        self.assertEqual(manifest["coverage"]["sales_orders"], 1)
        self.assertEqual(manifest["coverage"]["so_lines"], 501)
        self.assertEqual(manifest["coverage"]["so_lines_total"], 501)

    def test_customer_code_501_rows_paginates_to_exhaustion(self):
        calls, lines, manifest = self._run_export(
            ["--codes", "TT0031038K0062927", "--test", "--page-size", "500", "--no-excel"])
        self.assertEqual(len(lines), 501)
        self.assertEqual([c["offset"] for c in calls], [0, 500])
        self.assertFalse(manifest["truncated"])
        self.assertEqual(manifest["termination_reason"], "exhausted")
        self.assertEqual(manifest["coverage"]["sales_orders"], 1)
        self.assertEqual(manifest["coverage"]["so_lines"], 501)
        self.assertEqual(manifest["coverage"]["so_lines_total"], 501)

    def test_paginate_rejects_stalled_offset(self):
        def fake_call(base, key, secret, filters):
            batch = self._batch_page(10, 0, 500)
            batch["pagination"].update({"exhausted": False, "next_offset": 0})
            return batch

        with mock.patch.object(self.exporter, "_call", side_effect=fake_call):
            with self.assertRaises(self.exporter.iss.QueryError):
                self.exporter._paginate("base", "k", "s", {}, 500, max_pages=2)

    def test_paginate_rejects_duplicate_unique_key(self):
        first = self._batch_page(501, 0, 500)
        second = self._batch_page(501, 500, 500)
        second["lines"][0]["so_line"] = first["lines"][0]["so_line"]

        with mock.patch.object(self.exporter, "_call", side_effect=[first, second]):
            with self.assertRaises(self.exporter.iss.QueryError):
                self.exporter._paginate("base", "k", "s", {}, 500, max_pages=2)

    def test_paginate_reports_page_limit(self):
        def fake_call(base, key, secret, filters):
            offset = int(filters.get("offset") or 0)
            batch = self._batch_page(10**9, offset, 500)
            batch["pagination"].update({"exhausted": False, "next_offset": offset + 500})
            return batch

        with mock.patch.object(self.exporter, "_call", side_effect=fake_call):
            batches, term = self.exporter._paginate("base", "k", "s", {}, 500, max_pages=2)
        self.assertEqual(len(batches), 2)
        self.assertEqual(term, "page_limit")


    def test_paginate_rejects_gap_in_next_offset(self):
        first = self._batch_page(501, 0, 500)
        first["pagination"]["next_offset"] = 501  # skips line #500
        with mock.patch.object(self.exporter, "_call", side_effect=[first]):
            with self.assertRaises(self.exporter.iss.QueryError):
                self.exporter._paginate("base", "k", "s", {}, 500)

    def test_paginate_rejects_exhausted_with_missing_rows(self):
        first = self._batch_page(501, 0, 500)
        final = self._batch_page(501, 500, 500)
        final["lines"] = []
        final["pagination"].update({"returned": 0, "next_offset": None, "exhausted": True})
        with mock.patch.object(self.exporter, "_call", side_effect=[first, final]):
            with self.assertRaises(self.exporter.iss.QueryError):
                self.exporter._paginate("base", "k", "s", {}, 500)

    def test_paginate_rejects_total_change(self):
        first = self._batch_page(501, 0, 500)
        second = self._batch_page(501, 500, 500)
        second["pagination"]["total"] = 502
        with mock.patch.object(self.exporter, "_call", side_effect=[first, second]):
            with self.assertRaises(self.exporter.iss.QueryError):
                self.exporter._paginate("base", "k", "s", {}, 500)

    def test_paginate_rejects_returned_mismatch(self):
        first = self._batch_page(501, 0, 500)
        first["pagination"]["returned"] = 499
        with mock.patch.object(self.exporter, "_call", side_effect=[first]):
            with self.assertRaises(self.exporter.iss.QueryError):
                self.exporter._paginate("base", "k", "s", {}, 500)

    def test_merge_coverage_dedupes_entities_by_identity(self):
        target = _empty_target()

        def batch(line_key):
            item = _batch()
            item["lines"] = [{"so": "SO-1", "so_line": line_key}]
            item["links"] = [{
                "source_doctype": "Delivery Note", "source_name": "DN-1",
                "target_doctype": "Shipment", "target_name": "SHIP-1",
                "target_is_record": True,
            }]
            item["pagination"] = {"total": 2}
            item["manifest"]["coverage"] = {
                "sales_orders": 1, "so_lines": 1, "so_lines_total": 2,
                "delivery_notes": 1, "shipments": 1,
            }
            return item

        self.exporter._merge(target, batch("L1"), scope_key="SO-1")
        self.exporter._merge(target, batch("L2"), scope_key="SO-1")
        self.exporter._finalize_coverage(target)
        self.assertEqual(target["coverage"]["sales_orders"], 1)
        self.assertEqual(target["coverage"]["so_lines"], 2)
        self.assertEqual(target["coverage"]["so_lines_total"], 2)
        self.assertEqual(target["coverage"]["delivery_notes"], 1)
        self.assertEqual(target["coverage"]["shipments"], 1)
        self.assertNotIn("_coverage_state", target)


    def test_failed_pagination_preserves_latest_pointer(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            out.mkdir(parents=True, exist_ok=True)
            old_pointer = {"run_id": "previous-success"}
            (out / "latest.json").write_text(json.dumps(old_pointer), encoding="utf-8")
            bad = self._batch_page(501, 0, 500)
            bad["pagination"]["next_offset"] = 501
            with mock.patch.object(self.exporter.iss, "load_credentials",
                                   return_value=("k", "s")), \
                 mock.patch.object(self.exporter.iss, "resolve_env_files",
                                   return_value=None), \
                 mock.patch.object(self.exporter.iss, "ENV_URLS",
                                   {"test": "http://test"}), \
                 mock.patch.object(self.exporter, "_call", return_value=bad):
                rc = self.exporter.main(
                    ["--so", "SO-1", "--test", "--page-size", "500",
                     "--no-excel", "--out", td])
            self.assertEqual(rc, 1)
            self.assertEqual(json.loads((out / "latest.json").read_text(encoding="utf-8")),
                             old_pointer)
            self.assertFalse((out / "runs").exists())


    def test_merge_coverage_waybill_only_is_not_a_shipment(self):
        target = _empty_target()
        batch = _batch()
        batch["lines"] = [{"so": "SO-1", "so_line": "L1"}]
        batch["links"] = [{
            "source_doctype": "Delivery Note", "source_name": "DN-1",
            "target_doctype": "Shipment", "target_name": "WAYBILL-1",
            "target_is_record": False,
        }]
        batch["pagination"] = {"total": 1}
        batch["manifest"]["coverage"] = {
            "sales_orders": 1, "so_lines": 1, "so_lines_total": 1,
            "delivery_notes": 1, "shipments": 0,
        }
        self.exporter._merge(target, batch, scope_key="SO-1")
        self.exporter._finalize_coverage(target)
        self.assertEqual(target["coverage"]["shipments"], 0)
        self.assertEqual(target["coverage"]["delivery_notes"], 1)

    def test_merge_coverage_matches_only_real_shipments_and_dedupes(self):
        target = _empty_target()
        batch = _batch()
        batch["lines"] = [{"so": "SO-1", "so_line": "L1"}]

        def link(dn, ship, is_record):
            return {
                "source_doctype": "Delivery Note", "source_name": dn,
                "target_doctype": "Shipment", "target_name": ship,
                "target_is_record": is_record,
            }

        batch["links"] = [
            link("DN-1", "SHIP-1", False),  # waybill placeholder, not a record
            link("DN-1", "SHIP-1", True),   # real record
            link("DN-2", "SHIP-1", True),   # duplicate real record
            link("DN-2", "SHIP-2", True),
        ]
        batch["pagination"] = {"total": 1}
        batch["manifest"]["coverage"] = {
            "sales_orders": 1, "so_lines": 1, "so_lines_total": 1,
            "delivery_notes": 2, "shipments": 0,
        }
        self.exporter._merge(target, batch, scope_key="SO-1")
        self.exporter._finalize_coverage(target)
        self.assertEqual(target["coverage"]["shipments"], 2)
        self.assertEqual(target["coverage"]["delivery_notes"], 2)
        self.assertEqual(target["coverage"]["sales_orders"], 1)
        self.assertEqual(target["coverage"]["so_lines"], 1)
        self.assertEqual(target["coverage"]["so_lines_total"], 1)


class ProductionEvidenceExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.exporter = _load_exporter()

    def _evidence(self, qty=10):
        return {"qty_proven": qty, "qty_estimated": None, "confidence": "proven",
                "started_at": "2026-10-20 08:00:00",
                "completed_at": "2026-10-20 17:00:00",
                "time_status": "complete", "reasons": [], "conflicts": [],
                "components": {}, "source_records": []}

    def _pe_line(self, so_line):
        stock = self._evidence()
        stock["components"] = {"gross_qty": 10.0, "return_qty": 0.0,
                               "net_qty": 10.0, "unallocated_return_qty": 0.0}
        return {
            "so": "SO-1", "so_line": so_line, "line_item_code": "KS0001",
            "work_orders": [{
                "wo": "WO-1", "production_item": "KS0001", "qty": 10.0,
                "uom": "个", "docstatus": 1, "wo_cancelled": False,
                "first_process_done": self._evidence(),
                "all_process_done": self._evidence(),
                "qc_performed": self._evidence(),
                "qc_passed": {"class": "qc_passed", "status": "Accepted",
                              "confidence": "unknown", "qty_proven": None,
                              "qty_estimated": None, "started_at": None,
                              "completed_at": None, "time_status": "not_applicable",
                              "reasons": [], "conflicts": [],
                              "components": {}, "source_records": []},
                "stock_in_posted": stock,
                "conflicts": [], "reasons": [],
            }],
        }

    def _page(self, total, offset, limit):
        end = min(offset + limit, total)
        lines = [{"so": "SO-1", "so_line": "SOI-%05d" % i} for i in range(offset, end)]
        return {
            "lines": lines, "stage_balances": [], "milestones": [],
            "availability_schedule": [], "links": [],
            "production_evidence": [self._pe_line("SOI-%05d" % i)
                                    for i in range(offset, end)],
            "shipment_progress": [{"so": "SO-1",
                                   "so_line": "SOI-%05d" % i,
                                   "shipment": "SHIP-1",
                                   "stage": "shipped_in_transit"}
                                  for i in range(offset, end)],
            "receipt_reconciliation": [],
            "inventory_bridge": [],
            "manifest": {
                "schema_version": "supply-facts/v1", "rule_version": "v1",
                "implementation_version": "impl-v1", "as_of": "2026-10-29",
                "partial": True, "counts": {"input": len(lines)},
                "coverage": {}, "notes": [],
            },
            "pagination": {
                "limit": limit, "offset": offset, "returned": len(lines),
                "total": total,
                "next_offset": None if end >= total else end,
                "next_cursor": lines[-1]["so_line"] if lines else None,
                "exhausted": end >= total, "partial": end < total,
                "termination_reason": "exhausted" if end >= total else "page_limit",
            },
        }

    def _run(self, total=2, excel=False, extra_argv=()):
        def fake_call(base, key, secret, filters):
            return self._page(total, int(filters.get("offset") or 0),
                              int(filters["limit"]))

        argv = ["--so", "SO-1", "--test", "--page-size", "1"]
        if not excel:
            argv.append("--no-excel")
        argv.extend(extra_argv)
        with mock.patch.object(self.exporter.iss, "load_credentials",
                               return_value=("k", "s")), \
             mock.patch.object(self.exporter.iss, "resolve_env_files",
                               return_value=None), \
             mock.patch.object(self.exporter.iss, "ENV_URLS",
                               {"test": "http://test"}), \
             mock.patch.object(self.exporter, "_call", side_effect=fake_call):
            with tempfile.TemporaryDirectory() as td:
                rc = self.exporter.main(argv + ["--out", td])
                self.assertEqual(rc, 0)
                run = list((Path(td) / "runs").glob("run-*"))[0]
                payload = json.loads((run / "production_evidence.json")
                                     .read_text(encoding="utf-8"))
                manifest = json.loads((run / "manifest.json")
                                      .read_text(encoding="utf-8"))
                digest = hashlib.sha256(
                    (run / "production_evidence.json").read_bytes()).hexdigest()
                datasets = {name: json.loads((run / (name + ".json"))
                                               .read_text(encoding="utf-8"))
                            for name in ("shipment_progress", "receipt_reconciliation",
                                         "inventory_bridge")}
                xlsx_rows = None
                if excel:
                    from openpyxl import load_workbook
                    wb = load_workbook(run / "supply_facts.xlsx", read_only=True)
                    xlsx_rows = list(wb["production_evidence"].iter_rows(
                        values_only=True))
                    wb.close()
                return payload, manifest, digest, xlsx_rows, datasets

    def test_production_evidence_is_exported_with_checksum_and_rows(self):
        payload, manifest, digest, _, _ = self._run(total=2)
        self.assertEqual(len(payload), 2)
        entry = manifest["files"]["production_evidence.json"]
        self.assertEqual(entry["rows"], 2)
        self.assertEqual(entry["sha256"], digest)
        flat = self.exporter._flatten_production_evidence(payload)
        self.assertEqual(manifest["datasets"]["production_evidence"]["rows"],
                         len(flat))
        self.assertEqual(manifest["datasets"]["production_evidence"]["lines"], 2)

    def test_flatten_is_one_row_per_wo_and_evidence_class(self):
        flat = self.exporter._flatten_production_evidence([self._pe_line("SOI-1")])
        self.assertEqual(len(flat), 5)
        classes = {row["evidence_class"] for row in flat}
        self.assertEqual(classes, {"first_process_done", "all_process_done",
                                   "qc_performed", "qc_passed", "stock_in_posted"})
        stock = next(row for row in flat if row["evidence_class"] == "stock_in_posted")
        self.assertEqual(stock["gross_qty"], 10.0)
        self.assertEqual(stock["net_qty"], 10.0)
        self.assertEqual(stock["line_role"], "primary")

    def test_excel_sheet_matches_json_row_count(self):
        payload, manifest, _, xlsx_rows, _ = self._run(total=2, excel=True)
        self.assertEqual(len(xlsx_rows) - 1,
                         manifest["datasets"]["production_evidence"]["rows"])
        self.assertEqual(len(xlsx_rows) - 1,
                         len(self.exporter._flatten_production_evidence(payload)))

    def test_shipment_progress_and_unintegrated_sources_are_exported(self):
        _, manifest, _, _, datasets = self._run(total=2)
        self.assertEqual(manifest["datasets"]["shipment_progress"]["rows"], 2)
        self.assertEqual(len(datasets["shipment_progress"]), 2)
        self.assertIn("shipment_progress.json", manifest["files"])
        self.assertEqual(
            manifest["datasets"]["receipt_reconciliation"]["source"],
            "not_integrated")
        self.assertEqual(manifest["datasets"]["inventory_bridge"]["source"],
                         "not_integrated")
        self.assertFalse(manifest["datasets"]["inventory_bridge"]["eligible"])

    def test_local_receipt_and_inventory_inputs_are_exported(self):
        state = {
            "schema_version": self.exporter.RECEIPT_STATE_SCHEMA,
            "coverage": _coverage(),
            "links": [{"shipment": "SHIP-1", "dn": "DN-1",
                       "po_codes": ["PO-1"],
                       "signed_at": "2026-10-20T10:00:00+08:00",
                       "qty": 10, "uom": "个", "sku": "SKU-1",
                       "warehouse": "WH-1", "identity_confidence": "proven",
                       "status_semantics": "unproven"}],
            "receipts": [{"receipt_id": "R1", "line_id": "L1",
                          "po_code": "PO-1", "sku": "SKU-1",
                          "warehouse": "WH-1", "uom": "个", "qty": 10,
                          "effective_at": "2026-10-25T10:00:00+08:00",
                          "quantity_semantics": "increment",
                          "uom_proven": True, "identity_proven": True}],
        }
        bridge = {
            "schema_version": self.exporter.INVENTORY_BRIDGE_SCHEMA,
            "coverage": _coverage(),
            "snapshots": [{"inventory_snapshot_id": "SNAP-1", "sku": "SKU-1",
                           "warehouse": "WH-1", "uom": "个", "available": 100,
                           "as_of": "2026-10-05T00:00:00+00:00",
                           "window_start": "2026-10-01T00:00:00+00:00",
                           "window_end": "2026-10-05T23:59:59+00:00"}],
            "receipts": [{"receipt_id": "R1", "line_id": "L1",
                          "po_code": "PO-1", "sku": "SKU-1",
                          "warehouse": "WH-1", "uom": "个", "qty": 10,
                          "receipt_effective_at": "2026-10-10T10:00:00+00:00",
                          "uom_proven": True,
                          "quantity_semantics_proven": True}],
        }
        with tempfile.TemporaryDirectory() as inputs:
            state_path = Path(inputs) / "state.json"
            bridge_path = Path(inputs) / "bridge.json"
            state_path.write_text(json.dumps(state), encoding="utf-8")
            bridge_path.write_text(json.dumps(bridge), encoding="utf-8")
            _, manifest, _, _, datasets = self._run(
                total=1, extra_argv=["--receipt-state-json", str(state_path),
                                     "--inventory-bridge-json", str(bridge_path)])
        self.assertEqual(datasets["receipt_reconciliation"][0]["status"],
                         "tongtool_received")
        self.assertEqual(datasets["inventory_bridge"][0]["receipt_inclusion_status"],
                         "after_snapshot")
        self.assertEqual(manifest["datasets"]["receipt_reconciliation"]["source"],
                         "local_receipt_state")
        self.assertEqual(manifest["datasets"]["inventory_bridge"]["source"],
                         "local_inventory_bridge")

    def test_write_failure_does_not_switch_latest_pointer(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            old = {"run_id": "previous-success"}
            (out / "latest.json").write_text(json.dumps(old), encoding="utf-8")

            def fake_call(base, key, secret, filters):
                return self._page(2, int(filters.get("offset") or 0),
                                  int(filters["limit"]))

            with mock.patch.object(self.exporter.iss, "load_credentials",
                                   return_value=("k", "s")), \
                 mock.patch.object(self.exporter.iss, "resolve_env_files",
                                   return_value=None), \
                 mock.patch.object(self.exporter.iss, "ENV_URLS",
                                   {"test": "http://test"}), \
                 mock.patch.object(self.exporter, "_call", side_effect=fake_call), \
                 mock.patch.object(self.exporter, "_write_xlsx",
                                   side_effect=RuntimeError("boom")):
                rc = self.exporter.main(
                    ["--so", "SO-1", "--test", "--page-size", "1", "--out", td])
            self.assertEqual(rc, 1)
            self.assertEqual(json.loads((out / "latest.json").read_text(encoding="utf-8")),
                             old)
            self.assertEqual(list((out / "runs").glob("run-*")), [])


class LocalInputValidationTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.exporter = _load_exporter()

    def _load(self, payload, kind="receipt_state", name="state.json"):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / name
            path.write_text(json.dumps(payload), encoding="utf-8")
            return self.exporter._load_local_input(str(path), kind)

    def test_valid_payload_passes(self):
        self.assertEqual(self._load(_state_payload())["schema_version"],
                         self.exporter.RECEIPT_STATE_SCHEMA)
        self.assertEqual(self._load(_bridge_payload(), "inventory_bridge"),
                         _bridge_payload())

    def test_string_boolean_is_rejected(self):
        payload = _state_payload()
        payload["coverage"]["query_range_complete"] = "true"
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_unknown_field_is_rejected(self):
        payload = _state_payload()
        payload["links"][0]["extra"] = 1
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_unknown_top_level_field_is_rejected(self):
        payload = _state_payload()
        payload["collection_windows"] = []
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_missing_coverage_field_is_rejected(self):
        payload = _state_payload()
        del payload["coverage"]["not_collected"]
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_complete_with_empty_windows_is_rejected(self):
        payload = _state_payload()
        payload["coverage"]["requested_windows"] = []
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_reversed_window_is_rejected(self):
        payload = _state_payload()
        payload["coverage"]["requested_windows"] = [{
            "start": "2026-10-31T00:00:00+08:00",
            "end": "2026-10-01T00:00:00+08:00"}]
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_source_as_of_before_window_end_is_rejected(self):
        payload = _state_payload()
        payload["coverage"]["source_as_of"] = "2026-10-15T00:00:00+08:00"
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_link_po_outside_coverage_scope_is_rejected(self):
        payload = _state_payload()
        payload["links"][0]["po_codes"] = ["PO-B"]
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_po_full_history_mode_accepts_empty_windows(self):
        payload = _state_payload()
        payload["coverage"].update({
            "requested_windows": [],
            "coverage_mode": "po_full_history",
            "po_full_history_proven": True,
            "po_full_history_source": "tongtool_receipt_history",
            "po_full_history_as_of": "2026-10-31T00:00:00+08:00",
        })
        self.assertEqual(self._load(payload)["coverage"]["coverage_mode"],
                         "po_full_history")

    def test_po_full_history_as_of_must_match_source_as_of(self):
        payload = _state_payload()
        payload["coverage"].update({
            "requested_windows": [],
            "coverage_mode": "po_full_history",
            "po_full_history_proven": True,
            "po_full_history_source": "tongtool_receipt_history",
            "po_full_history_as_of": "2026-10-30T00:00:00+08:00",
        })
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_time_mode_rejects_full_history_proof_fields(self):
        payload = _state_payload()
        payload["coverage"]["po_full_history_proven"] = True
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_signed_at_may_be_null_when_status_is_proven(self):
        payload = _state_payload()
        payload["links"][0]["signed_at"] = None
        payload["links"][0]["signed_status_proven"] = True
        self.assertEqual(self._load(payload)["links"][0]["signed_status_proven"], True)

    def test_signed_status_proven_must_be_boolean(self):
        payload = _state_payload()
        payload["links"][0]["signed_status_proven"] = "true"
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_receipt_effective_at_may_be_null_for_unverified_rows(self):
        payload = _state_payload()
        payload["receipts"][0]["effective_at"] = None
        self.assertIsNone(self._load(payload)["receipts"][0]["effective_at"])

    def test_schema_mismatch_is_rejected(self):
        payload = _state_payload()
        payload["schema_version"] = "supply-facts/receipt-state/v2"
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_naive_time_is_rejected(self):
        payload = _state_payload()
        payload["links"][0]["signed_at"] = "2026-10-20"
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_duplicate_identity_with_different_payload_is_rejected(self):
        payload = _state_payload()
        second = dict(payload["receipts"][0], qty=6)
        payload["receipts"].append(second)
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload)

    def test_bridge_inclusion_proof_must_be_structured(self):
        payload = _bridge_payload()
        payload["receipts"][0]["inclusion_proof"] = True
        with self.assertRaises(self.exporter.LocalInputError):
            self._load(payload, "inventory_bridge")

    def test_invalid_input_does_not_create_run_or_switch_latest(self):
        bad = _state_payload()
        bad["coverage"]["truncated"] = "false"
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "out"
            out.mkdir()
            old = {"run_id": "previous-success"}
            (out / "latest.json").write_text(json.dumps(old), encoding="utf-8")
            bad_path = Path(td) / "bad.json"
            bad_path.write_text(json.dumps(bad), encoding="utf-8")
            with mock.patch.object(self.exporter.iss, "load_credentials",
                                   return_value=("k", "s")), \
                 mock.patch.object(self.exporter.iss, "resolve_env_files",
                                   return_value=None), \
                 mock.patch.object(self.exporter.iss, "ENV_URLS",
                                   {"test": "http://test"}), \
                 mock.patch.object(self.exporter, "_call",
                                   side_effect=AssertionError("must not query")):
                rc = self.exporter.main(
                    ["--so", "SO-1", "--test", "--no-excel", "--out", str(out),
                     "--receipt-state-json", str(bad_path)])
            self.assertEqual(rc, 2)
            self.assertEqual(json.loads((out / "latest.json").read_text(
                encoding="utf-8")), old)
            self.assertFalse((out / "runs").exists())


class AutoLinkTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.exporter = _load_exporter()

    def test_auto_links_use_server_shipment_progress(self):
        result = {
            "stage_balances": [{
                "stage": "signed_pending_tongtool_receipt", "batch_key": "DN-1",
                "item_code": "KS0001", "uom": "个", "warehouse": "WH-1",
                "qty": 10.0}],
            "shipment_progress": [{
                "stage": "signed_pending_tongtool_receipt",
                "shipment": "SHIP-1", "dn": "DN-1", "tongtool_po": "PO-1",
                "item_code": "KS0001", "uom": "个", "warehouse": "WH-1",
                "shipment_signed_at": "2026-10-20"}],
        }
        links = self.exporter._auto_shipment_links(result)
        self.assertEqual(len(links), 1)
        self.assertEqual(links[0].shipment, "SHIP-1")
        self.assertEqual(links[0].po_codes, ("PO-1",))
        self.assertEqual(links[0].qty, 10.0)
        self.assertEqual(links[0].identity_confidence, "proven")

    def test_auto_links_drive_export_without_hand_written_links(self):
        page = {
            "lines": [{"so": "SO-1", "so_line": "SOI-1"}],
            "stage_balances": [{
                "stage": "signed_pending_tongtool_receipt", "batch_key": "DN-1",
                "item_code": "SKU-1", "uom": "个", "warehouse": "WH-1",
                "qty": 10.0}],
            "shipment_progress": [{
                "stage": "signed_pending_tongtool_receipt",
                "shipment": "SHIP-1", "dn": "DN-1", "tongtool_po": "PO-1",
                "item_code": "SKU-1", "uom": "个", "warehouse": "WH-1",
                "is_final": True,
                "shipment_signed_at": "2026-10-20T10:00:00+08:00"}],
            "manifest": {
                "schema_version": "supply-facts/v1", "rule_version": "v1",
                "implementation_version": "impl-v1", "as_of": "2026-10-29",
                "partial": True, "counts": {"input": 1},
                "coverage": {}, "notes": [],
            },
            "pagination": {"limit": 500, "offset": 0, "returned": 1,
                           "total": 1, "next_offset": None, "exhausted": True,
                           "partial": False, "termination_reason": "exhausted"},
        }
        state = _state_payload()
        del state["links"]
        with tempfile.TemporaryDirectory() as td:
            state_path = Path(td) / "state.json"
            state_path.write_text(json.dumps(state), encoding="utf-8")
            with mock.patch.object(self.exporter.iss, "load_credentials",
                                   return_value=("k", "s")), \
                 mock.patch.object(self.exporter.iss, "resolve_env_files",
                                   return_value=None), \
                 mock.patch.object(self.exporter.iss, "ENV_URLS",
                                   {"test": "http://test"}), \
                 mock.patch.object(self.exporter, "_call", return_value=page):
                rc = self.exporter.main(
                    ["--so", "SO-1", "--test", "--no-excel", "--out", td,
                     "--receipt-state-json", str(state_path)])
            self.assertEqual(rc, 0)
            run = list((Path(td) / "runs").glob("run-*"))[0]
            rows = json.loads((run / "receipt_reconciliation.json").read_text(
                encoding="utf-8"))
            manifest = json.loads((run / "manifest.json").read_text(
                encoding="utf-8"))
        self.assertEqual(rows[0]["status"], "tongtool_received")
        self.assertEqual(
            manifest["datasets"]["receipt_reconciliation"]["links_source"],
            "server_auto")


class DeliveryExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.exporter = _load_exporter()

    def _delivery_page(self, total, offset, limit):
        end = min(offset + limit, total)
        lines = [{
            "delivery_note": "DN-%s" % i, "delivery_note_item": "DNI-%s" % i,
            "item_code": "SKU-1", "qty": 1, "uom": "个",
            "actual_qty": 1, "association_status": "unlinked",
            "relevance_status": "in_scope",
        } for i in range(offset, end)]
        return {
            "schema_version": "supply-facts/v3", "rule_version": "r3",
            "implementation_version": "impl-v3",
            "delivery_lines": lines, "association_edges": [],
            "coverage_exceptions": [], "so_fulfillment": [],
            "supply_ledger": [], "physical_chain": [],
            "scope": {"status": "pending", "source": None, "version": None,
                      "hash": None, "items": []},
            "snapshot_id": "snap-%s" % total,
            "coverage": {"delivery_lines": {"total": total},
                         "as_of": "2026-10-30",
                         "shipment_logistics_event": {"available": True},
                         "shipment_permission": True},
            "pagination": {
                "limit": limit, "offset": offset, "returned": len(lines),
                "total": total, "next_offset": None if end >= total else end,
                "next_cursor": lines[-1]["delivery_note_item"] if lines and end < total else None,
                "exhausted": end >= total, "partial": end < total,
                "termination_reason": "exhausted" if end >= total else "page_limit",
            },
        }

    def _run(self, td, total=3, page_size=2):
        calls = []

        def fake(base, key, secret, filters, scope=None):
            calls.append(dict(filters))
            return self._delivery_page(total, int(filters["offset"]),
                                       int(filters["limit"]))

        with mock.patch.object(self.exporter.iss, "load_credentials",
                               return_value=("k", "s")), \
             mock.patch.object(self.exporter.iss, "resolve_env_files",
                               return_value=None), \
             mock.patch.object(self.exporter.iss, "ENV_URLS",
                               {"test": "http://test"}), \
             mock.patch.object(self.exporter, "_call_delivery", side_effect=fake):
            rc = self.exporter.main(["--delivery-all", "--test", "--no-excel",
                                     "--page-size", str(page_size), "--out", td])
        return rc, calls

    def test_delivery_export_paginates_and_writes_manifest(self):
        with tempfile.TemporaryDirectory() as td:
            rc, calls = self._run(td)
            self.assertEqual(rc, 0)
            self.assertEqual([call["offset"] for call in calls], [0, 2])
            run = list((Path(td) / "runs").glob("run-*"))[0]
            lines = json.loads((run / "delivery_lines.json").read_text(encoding="utf-8"))
            manifest = json.loads((run / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(len(lines), 3)
            self.assertEqual(manifest["schema_version"], "supply-facts/v3")
            self.assertEqual(manifest["delivery_schema_version"], "supply-facts/v3")
            self.assertEqual(manifest["datasets"]["delivery_lines"]["rows"], 3)
            self.assertFalse(manifest["truncated"])
            for name in ("delivery_lines.json", "association_edges.json",
                         "coverage_exceptions.json", "so_fulfillment.json",
                         "supply_ledger.json", "physical_chain.json"):
                self.assertIn(name, manifest["files"])
                digest = hashlib.sha256((run / name).read_bytes()).hexdigest()
                self.assertEqual(manifest["files"][name]["sha256"], digest)

    def test_delivery_duplicate_identity_fails_before_writing(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            old = {"run_id": "previous-success"}
            (out / "latest.json").write_text(json.dumps(old), encoding="utf-8")

            def fake(base, key, secret, filters, scope=None):
                batch = self._delivery_page(3, int(filters["offset"]),
                                            int(filters["limit"]))
                if int(filters["offset"]) > 0 and batch["delivery_lines"]:
                    batch["delivery_lines"][0]["delivery_note_item"] = "DNI-0"
                return batch

            with mock.patch.object(self.exporter.iss, "load_credentials",
                                   return_value=("k", "s")), \
                 mock.patch.object(self.exporter.iss, "resolve_env_files",
                                   return_value=None), \
                 mock.patch.object(self.exporter.iss, "ENV_URLS",
                                   {"test": "http://test"}), \
                 mock.patch.object(self.exporter, "_call_delivery",
                                   side_effect=fake):
                rc = self.exporter.main(["--delivery-all", "--test", "--no-excel",
                                         "--page-size", "2", "--out", td])
            self.assertEqual(rc, 1)
            self.assertEqual(json.loads((out / "latest.json").read_text(
                encoding="utf-8")), old)
            self.assertFalse((out / "runs").exists())


    def test_delivery_mixed_version_across_pages_is_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            old = {"run_id": "previous-success"}
            (out / "latest.json").write_text(json.dumps(old), encoding="utf-8")

            def fake(base, key, secret, filters, scope=None):
                batch = self._delivery_page(3, int(filters["offset"]),
                                            int(filters["limit"]))
                if int(filters["offset"]) > 0:
                    batch["rule_version"] = "other-rule"
                return batch

            with mock.patch.object(self.exporter.iss, "load_credentials",
                                   return_value=("k", "s")), \
                 mock.patch.object(self.exporter.iss, "resolve_env_files",
                                   return_value=None), \
                 mock.patch.object(self.exporter.iss, "ENV_URLS",
                                   {"test": "http://test"}), \
                 mock.patch.object(self.exporter, "_call_delivery",
                                   side_effect=fake):
                rc = self.exporter.main(["--delivery-all", "--test", "--no-excel",
                                         "--page-size", "2", "--out", td])
            self.assertEqual(rc, 1)
            self.assertFalse((out / "runs").exists())
            self.assertEqual(json.loads((out / "latest.json").read_text(
                encoding="utf-8")), old)

    def test_delivery_snapshot_drift_across_pages_is_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            old = {"run_id": "previous-success"}
            (out / "latest.json").write_text(json.dumps(old), encoding="utf-8")

            def fake(base, key, secret, filters, scope=None):
                batch = self._delivery_page(3, int(filters["offset"]),
                                            int(filters["limit"]))
                if int(filters["offset"]) > 0:
                    batch["snapshot_id"] = "snap-drifted"
                return batch

            with mock.patch.object(self.exporter.iss, "load_credentials",
                                   return_value=("k", "s")), \
                 mock.patch.object(self.exporter.iss, "resolve_env_files",
                                   return_value=None), \
                 mock.patch.object(self.exporter.iss, "ENV_URLS",
                                   {"test": "http://test"}), \
                 mock.patch.object(self.exporter, "_call_delivery",
                                   side_effect=fake):
                rc = self.exporter.main(["--delivery-all", "--test", "--no-excel",
                                         "--page-size", "2", "--out", td])
            self.assertEqual(rc, 1)
            self.assertFalse((out / "runs").exists())


if __name__ == "__main__":
    unittest.main()
