#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 stock_daily_backup.py 的 JSONL 备份转成 g1b_coverage.py 可读的快照 JSON。

输入：D:/Work/库存备份/YYYY-MM-DD/warehouse=*.jsonl + manifest.json
输出：out/stock_snapshot_from_daily.json（聚合，无 PII）

严格口径：
  - 空/非法/负数/非整数库存量 → bad row，不转 0；
  - 同仓同 SKU 重复：同值计 duplicate，异值计 conflict 并整键丢弃，不 first-wins；
  - 缺文件/坏行/冲突 → 该仓 complete=false；
  - complete 只绑定实际选中并完整读取的仓。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from collections import defaultdict
from pathlib import Path

import openpyxl

BOM = Path(r"D:/Work/王忠于/成本核算/EN产品BOM成本列表20261016.xlsx")
FAMILIES = ("KS0001", "KS0248")
QTY_FIELDS = ("available", "intransit", "defects", "waiting")


def load_family_codes(path: Path = BOM) -> set[str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    i_item = hdr.index("产品编号")
    i_code = hdr.index("客户物料号")
    codes: set[str] = set()
    for r in rows:
        item = str(r[i_item] or "").strip()
        code = str(r[i_code] or "").strip()
        if item.startswith(FAMILIES) and code and "已删除" not in code:
            codes.add(code)
    wb.close()
    return codes


def parse_qty(value):
    if value is None or value == "":
        return None
    try:
        f = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(f) or f < 0 or f != int(f):
        return None
    return f


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--backup-dir", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--warehouses", default="", help="逗号分隔仓名；空=manifest 全部")
    ap.add_argument("--bom", default=str(BOM))
    ap.add_argument("--all-skus", action="store_true", help="不过滤 family 客户码")
    ap.add_argument("--allow-missing-hash", action="store_true",
                    help="兼容旧 manifest（无 file_sha256）；默认拒绝")
    args = ap.parse_args()

    family = None if args.all_skus else load_family_codes(Path(args.bom))
    root = Path(args.backup_dir)
    manifest_path = root / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    wanted = {w.strip() for w in args.warehouses.split(",") if w.strip()}
    all_entries = manifest.get("warehouses", [])
    by_name = {e.get("warehouseName"): e for e in all_entries}
    requested = sorted(wanted or set(by_name))
    entries = [by_name[n] for n in requested if n in by_name]
    missing_requested = [n for n in requested if n not in by_name]

    rows_by_key: dict[tuple[str, str], list[dict]] = defaultdict(list)
    row_meta: dict[tuple[str, str], dict] = {}
    bad_rows: dict[str, int] = defaultdict(int)
    bad_samples = []
    missing_files = []
    read_files = {}
    file_hashes = {}
    hash_missing = []
    hash_mismatches = {}

    for e in entries:
        name = e.get("warehouseName")
        path = Path(e.get("file") or "")
        if not path.is_file():
            missing_files.append(name)
            continue
        expected = e.get("file_sha256")
        actual = sha256(path)
        if expected:
            if actual != expected:
                hash_mismatches[name] = {"expected": expected, "actual": actual}
                continue
        elif not args.allow_missing_hash:
            hash_missing.append(name)
            continue
        read_files[name] = str(path)
        file_hashes[name] = actual
        with path.open(encoding="utf-8") as f:
            for line_no, line in enumerate(f, 1):
                line = line.strip()
                if not line:
                    continue
                try:
                    r = json.loads(line)
                except json.JSONDecodeError:
                    bad_rows[name] += 1
                    if len(bad_samples) < 20:
                        bad_samples.append({"warehouse": name, "file": str(path),
                                            "line": line_no, "reason": "bad_json"})
                    continue
                sku = str(r.get("goodsSku") or "").strip()
                if not sku:
                    bad_rows[name] += 1
                    if len(bad_samples) < 20:
                        bad_samples.append({"warehouse": name, "file": str(path),
                                            "line": line_no, "reason": "missing_sku"})
                    continue
                if family is not None and sku not in family:
                    continue
                vals = {
                    "available": parse_qty(r.get("availableStockQuantity")),
                    "intransit": parse_qty(r.get("intransitStockQuantity")),
                    "defects": parse_qty(r.get("defectsStockQuantity")),
                    "waiting": parse_qty(r.get("waitingShipmentStockQuantity")),
                }
                if any(v is None for v in vals.values()):
                    bad_rows[name] += 1
                    if len(bad_samples) < 20:
                        bad_samples.append({"warehouse": name, "file": str(path),
                                            "line": line_no, "reason": "bad_qty", "sku": sku})
                    continue
                key = (name, sku)
                rows_by_key[key].append(vals)
                row_meta.setdefault(key, {"file": str(path), "line": line_no})

    by_sku: dict[str, dict] = {}
    by_wh: dict[str, dict] = {}
    duplicate_same = 0
    conflicts = []
    conflict_warehouses: set[str] = set()
    for (wh_name, sku), rows in rows_by_key.items():
        if len(rows) > 1:
            if all(vals == rows[0] for vals in rows[1:]):
                duplicate_same += len(rows) - 1
            else:
                conflict_warehouses.add(wh_name)
                conflicts.append({"warehouse": wh_name, "sku": sku, "rows": len(rows),
                                  "source": row_meta.get((wh_name, sku))})
                continue
        vals = rows[0]
        entry = by_sku.setdefault(
            sku, {"sku": sku, "warehouses": {}, "available": 0.0, "intransit": 0.0,
                  "defects": 0.0, "waiting": 0.0}
        )
        entry["warehouses"][wh_name] = vals
        for k in QTY_FIELDS:
            entry[k] += vals[k]
        w = by_wh.setdefault(wh_name, {
            "available": 0.0, "intransit": 0.0, "defects": 0.0, "waiting": 0.0,
            "rows": 0, "complete": False,
        })
        for k in QTY_FIELDS:
            w[k] += vals[k]
        w["rows"] += 1

    for name in missing_requested:
        by_wh.setdefault(name, {
            "available": 0.0, "intransit": 0.0, "defects": 0.0, "waiting": 0.0,
            "rows": 0, "complete": False,
        })
        by_wh[name]["error"] = "requested_but_not_in_manifest"
    for e in entries:
        name = e.get("warehouseName")
        w = by_wh.setdefault(name, {
            "available": 0.0, "intransit": 0.0, "defects": 0.0, "waiting": 0.0,
            "rows": 0, "complete": False,
        })
        complete = (
            e.get("complete") is True
            and name not in missing_files
            and name not in conflict_warehouses
            and name not in hash_missing
            and name not in hash_mismatches
            and bad_rows.get(name, 0) == 0
        )
        w["complete"] = complete
        w["bad_rows"] = bad_rows.get(name, 0)
        w["conflict"] = name in conflict_warehouses
        w["file"] = read_files.get(name)

    complete_warehouses = sorted(
        name for name, w in by_wh.items() if w.get("complete") is True
    )
    presence: dict[str, set[str]] = defaultdict(set)
    for (wh_name, sku) in rows_by_key:
        presence[wh_name].add(sku)
    global_complete = bool(
        requested
        and not missing_requested
        and not missing_files
        and not hash_missing
        and not hash_mismatches
        and all(by_wh.get(name, {}).get("complete") is True for name in requested)
    )

    result = {
        "run_id": manifest.get("run_id"),
        "integrity_mode": ("legacy_unverified" if args.allow_missing_hash else "strict"),
        "integrity_exception": ("--allow-missing-hash" if args.allow_missing_hash else None),
        "fetched_at": manifest.get("snapshot_at"),
        "capture_start": min((e.get("capture_start") for e in entries if e.get("capture_start")),
                             default=None),
        "capture_end": max((e.get("capture_end") for e in entries if e.get("capture_end")),
                           default=None),
        "source": "stock_daily_backup.py (tongtool API)",
        "requested_warehouses": requested,
        "missing_requested": missing_requested,
        "read_warehouses": sorted(read_files),
        "missing_files": missing_files,
        "hash_missing": hash_missing,
        "hash_mismatches": hash_mismatches,
        "file_hashes": file_hashes,
        "warehouses": sorted(read_files),
        "complete": global_complete,
        "complete_warehouses": complete_warehouses,
        "incomplete_warehouses": sorted(set(requested) - set(complete_warehouses)),
        "sku_presence": {wh: sorted(skus) for wh, skus in presence.items()},
        "bad_rows_total": sum(bad_rows.values()),
        "bad_samples": bad_samples,
        "duplicate_same": duplicate_same,
        "duplicate_conflicts": len(conflicts),
        "conflict_samples": conflicts[:20],
        "by_warehouse": {
            k: {**v, **{f: round(v[f], 2) for f in QTY_FIELDS}}
            for k, v in by_wh.items()
        },
        "by_sku": sorted(by_sku.values(), key=lambda x: x["sku"]),
        "notes": [
            "由每日 API JSONL 聚合；不重复请求通途。",
            "空/非法/负数/非整数量→bad row，不转 0。",
            "同仓同 SKU 重复：同值计 duplicate；异值 conflict 并整键丢弃，不 first-wins。",
            "缺文件/坏行/冲突→该仓 complete=false；complete 只绑定实际选中并完整读取的仓。",
        ],
    }
    Path(args.out).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"OK: {args.out}")
    print(f"  requested={len(requested)} read={len(read_files)} complete={global_complete} "
          f"skus={len(by_sku)} bad={result['bad_rows_total']} conflicts={len(conflicts)}")


if __name__ == "__main__":
    main()
