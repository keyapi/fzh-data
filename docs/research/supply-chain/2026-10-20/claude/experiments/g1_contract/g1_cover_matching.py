#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""G1a corrected: match by EN BOM customer material code only.

Family: KS0001 / KS0248.
Codes: 客户物料号 in EN产品BOM成本列表20261016.xlsx.
Classes: -Cover (cover), -out, -Foam, finished (base).
No category-text inclusion; text is diagnostic only.
"""

from __future__ import annotations

import glob
import json
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
OUT = BASE / "out"
BOM = Path(r"D:/Work/王忠于/成本核算/EN产品BOM成本列表20261016.xlsx")
ORDER_GLOB = r"D:/Work/王忠于/成本核算/2026年{month}月订单_order_cost_*.xlsx"
MONTHS = range(4, 10)
FAMILIES = ("KS0001", "KS0248")


def cell(row, idx):
    return None if idx is None or idx >= len(row) else row[idx]


def norm_wh(w):
    u = (w or "").strip().upper()
    if "CENTRADE" in u:
        return "美东-CENTRADE"
    if "DANEEY" in u:
        if "退货" in u:
            return "DANEEY-退货仓"
        if "皮壳" in u or "半成品" in u:
            return "DANEEY-皮壳/半成品"
        return "美中-FZH-DANEEY"
    if "POLAND" in u or "波兰" in u:
        return "波兰"
    return "其他"


def classify(code: str) -> str:
    c = code.strip()
    if c.endswith("-Cover"):
        return "cover"
    if c.endswith("-out"):
        return "out"
    if c.endswith("-Foam"):
        return "foam"
    return "finished"


def load_bom():
    wb = openpyxl.load_workbook(BOM, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    idx = {n: (hdr.index(n) if n in hdr else None) for n in ("产品编号", "客户物料号")}
    codes = {}  # code -> (item, class)
    empty = deleted = family_rows = 0
    for r in rows:
        if not r:
            continue
        item = str(cell(r, idx["产品编号"]) or "").strip()
        if not item.startswith(FAMILIES):
            continue
        family_rows += 1
        code = str(cell(r, idx["客户物料号"]) or "").strip()
        if not code:
            empty += 1
            continue
        if code.startswith("已删除"):
            deleted += 1
            continue
        cls = classify(code)
        if code in codes and codes[code][0] != item:
            # 同一客户码映射多个 EN 物料：保留但标记
            codes[code] = (codes[code][0], "ambiguous")
        else:
            codes[code] = (item, cls)
    wb.close()
    by_class = Counter(v[1] for v in codes.values())
    return {
        "family_rows": family_rows,
        "active_codes": len(codes),
        "empty_codes": empty,
        "deleted_codes": deleted,
        "by_class": dict(by_class),
        "cover_codes": sum(1 for v in codes.values() if v[1] == "cover"),
        "finished_codes": sum(1 for v in codes.values() if v[1] == "finished"),
    }, codes


def latest(m):
    fs = sorted(glob.glob(ORDER_GLOB.format(month=m)))
    return fs[-1] if fs else None


def main():
    bom_summary, codes = load_bom()
    cover_codes = {k: v for k, v in codes.items() if v[1] == "cover"}
    finished_codes = {k: v for k, v in codes.items() if v[1] == "finished"}
    out_codes = {k: v for k, v in codes.items() if v[1] == "out"}
    foam_codes = {k: v for k, v in codes.items() if v[1] == "foam"}

    monthly = []
    cover_wh = Counter()
    finished_wh = Counter()
    out_wh = Counter()
    foam_wh = Counter()
    cover_qty = finished_qty = out_qty = foam_qty = 0.0
    cover_rows = finished_rows = out_rows = foam_rows = 0
    text_tri_unmatched_rows = 0
    text_tri_unmatched_qty = 0.0
    total_rows = 0

    for m in MONTHS:
        p = latest(m)
        if not p:
            monthly.append({"month": f"2026-{m:02d}", "file": None})
            continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        hdr = list(next(rows))
        ix = lambda n: hdr.index(n) if n in hdr else None
        cols = {k: ix(k) for k in ("通途SKU", "产品名称", "品类orig", "品类",
                                   "品类统一名", "发货数量", "发货仓库")}
        mc = mf = mo = mfm = 0
        qc = qf = qo = qfm = 0.0
        for r in rows:
            if not r:
                continue
            total_rows += 1
            sku = str(cell(r, cols["通途SKU"]) or "").strip()
            wh = norm_wh(cell(r, cols["发货仓库"]))
            try:
                q = float(cell(r, cols["发货数量"]) or 0)
            except (TypeError, ValueError):
                q = 0.0
            if sku in cover_codes:
                cover_rows += 1; cover_qty += q; mc += 1; qc += q; cover_wh[wh] += q
            elif sku in finished_codes:
                finished_rows += 1; finished_qty += q; mf += 1; qf += q; finished_wh[wh] += q
            elif sku in out_codes:
                out_rows += 1; out_qty += q; mo += 1; qo += q; out_wh[wh] += q
            elif sku in foam_codes:
                foam_rows += 1; foam_qty += q; mfm += 1; qfm += q; foam_wh[wh] += q
            else:
                texts = [str(cell(r, cols[k]) or "") for k in
                         ("产品名称", "品类orig", "品类", "品类统一名")]
                if any("三角" in t for t in texts):
                    text_tri_unmatched_rows += 1
                    text_tri_unmatched_qty += q
        wb.close()
        monthly.append({
            "month": f"2026-{m:02d}", "file": p.split("/")[-1],
            "cover_rows": mc, "cover_qty": qc,
            "finished_rows": mf, "finished_qty": qf,
            "out_rows": mo, "out_qty": qo,
            "foam_rows": mfm, "foam_qty": qfm,
        })

    out = {
        "bom": bom_summary,
        "orders": {
            "total_rows": total_rows,
            "cover_rows": cover_rows, "cover_qty": round(cover_qty, 2),
            "finished_rows": finished_rows, "finished_qty": round(finished_qty, 2),
            "out_rows": out_rows, "out_qty": round(out_qty, 2),
            "foam_rows": foam_rows, "foam_qty": round(foam_qty, 2),
            "cover_by_warehouse_qty": dict(cover_wh.most_common()),
            "finished_by_warehouse_qty": dict(finished_wh.most_common()),
            "out_by_warehouse_qty": dict(out_wh.most_common()),
            "foam_by_warehouse_qty": dict(foam_wh.most_common()),
            "monthly": monthly,
            "text_tri_unmatched_rows": text_tri_unmatched_rows,
            "text_tri_unmatched_qty": round(text_tri_unmatched_qty, 2),
            "note": "family 只由 BOM 客户物料号精确匹配决定；品类文本只用于诊断未匹配，不纳入 family。",
        },
    }
    (OUT / "cover_matching.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(out, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
