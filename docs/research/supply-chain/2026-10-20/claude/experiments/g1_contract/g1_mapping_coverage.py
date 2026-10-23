#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""G1a：三角靠枕 local mapping coverage.

Read-only local files:
  - EN BOM cost list 20261016 (产品编号 + 客户物料号)
  - monthly order_cost 2026-04..09 (通途SKU + 品类/仓库/数量)

Outputs aggregate only (no customer PII) to out/.
"""

from __future__ import annotations

import glob
import json
import os
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
OUT = BASE / "out"
OUT.mkdir(parents=True, exist_ok=True)
BOM = Path(r"D:/Work/王忠于/成本核算/EN产品BOM成本列表20261016.xlsx")
ORDER_GLOB = r"D:/Work/王忠于/成本核算/2026年{month}月订单_order_cost_*.xlsx"
MONTHS = list(range(4, 10))
FAMILIES = ("KS0001", "KS0248")


def cell(row, idx):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def load_bom():
    wb = openpyxl.load_workbook(BOM, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    idx = {name: (hdr.index(name) if name in hdr else None)
           for name in ("产品编号", "客户物料号", "产品名称")}
    reverse = defaultdict(set)
    item_codes = defaultdict(set)
    deleted = set()
    family_rows = 0
    all_rows = 0
    for r in rows:
        if not r:
            continue
        all_rows += 1
        item = str(cell(r, idx["产品编号"]) or "").strip()
        if not item.startswith(FAMILIES):
            continue
        family_rows += 1
        code = str(cell(r, idx["客户物料号"]) or "").strip()
        if not code:
            continue
        if code.startswith("已删除"):
            deleted.add(code)
            continue
        reverse[code].add(item)
        item_codes[item].add(code)
    wb.close()
    ambiguous = {k: sorted(v) for k, v in reverse.items() if len(v) > 1}
    return {
        "all_rows": all_rows,
        "family_rows": family_rows,
        "active_codes": len(reverse),
        "deleted_codes": len(deleted),
        "en_items": len(item_codes),
        "ambiguous_codes": len(ambiguous),
        "ambiguous_sample": dict(list(sorted(ambiguous.items()))[:20]),
        "reverse": reverse,
        "deleted": deleted,
    }


def latest_order_file(month):
    fs = sorted(glob.glob(ORDER_GLOB.format(month=month)))
    return fs[-1] if fs else None


def normalize_warehouse(wh: str) -> str:
    w = (wh or "").strip().upper()
    if "CENTRADE" in w:
        return "美东-CENTRADE"
    if "DANEEY" in w:
        if "退货" in w:
            return "DANEEY-退货仓"
        if "皮壳" in w or "半成品" in w:
            return "DANEEY-皮壳/半成品"
        return "美中-FZH-DANEEY"
    if "POLAND" in w or "波兰" in w:
        return "波兰"
    return "其他"


G1_WAREHOUSES = {"美东-CENTRADE", "美中-FZH-DANEEY"}


def load_orders(reverse, deleted):
    by_month = []
    by_warehouse = Counter()
    by_family = Counter()
    by_family_qty = Counter()
    by_channel = Counter()
    missing_prefix = Counter()
    total_rows = tri_rows = tri_qty = 0
    mapped_rows = mapped_qty = 0
    ambiguous_rows = missing_rows = deleted_only_rows = 0
    missing_qty = 0
    g1 = {"rows": 0, "qty": 0.0, "mapped_rows": 0, "mapped_qty": 0.0,
          "ambiguous_rows": 0, "missing_rows": 0, "missing_qty": 0.0,
          "deleted_only_rows": 0, "by_warehouse": Counter()}
    for m in MONTHS:
        p = latest_order_file(m)
        if not p:
            by_month.append({"month": f"2026-{m:02d}", "file": None})
            continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        hdr = list(next(rows))
        def ix(name):
            return hdr.index(name) if name in hdr else None
        cols = {k: ix(k) for k in
                ("发货日期", "通途SKU", "产品名称", "品类orig", "品类",
                 "品类统一名", "发货数量", "渠道", "发货仓库")}
        m_rows = m_tri = m_mapped = m_amb = m_miss = m_del = 0
        m_qty = m_mqty = 0
        for r in rows:
            if not r:
                continue
            total_rows += 1
            texts = [str(cell(r, cols[k]) or "") for k in
                     ("产品名称", "品类orig", "品类", "品类统一名")]
            sku = str(cell(r, cols["通途SKU"]) or "").strip()
            is_tri = any("三角" in t for t in texts) or (sku in reverse)
            if not is_tri:
                continue
            tri_rows += 1
            m_tri += 1
            try:
                qty = float(cell(r, cols["发货数量"]) or 0)
            except (TypeError, ValueError):
                qty = 0.0
            tri_qty += qty
            m_qty += qty
            wh = str(cell(r, cols["发货仓库"]) or "").strip()
            ch = str(cell(r, cols["渠道"]) or "").strip()
            nwh = normalize_warehouse(wh)
            by_warehouse[nwh] += qty
            by_channel[ch] += qty
            is_g1 = nwh in G1_WAREHOUSES
            if is_g1:
                g1["rows"] += 1
                g1["qty"] += qty
                g1["by_warehouse"][nwh] += qty
            if sku in reverse:
                items = reverse[sku]
                if len(items) == 1:
                    mapped_rows += 1
                    mapped_qty += qty
                    m_mapped += 1
                    m_mqty += qty
                    if is_g1:
                        g1["mapped_rows"] += 1
                        g1["mapped_qty"] += qty
                    item = next(iter(items))
                    fam = item[:6]
                    by_family[fam] += 1
                    by_family_qty[fam] += qty
                else:
                    ambiguous_rows += 1
                    m_amb += 1
                    if is_g1:
                        g1["ambiguous_rows"] += 1
            elif sku in deleted:
                deleted_only_rows += 1
                m_del += 1
                if is_g1:
                    g1["deleted_only_rows"] += 1
            else:
                missing_rows += 1
                missing_qty += qty
                m_miss += 1
                if is_g1:
                    g1["missing_rows"] += 1
                    g1["missing_qty"] += qty
                prefix = sku.split("-")[0] if sku else "(empty)"
                missing_prefix[prefix] += 1
        wb.close()
        by_month.append({
            "month": f"2026-{m:02d}", "file": os.path.basename(p),
            "tri_rows": m_tri, "tri_qty": round(m_qty, 2),
            "mapped_rows": m_mapped, "mapped_qty": round(m_mqty, 2),
            "ambiguous_rows": m_amb, "missing_rows": m_miss,
            "deleted_only_rows": m_del,
        })
    return {
        "total_rows": total_rows,
        "tri_rows": tri_rows,
        "tri_qty": round(tri_qty, 2),
        "mapped_rows": mapped_rows,
        "mapped_qty": round(mapped_qty, 2),
        "mapped_row_rate": round(mapped_rows / tri_rows, 4) if tri_rows else None,
        "mapped_qty_rate": round(mapped_qty / tri_qty, 4) if tri_qty else None,
        "ambiguous_rows": ambiguous_rows,
        "missing_rows": missing_rows,
        "missing_qty": round(missing_qty, 2),
        "deleted_only_rows": deleted_only_rows,
        "g1_scope": {
            "warehouses": sorted(G1_WAREHOUSES),
            "rows": g1["rows"],
            "qty": round(g1["qty"], 2),
            "mapped_rows": g1["mapped_rows"],
            "mapped_qty": round(g1["mapped_qty"], 2),
            "mapped_row_rate": (round(g1["mapped_rows"] / g1["rows"], 4)
                                if g1["rows"] else None),
            "mapped_qty_rate": (round(g1["mapped_qty"] / g1["qty"], 4)
                                if g1["qty"] else None),
            "ambiguous_rows": g1["ambiguous_rows"],
            "missing_rows": g1["missing_rows"],
            "missing_qty": round(g1["missing_qty"], 2),
            "deleted_only_rows": g1["deleted_only_rows"],
            "by_warehouse_qty": {k: round(v, 2) for k, v in g1["by_warehouse"].items()},
        },
        "by_month": by_month,
        "by_warehouse_qty": dict(by_warehouse.most_common()),
        "by_channel_qty": dict(by_channel.most_common()),
        "by_family_rows": dict(by_family),
        "by_family_qty": {k: round(v, 2) for k, v in by_family_qty.items()},
        "missing_top_prefix": dict(missing_prefix.most_common(15)),
    }


def main():
    bom = load_bom()
    reverse = bom.pop("reverse")
    deleted = bom.pop("deleted")
    orders = load_orders(reverse, deleted)
    out = {"bom": bom, "orders": orders}
    (OUT / "mapping_coverage.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(out, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
