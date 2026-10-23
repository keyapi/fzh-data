#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""G1b：三角靠枕 family 两仓当前库存快照（通途 ERP2，只读）。

只查 family 客户物料号（BOM 20261016，KS0001/KS0248），按仓库分块查询，
输出聚合 JSON（SKU × 仓 × 可用/在途/次品/待发），不落原始客户数据。

用法：
  uv run python g1b_stock_snapshot.py
  uv run python g1b_stock_snapshot.py --warehouses 美东-CENTRADE,美中-FZH-DANEEY
  uv run python g1b_stock_snapshot.py --chunk-size 50 --min-interval 13

注意：通途库存查询限制 5 次/分钟（商户级），脚本默认 13 秒节流。
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
OUT = BASE / "out"
OUT.mkdir(parents=True, exist_ok=True)

BOM = Path(r"D:/Work/王忠于/成本核算/EN产品BOM成本列表20261016.xlsx")
FAMILIES = ("KS0001", "KS0248")
DEFAULT_WAREHOUSES = [
    "美东-CENTRADE",
    "美中-FZH-DANEEY",
    "美中-FZH-DANEEY-皮壳仓库",
]

REPO_ROOT = next(
    (p for p in Path(__file__).resolve().parents if (p / "tongtool_api" / "mcp_http.py").exists()),
    Path(__file__).resolve().parents[7],
)
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))
from tongtool_api.mcp_http import McpClient, load_env  # noqa: E402

STOCK_QUERY_TOOL = "erp2_stocks_stocksquery"
CODE_OK = 200
CODE_RATE_LIMIT = 526


def default_env_file() -> Path:
    env = os.environ.get("TONGTOOL_ENV_FILE")
    if env:
        return Path(env)
    return Path(r"D:/Work/赛狐/Cursor/tongtool_api/.env")


def load_family_codes(path: Path = BOM) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    i_item = hdr.index("产品编号")
    i_code = hdr.index("客户物料号")
    codes: set[str] = set()
    for r in rows:
        item = str(r[i_item] or "").strip()
        code = str(r[i_code] or "").strip()
        if not item.startswith(FAMILIES):
            continue
        if not code or "已删除" in code:
            continue
        codes.add(code)
    wb.close()
    return sorted(codes)


def _to_num(value):
    """严格数值：空/非法返回 None，不把坏数静默转 0。"""
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def extract_payload(resp: dict) -> dict:
    if "error" in resp:
        raise RuntimeError(f"MCP error: {resp['error']}")
    texts = [
        c.get("text", "")
        for c in resp.get("result", {}).get("content", [])
        if c.get("type") == "text"
    ]
    raw = "\n".join(texts).strip()
    if not raw:
        return {}
    if raw[0] in "{[":
        parsed = json.loads(raw)
        return parsed if isinstance(parsed, dict) else {"datas": parsed}
    return {"text": raw}


def business_code(payload: dict):
    if not isinstance(payload, dict):
        return None
    for key in ("code", "businessCode", "statusCode"):
        if key in payload:
            try:
                return int(payload[key])
            except (TypeError, ValueError):
                return None
    return None


def rows_of(payload: dict) -> list:
    datas = payload.get("datas") if isinstance(payload, dict) else None
    if isinstance(datas, list):
        return datas
    if isinstance(datas, dict):
        for key in ("array", "list", "records", "dataList", "items", "rows", "stocks"):
            value = datas.get(key)
            if isinstance(value, list):
                return value
    return []


class StockClient:
    def __init__(self, mcp: McpClient, min_interval: float = 13.0, max_retries: int = 4):
        self.mcp = mcp
        self.min_interval = min_interval
        self.max_retries = max_retries
        self.last_call = 0.0
        self.requests = 0

    def _throttle(self) -> None:
        wait = self.min_interval - (time.monotonic() - self.last_call)
        if wait > 0:
            time.sleep(wait)
        self.last_call = time.monotonic()

    def _call(self, arguments: dict) -> dict:
        self._throttle()
        self.requests += 1
        return extract_payload(self.mcp.call(STOCK_QUERY_TOOL, arguments))

    def _call_with_retry(self, arguments: dict) -> dict:
        payload = self._call(arguments)
        if business_code(payload) != CODE_RATE_LIMIT:
            return payload
        delay = 60.0
        for _ in range(self.max_retries):
            print(f"  [526] 限流，等待 {delay:.0f}s 重试…", file=sys.stderr)
            time.sleep(delay)
            payload = self._call(arguments)
            if business_code(payload) != CODE_RATE_LIMIT:
                return payload
            delay *= 2
        raise RuntimeError("通途限流(526)重试后仍被拒")

    def query_chunk(self, warehouse: str, skus: list[str], page_size: int, max_pages: int):
        rows: list[dict] = []
        pages = 0
        last_len = None
        complete = False
        truncated = False
        for page in range(1, max_pages + 1):
            payload = self._call_with_retry(
                {
                    "warehouseName": warehouse,
                    "pageNo": page,
                    "pageSize": page_size,
                    "skus": skus,
                }
            )
            code = business_code(payload)
            if code not in (None, CODE_OK):
                raise RuntimeError(
                    f"{warehouse} 业务码 {code}: {json.dumps(payload, ensure_ascii=False)[:300]}"
                )
            if code is None and not isinstance(payload.get("datas"), (dict, list)):
                raise RuntimeError(f"{warehouse} 响应结构无法识别: {str(payload)[:300]}")
            batch = rows_of(payload)
            pages += 1
            last_len = len(batch)
            rows.extend(batch)
            if len(batch) < page_size:
                complete = True
                break
        else:
            if last_len == page_size:
                truncated = True
        return rows, {
            "pages": pages,
            "last_len": last_len,
            "complete": complete,
            "truncated": truncated,
        }


def main() -> int:
    ap = argparse.ArgumentParser(description="G1b family 两仓库存快照（只读）")
    ap.add_argument("--warehouses", default=",".join(DEFAULT_WAREHOUSES))
    ap.add_argument("--chunk-size", type=int, default=100)
    ap.add_argument("--page-size", type=int, default=100)
    ap.add_argument("--max-pages", type=int, default=20)
    ap.add_argument("--min-interval", type=float, default=13.0)
    ap.add_argument("--env-file", default=str(default_env_file()))
    ap.add_argument("--out", default="")
    args = ap.parse_args()

    warehouses = [w.strip() for w in args.warehouses.split(",") if w.strip()]
    codes = load_family_codes()
    print(f"family 客户码: {len(codes)}；仓库: {warehouses}")

    env = load_env(Path(args.env_file))
    key = env.get("TONGTOOL_ERP2_PRIMARY_KEY", "")
    secret = env.get("TONGTOOL_ERP2_PRIMARY_SECRET", "")
    if not key or not secret:
        print(f"✗ 缺少通途凭证：{args.env_file}", file=sys.stderr)
        return 2

    mcp = McpClient(key, secret, client_name="fzh-g1b-stock-snapshot")
    client = StockClient(mcp, min_interval=args.min_interval)

    started = time.monotonic()
    capture_start = datetime.now()
    by_warehouse: dict[str, dict] = {}
    by_sku: dict[str, dict] = {}
    errors: list[str] = []
    code_set = set(codes)
    chunks = [codes[i:i + args.chunk_size] for i in range(0, len(codes), args.chunk_size)]

    for wh in warehouses:
        wh_start = datetime.now()
        agg = {
            "rows": 0,
            "skus": 0,
            "available": 0.0,
            "intransit": 0.0,
            "defects": 0.0,
            "waiting": 0.0,
            "bad_rows": 0,
            "bad_fields": 0,
            "duplicates_ignored": 0,
            "duplicate_conflicts": 0,
            "bad_samples": [],
            "duplicate_samples": [],
            "page_meta": [],
            "truncated": False,
            "complete": True,
            "capture_start": wh_start.isoformat(timespec="seconds"),
            "capture_end": None,
        }
        seen: dict[tuple[str, str], dict] = {}
        print(f"\n仓库 {wh}：{len(chunks)} 块")
        for idx, chunk in enumerate(chunks, 1):
            try:
                raw_rows, meta = client.query_chunk(wh, chunk, args.page_size, args.max_pages)
            except Exception as e:  # noqa: BLE001 - 逐块记录，不中断整仓
                msg = f"{wh} chunk {idx}: {e}"
                errors.append(msg)
                agg["complete"] = False
                print(f"  ✗ {msg}", file=sys.stderr)
                continue
            agg["page_meta"].append({"chunk": idx, **meta})
            if meta.get("truncated"):
                agg["truncated"] = True
                agg["complete"] = False
            for r in raw_rows:
                sku = str(r.get("goodsSku") or "").strip()
                if sku not in code_set:
                    continue
                vals = {
                    "available": _to_num(r.get("availableStockQuantity")),
                    "intransit": _to_num(r.get("intransitStockQuantity")),
                    "defects": _to_num(r.get("defectsStockQuantity")),
                    "waiting": _to_num(r.get("waitingShipmentStockQuantity")),
                }
                bad_fields = [k for k, v in vals.items() if v is None]
                if bad_fields:
                    agg["bad_rows"] += 1
                    agg["bad_fields"] += len(bad_fields)
                    if len(agg["bad_samples"]) < 20:
                        agg["bad_samples"].append({"sku": sku, "fields": bad_fields})
                    agg["complete"] = False
                    continue
                key = (wh, sku)
                if key in seen:
                    if seen[key] == vals:
                        agg["duplicates_ignored"] += 1
                    else:
                        agg["duplicate_conflicts"] += 1
                        agg["complete"] = False
                        if len(agg["duplicate_samples"]) < 20:
                            agg["duplicate_samples"].append(
                                {"sku": sku, "kept": seen[key], "discarded": vals}
                            )
                    continue
                seen[key] = vals
                agg["rows"] += 1
                for k in ("available", "intransit", "defects", "waiting"):
                    agg[k] += vals[k]
                sku_entry = by_sku.setdefault(
                    sku, {"sku": sku, "warehouses": {}, "available": 0.0, "intransit": 0.0,
                          "defects": 0.0, "waiting": 0.0}
                )
                sku_entry["warehouses"][wh] = vals
                for k in ("available", "intransit", "defects", "waiting"):
                    sku_entry[k] += vals[k]
            if idx % 5 == 0 or idx == len(chunks):
                print(f"  块 {idx}/{len(chunks)}，累计 family 行 {agg['rows']}，SKU {len(seen)}")
        agg["skus"] = len(seen)
        agg["capture_end"] = datetime.now().isoformat(timespec="seconds")
        agg["complete"] = bool(agg["complete"] and not agg["truncated"])
        for k in ("available", "intransit", "defects", "waiting"):
            agg[k] = round(agg[k], 2)
        by_warehouse[wh] = agg
        print(f"  完成：{agg}")

    result = {
        "okf": "v0.1",
        "type": "Experiment-Results",
        "title": "G1b family 两仓库存快照（通途只读）",
        "date": datetime.now().strftime("%Y-%m-%d"),
        "fetched_at": datetime.now().isoformat(timespec="seconds"),
        "capture_start": capture_start.isoformat(timespec="seconds"),
        "capture_end": datetime.now().isoformat(timespec="seconds"),
        "snapshot_is_transactional": False,
        "complete": bool(not errors and all(v.get("complete") for v in by_warehouse.values())
                         and all(v.get("truncated") is False for v in by_warehouse.values())),
        "bad_rows_total": sum(v.get("bad_rows", 0) for v in by_warehouse.values()),
        "duplicate_conflicts_total": sum(v.get("duplicate_conflicts", 0)
                                          for v in by_warehouse.values()),
        "source": "tongtool erp2_stocks_stocksquery",
        "warehouses": warehouses,
        "family_codes": len(codes),
        "chunk_size": args.chunk_size,
        "requests": client.requests,
        "elapsed_seconds": round(time.monotonic() - started, 1),
        "by_warehouse": by_warehouse,
        "by_sku": sorted(by_sku.values(), key=lambda x: (-x["available"], x["sku"])),
        "errors": errors,
        "notes": [
            "只统计 BOM 20261016 的 family 客户物料号（排除含「已删除」的码）。",
            "available/intransit/defects/waiting 来自通途分仓库存；不是 EN Bin。",
            "WayfairCG 后台库存不在此文件；需 YTQ 手工导出。",
            "分块 + 13s 节流以避开通途 5 次/分钟限流。",
            "capture_start/end 是跨块采集窗口，不是严格瞬时全局事务。",
            "坏数/缺字段不转 0；重复 (warehouse, sku) 相同则忽略、不同则冲突并置 complete=false。",
            "满 max_pages 且末页仍满即 truncated=true、complete=false。",
        ],
    }

    out = Path(args.out) if args.out else OUT / f"stock_snapshot_{datetime.now():%Y%m%d_%H%M%S}.json"
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nOK: {out}")
    print(f"  requests={client.requests} elapsed={result['elapsed_seconds']}s errors={len(errors)}")
    return 0 if not errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
