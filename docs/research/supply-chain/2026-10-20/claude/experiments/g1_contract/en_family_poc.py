#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Family SO 批量只读 PoC（E01/E02 第一段，S02/S04 修复）。

- 日期口径显式：默认 delivery_date，可切 transaction_date；所有来源共用同一口径。
- family 角色：KS / PK# / ND# / group / customer；无法裁定记 pending，不直接算非 family。
- **S02 证据门**：PK#/ND# 前缀本身不是证据；只有证据表证明 KS 成品 BOM 消耗该组件
  （方向/比例/UOM/有效期/多 BOM 一致）才接纳为 family；结构证据不等于销量等价。
- **family 与 demand_role 分离**：family_membership/demand_role 分行记录，
  内胆独立销售不自动产生皮壳需求；qty 缺数/坏值不转 0。
- 来源 stats：候选行数、页数、终止依据；候选/选中/截断/详情失败逐层对平。
- **S04 逐行审计**：每个选中 SO 的每一行写审计 JSONL（family/pending/out_of_scope）；
  详情缺失（无 items/空 items/查询异常）不得静默成功；so/line 两层守恒入 manifest。
- 只读 EN；输出 family JSONL + 审计 JSONL + manifest。不含 DN/WO/JC/SE。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import sys
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

REPO_ROOT = next(
    (p for p in Path(__file__).resolve().parents if (p / "EN_API" / "item_shipment_status.py").exists()),
    Path(__file__).resolve().parents[7],
)
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import EN_API.item_shipment_status as iss  # noqa: E402

DEFAULT_GROUPS = ["三角靠枕", "三角靠枕无扣"]
FAMILY_PREFIXES = ("KS0001", "KS0248")
BOM = Path(r"D:/Work/王忠于/成本核算/EN产品BOM成本列表20261016.xlsx")
PILOT_CODES = {
    "CENCord995-DKGrey-97", "CEN003-DPBLUE-100",
    "C/Linen-Natural-100-688-wow", "C/Cord-DKBlue-100-713-wow",
    "C/Linen-Natural-194-688-wow", "C/Linen-Natural-153-688-wow",
    "C/Velvet-Grey-100-705-wow", "C/Cord-Grey-100-710-wow",
    "C/Linen-White-100-662-WOW", "TT0000671K0063721",
}


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def _no_dup_object(pairs):
    """json object hook：重复 key 直接报错，不静默覆盖（含 items 内的重复 code）。"""
    out = {}
    for key, value in pairs:
        if key in out:
            raise ValueError(f"duplicate JSON key: {key}")
        out[key] = value
    return out


def _canonical_json(obj) -> bytes:
    return json.dumps(obj, ensure_ascii=False, sort_keys=True,
                      separators=(",", ":")).encode("utf-8")


def _parse_iso_date(value):
    try:
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def load_family_manifest(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    i_item = hdr.index("产品编号")
    i_code = hdr.index("客户物料号")
    items: set[str] = set()
    codes: set[str] = set()
    for r in rows:
        item = str(r[i_item] or "").strip()
        code = str(r[i_code] or "").strip()
        if not item.startswith(FAMILY_PREFIXES):
            continue
        items.add(item)
        if code and "已删除" not in code:
            codes.add(code)
    wb.close()
    variants = set()
    for item in items:
        variants.update({item, f"PK#{item}", f"ND#{item}"})
    manifest = {
        "version": "family_manifest/20261016",
        "ks_items": sorted(items),
        "customer_codes": sorted(codes),
        "query_item_codes": sorted(variants),
        "roles": {"KS": "ks_finished (bare KS item)",
                  "PK#": "cover (皮壳); prefix only, requires structural evidence",
                  "ND#": "inner (内胆); prefix only, requires structural evidence"},
    }
    blob = json.dumps(manifest, ensure_ascii=False, sort_keys=True).encode("utf-8")
    manifest["sha256"] = sha256_bytes(blob)
    return manifest


def load_variant_evidence(path: Path) -> dict:
    """读取并校验 variant→KS 证据表（producer/consumer 信任契约的文件层）。

    `payload_sha256` 覆盖 schema_version/version/source/site/as_of/producer/
    rule_version/items；文件层不可信即 exit 2：不可读、重复 key、缺 producer/
    schema/site/as_of/rule_version/payload_sha256、payload_sha256 不匹配。
    条目级问题留给 classify_line 记 pending。
    """
    try:
        raw = path.read_bytes()
        doc = json.loads(raw.decode("utf-8"), object_pairs_hook=_no_dup_object)
    except (OSError, ValueError) as e:
        raise ValueError(f"variant manifest 不可读: {e}") from e
    if not isinstance(doc, dict):
        raise ValueError("variant manifest 顶层必须是对象")
    for field in ("schema_version", "site", "as_of"):
        if not str(doc.get(field) or "").strip():
            raise ValueError(f"variant manifest 缺 {field}")
    if not _parse_iso_date(doc.get("as_of")):
        raise ValueError("variant manifest as_of 必须是 YYYY-MM-DD")
    producer = doc.get("producer")
    if (not isinstance(producer, dict) or not str(producer.get("name") or "").strip()
            or not str(producer.get("version") or "").strip()):
        raise ValueError("variant manifest 缺 producer.name/version")
    rule_version = str(doc.get("rule_version") or "").strip()
    if not rule_version:
        raise ValueError("variant manifest 缺 rule_version")
    items = doc.get("items")
    if not isinstance(items, dict):
        raise ValueError("variant manifest 缺少 items 对象")
    payload_hash = str(doc.get("payload_sha256") or "").strip()
    payload_obj = {k: doc.get(k) for k in
                   ("schema_version", "version", "source", "site", "as_of",
                    "producer", "rule_version", "items")}
    actual_hash = sha256_bytes(_canonical_json(payload_obj))
    if not payload_hash:
        raise ValueError("variant manifest 缺 payload_sha256")
    if payload_hash != actual_hash:
        raise ValueError("variant manifest payload_sha256 不匹配")
    by_evidence = defaultdict(int)
    by_scope = defaultdict(int)
    for rec in items.values():
        if not isinstance(rec, dict):
            by_evidence["invalid_record"] += 1
            continue
        by_evidence[str(rec.get("evidence") or "missing")] += 1
        by_scope[str(rec.get("evidence_scope") or "missing")] += 1
    return {
        "path": str(path),
        "sha256": sha256_bytes(raw),
        "schema_version": str(doc.get("schema_version")),
        "version": str(doc.get("version") or ""),
        "source": str(doc.get("source") or ""),
        "site": str(doc.get("site")),
        "as_of": str(doc.get("as_of")),
        "producer": producer,
        "rule_version": rule_version,
        "payload_sha256": actual_hash,
        "items": items,
        "entries": len(items),
        "by_evidence": dict(sorted(by_evidence.items())),
        "by_scope": dict(sorted(by_scope.items())),
    }


def normalize_item(item: str) -> tuple[str, str]:
    for prefix, role in (("PK#", "PK"), ("ND#", "ND")):
        if item.startswith(prefix):
            return item[len(prefix):], role
    return item, "KS"


VARIANT_DIRECTIONS = {"ks_parent_variant_component"}
MEMBERSHIP_BY_ROLE = {"KS": "ks_finished", "PK": "cover", "ND": "inner"}


def _finite_positive(value) -> float | None:
    try:
        num = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(num) or num <= 0:
        return None
    return num


def _check_quantity(part, label: str) -> tuple[bool, str, str]:
    """校验一个 BOM 侧（component/output）的 stock 数量与 UOM 换算。"""
    if not isinstance(part, dict):
        return False, "variant_record_invalid", label
    for field in ("item_code", "uom", "stock_uom"):
        if not str(part.get(field) or "").strip():
            return False, "variant_record_invalid", f"{label}.{field}"
    qty = _finite_positive(part.get("qty"))
    stock_qty = _finite_positive(part.get("stock_qty"))
    if qty is None or stock_qty is None:
        return False, "variant_quantity_invalid", f"{label}.qty"
    uom = str(part["uom"]).strip()
    stock_uom = str(part["stock_uom"]).strip()
    if uom != stock_uom:
        factor = _finite_positive(part.get("conversion_factor"))
        if factor is None:
            return False, "variant_uom_conversion_missing", f"{label}.conversion_factor"
        if abs(qty * factor - stock_qty) > max(1e-9, abs(stock_qty) * 1e-9):
            return False, "variant_uom_conversion_mismatch", f"{label}.stock_qty"
    return True, "", ""


def _check_candidate(cand, code: str, base: str, direction: str, uom: str
                     ) -> tuple[bool, str, str]:
    """候选 BOM 必须已提交、active，且与 canonical/角色/方向/比例一致。"""
    if not isinstance(cand, dict):
        return False, "variant_record_invalid", "candidate"
    for field in ("name", "modified", "sha256", "parent_item", "child_item", "uom"):
        if not str(cand.get(field) or "").strip():
            return False, "variant_record_invalid", f"candidate.{field}"
    name = str(cand.get("name"))
    if int(cand.get("docstatus") or 0) != 1:
        return False, "variant_bom_not_submitted", f"candidate.docstatus:{name}"
    # 历史证据看“当时是否有效”，不是今天的 is_active；current planning 由
    # current_planning_eligible 与当前 active 另判。
    if cand.get("eligible_at_evidence") is not True:
        return False, "variant_bom_conflict", f"candidate.eligible_at_evidence:{name}"
    if str(cand.get("direction") or "") != direction:
        return False, "variant_bom_conflict", f"candidate.direction:{name}"
    if str(cand.get("parent_item") or "") != base:
        return False, "variant_bom_conflict", f"candidate.parent:{name}"
    if str(cand.get("child_item") or "") != code:
        return False, "variant_bom_conflict", f"candidate.child:{name}"
    if str(cand.get("uom") or "").strip() != uom:
        return False, "variant_bom_conflict", f"candidate.uom:{name}"
    comp = _finite_positive(cand.get("component_stock_qty"))
    out = _finite_positive(cand.get("output_stock_qty"))
    if comp is None or out is None:
        return False, "variant_quantity_invalid", f"candidate.qty:{name}"
    if abs(comp / out - 1.0) > 1e-9:
        return False, "variant_bom_conflict", f"candidate.ratio:{name}"
    return True, "", ""


def validate_variant_record(code: str, base: str, prefix_role: str, rec,
                            family_items: set[str], date_from, date_to,
                            line_uom: str = "") -> tuple[bool, str, str, str]:
    """按可复核字段判定 variant 的结构关系；只证明结构，不证明销量等价。

    方向固定为 `ks_parent_variant_component`（KS 成品 BOM 消耗 PK/ND 组件）；
    反过来要求 PK/ND 的 BOM 必有 KS 是错误假设。ratio = 组件 stock 累计 /
    KS 输出 stock 数量；多 BOM 必须一致，default 不消解冲突。
    """
    if not isinstance(rec, dict):
        return False, "variant_record_invalid", "", "record"
    if not str(rec.get("schema_version") or "").strip():
        return False, "variant_schema_missing", "", "schema_version"
    canonical = str(rec.get("canonical_ks") or "")
    if not canonical:
        return False, "variant_record_invalid", "", "canonical_ks"
    if canonical != base:
        return False, "variant_base_conflict", "", "canonical_ks"
    if base not in family_items:
        return False, "variant_base_not_family", "", "canonical_ks"
    if str(rec.get("role") or "") != prefix_role:
        return False, "variant_role_mismatch", "", "role"
    if str(rec.get("evidence") or "") != "structural":
        return False, "variant_evidence_not_structural", "", "evidence"
    scope = str(rec.get("evidence_scope") or "")
    if scope == "current_structure_only":
        return False, "variant_current_structure_only", "", "evidence_scope"
    if scope != "historical_verified":
        return False, "variant_historical_unknown", "", "evidence_scope"
    if not str(rec.get("historical_basis") or "").strip():
        return False, "variant_historical_basis_missing", "", "historical_basis"
    direction = str(rec.get("direction") or "")
    if direction not in VARIANT_DIRECTIONS:
        return False, "variant_direction_not_cover_consumption", "", "direction"
    valid_from = _parse_iso_date(rec.get("valid_from"))
    valid_to = _parse_iso_date(rec.get("valid_to"))
    if not valid_from or not valid_to:
        return False, "variant_record_invalid", "", "validity"
    if valid_from > date_from or valid_to < date_to:
        return False, "variant_out_of_validity", "", "validity"
    if str(rec.get("resolution") or "") != "consistent":
        return False, "variant_bom_conflict", "", "resolution"
    uom = str(rec.get("uom") or "").strip()
    if not uom:
        return False, "variant_record_invalid", "", "uom"
    comp = rec.get("component")
    out = rec.get("output")
    for label, part in (("component", comp), ("output", out)):
        ok, reason, detail = _check_quantity(part, label)
        if not ok:
            return False, reason, "", detail
    if str(comp["item_code"]) != code:
        return False, "variant_component_item_mismatch", "", "component.item_code"
    if str(out["item_code"]) != base:
        return False, "variant_output_item_mismatch", "", "output.item_code"
    if (str(comp["stock_uom"]).strip() != uom
            or str(out["stock_uom"]).strip() != uom):
        return False, "variant_uom_mismatch", "", "stock_uom"
    if line_uom and str(line_uom).strip() != str(comp["uom"]).strip():
        return False, "variant_line_uom_mismatch", "", "line.uom"
    ratio_calc = float(comp["stock_qty"]) / float(out["stock_qty"])
    if abs(ratio_calc - 1.0) > 1e-9:
        return False, "variant_ratio_not_1", "", "ratio"
    rec_ratio = _finite_positive(rec.get("ratio"))
    if rec_ratio is None or abs(rec_ratio - ratio_calc) > 1e-9:
        return False, "variant_ratio_mismatch", "", "ratio"
    item = rec.get("item")
    if not isinstance(item, dict):
        return False, "variant_record_invalid", "", "item"
    if str(item.get("name") or "") != code:
        return False, "variant_item_mismatch", "", "item.name"
    if not isinstance(rec.get("current_planning_eligible"), bool):
        return False, "variant_current_planning_flag_missing", "", "current_planning_eligible"
    if str(item.get("stock_uom") or "").strip() != uom:
        return False, "variant_uom_mismatch", "", "item.stock_uom"
    for field in ("sha256", "modified"):
        if not str(item.get(field) or "").strip():
            return False, "variant_record_invalid", "", f"item.{field}"
    candidates = rec.get("candidate_boms")
    if not isinstance(candidates, list) or not candidates:
        return False, "variant_candidates_missing", "", "candidate_boms"
    selected = rec.get("bom")
    if not isinstance(selected, dict):
        return False, "variant_record_invalid", "", "bom"
    sel_name = str(selected.get("name") or "")
    sel_sha = str(selected.get("sha256") or "")
    if not sel_name or not sel_sha:
        return False, "variant_record_invalid", "", "bom.name/sha256"
    matched = False
    for cand in candidates:
        ok, reason, detail = _check_candidate(cand, code, base, direction, uom)
        if not ok:
            return False, reason, "", detail
        if str(cand.get("name")) == sel_name:
            matched = True
            if str(cand.get("sha256")) != sel_sha:
                return False, "variant_record_invalid", "", "bom.sha256"
            if str(cand.get("modified") or "") != str(selected.get("modified") or ""):
                return False, "variant_record_invalid", "", "bom.modified"
    if not matched:
        return False, "variant_selected_bom_not_candidate", "", "bom.name"
    return True, "variant_structural_1to1", f"{sel_name}@{sel_sha[:12]}", ""


def classify_line(item: str, group: str, cust: str, groups: list[str],
                  family_items: set[str], family_codes: set[str],
                  variant_items: dict, variant_mode: str,
                  date_from, date_to, line_uom: str = ""
                  ) -> tuple[str, str, str, str, str]:
    """返回 (disposition, role, reason, evidence_ref, detail)。

    PK#/ND# 不能凭前缀或 `proven` 标签接纳：必须通过 validate_variant_record
    的方向/比例/UOM/有效期/多 BOM 一致性检查；结构证据不等于销量等价。
    """
    base, prefix_role = normalize_item(item)
    if prefix_role in ("PK", "ND"):
        if base not in family_items:
            return "pending", prefix_role, "variant_base_not_family", "", "canonical_ks"
        rec = variant_items.get(item)
        if rec is None:
            reason = ("variant_manifest_absent" if variant_mode == "none"
                      else "variant_not_in_manifest")
            return "pending", prefix_role, reason, "", ""
        ok, reason, evidence_ref, detail = validate_variant_record(
            item, base, prefix_role, rec, family_items, date_from, date_to, line_uom)
        if ok:
            return "family", prefix_role, reason, evidence_ref, ""
        return "pending", prefix_role, reason, "", detail
    if base.startswith(FAMILY_PREFIXES) or base in family_items:
        return "family", "KS", "ks_exact", "", ""
    if group in groups:
        return "family", "group", "group_match", "", ""
    if cust in family_codes:
        return "family", "customer", "customer_match", "", ""
    return "out_of_scope", "other", "not_family", "", ""


def derive_demand(disposition: str, role: str, cust: str
                  ) -> tuple[str, str, float | None, str]:
    """family_membership 与 demand_role 分离；family 不是数量加总开关。"""
    membership = MEMBERSHIP_BY_ROLE.get(role, "unknown")
    if disposition != "family" or membership == "unknown":
        return membership, "unknown", None, "not_counted"
    if membership == "cover":
        return membership, "independent_cover", 1, "count_cover_1to1"
    if membership == "inner":
        return membership, "assembly_component", 0, "no_cover_demand"
    if membership == "ks_finished":
        if cust.endswith("-Cover"):
            return membership, "independent_cover", 1, "count_cover_1to1"
        return membership, "terminal_finished", 1, "count_cover_1to1_event_unverified"
    return membership, "unknown", None, "not_counted"


def parse_qty(value) -> tuple[float | None, str]:
    """缺数/坏值/NaN/Inf/负数一律 None，不转 0。"""
    if value is None or value == "":
        return None, "missing"
    try:
        num = float(value)
    except (TypeError, ValueError):
        return None, "bad"
    if not math.isfinite(num) or num < 0:
        return None, "bad"
    return num, "ok"


def paged_so(filters: list, base_url: str, key: str, secret: str,
             label: str, page_size: int = 100):
    rows, start, pages, complete = [], 0, 0, False
    fields = iss.SO_FIELDS
    fallback = ["name", "status", "docstatus", "customer", "delivery_date",
                "transaction_date"]
    while pages < 200:
        status, body = iss.api_get_status(
            iss._build_url("Sales Order", filters, fields, start, page_size, "name asc"),
            base_url, key, secret,
        )
        if status == 417:
            fields = fallback
            status, body = iss.api_get_status(
                iss._build_url("Sales Order", filters, fields, start, page_size, "name asc"),
                base_url, key, secret,
            )
        if status != 200:
            raise iss.QueryError(f"{label}: HTTP {status} {str(body)[:140]}")
        pages += 1
        data = body.get("data", []) or []
        if not data:
            complete = True
            break
        rows.extend(data)
        if len(data) < page_size:
            complete = True
            break
        start += len(data)
    return rows, {"rows": len(rows), "pages": pages, "complete": complete,
                  "termination": "empty" if not data else ("short_page" if complete else "max_pages")}


def main():
    ap = argparse.ArgumentParser(description="Family SO 批量只读 PoC")
    ap.add_argument("--date-from", required=True)
    ap.add_argument("--date-to", required=True)
    ap.add_argument("--date-basis", choices=("delivery", "transaction"), default="delivery")
    ap.add_argument("--groups", default=",".join(DEFAULT_GROUPS))
    ap.add_argument("--bom", default=str(BOM))
    ap.add_argument("--no-item-codes", action="store_true")
    ap.add_argument("--out", required=True)
    ap.add_argument("--audit-out", default=None,
                    help="逐行审计 JSONL；默认 <out>.audit.jsonl")
    ap.add_argument("--variant-manifest", default=None,
                    help="PK#/ND# 证据表；不传则变体一律 pending")
    ap.add_argument("--max-sos", type=int, default=500)
    args = ap.parse_args()

    groups = [g.strip() for g in args.groups.split(",") if g.strip()]
    date_field = "delivery_date" if args.date_basis == "delivery" else "transaction_date"
    family = load_family_manifest(Path(args.bom))
    family_items = set(family["ks_items"])
    family_codes = set(family["customer_codes"])

    variant_items: dict = {}
    variant_mode = "none"
    variant_manifest = None
    variant_query_added: list[str] = []
    query_codes = list(family["query_item_codes"])
    if args.variant_manifest:
        try:
            variant_manifest = load_variant_evidence(Path(args.variant_manifest))
        except ValueError as e:
            print(f"✗ {e}", file=sys.stderr)
            return 2
        variant_items = variant_manifest["items"]
        variant_mode = "evidence"
        # 证据表里实际出现的 code 必须进入候选查询，避免“证据有、查询没查”。
        variant_query_added = sorted(set(variant_items) - set(query_codes))
        query_codes = sorted(set(query_codes) | set(variant_items))

    date_from = _parse_iso_date(args.date_from)
    date_to = _parse_iso_date(args.date_to)
    if not date_from or not date_to or date_from > date_to:
        print("✗ --date-from/--date-to 必须是 YYYY-MM-DD 且 from<=to", file=sys.stderr)
        return 2

    key, secret = iss.load_credentials("prod", iss.resolve_env_files())
    if not key or not secret:
        print("✗ EN 生产凭证未配置", file=sys.stderr)
        return 2
    base = iss.ENV_URLS["prod"]

    source_stats = {}
    matched_sources: dict[str, set[str]] = defaultdict(set)
    raw_all = []

    group_filter = iss.child_filter("Sales Order Item", "item_group", "in", groups)
    rows, stats = paged_so(group_filter + [[date_field, ">=", args.date_from],
                                           [date_field, "<=", args.date_to]],
                           base, key, secret, "SO(group)")
    source_stats["group"] = stats
    raw_all.extend((r, "group") for r in rows)

    if not args.no_item_codes:
        item_rows = 0
        item_pages = 0
        item_complete = True
        for chunk in iss.chunks(query_codes, iss.IN_CHUNK):
            rows, stats = paged_so(
                [["Sales Order Item", "item_code", "in", chunk],
                 [date_field, ">=", args.date_from],
                 [date_field, "<=", args.date_to]],
                base, key, secret, f"SO(item {chunk[0]}…)",
            )
            item_rows += stats["rows"]
            item_pages += stats["pages"]
            item_complete = item_complete and stats["complete"]
            raw_all.extend((r, "item_code") for r in rows)
        source_stats["item_code"] = {"rows": item_rows, "pages": item_pages,
                                     "complete": item_complete,
                                     "termination": "sum_of_chunks"}

    pilot_rows = 0
    pilot_pages = 0
    pilot_complete = True
    for code in sorted(PILOT_CODES):
        rows, stats = paged_so(
            [["Sales Order Item", "customer_item_code", "like", f"%{code}%"],
             [date_field, ">=", args.date_from],
             [date_field, "<=", args.date_to]],
            base, key, secret, f"SO(pilot {code})",
        )
        pilot_rows += stats["rows"]
        pilot_pages += stats["pages"]
        pilot_complete = pilot_complete and stats["complete"]
        raw_all.extend((r, "pilot") for r in rows)
    source_stats["pilot"] = {"rows": pilot_rows, "pages": pilot_pages,
                             "complete": pilot_complete, "termination": "sum_of_codes"}

    by_name = {}
    for row, src in raw_all:
        name = row.get("name")
        if not name:
            continue
        by_name.setdefault(name, row)
        matched_sources[name].add(src)
    candidate_sos = len(by_name)
    selected = sorted(by_name.values(), key=lambda r: r.get("name") or "")
    truncated = candidate_sos > args.max_sos
    selected = selected[:args.max_sos]

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    audit_path = (Path(args.audit_out) if args.audit_out
                  else out_path.parent / f"{out_path.stem}.audit.jsonl")
    run_id = datetime.now().strftime("%Y%m%dT%H%M%S")
    counts = defaultdict(int)
    disposition_reasons = defaultdict(int)
    pending_reasons = defaultdict(int)
    so_counts = defaultdict(int)
    so_detail_failures = []
    audit_rows = 0
    pilot_seen = set()
    pilot_seen_family = set()
    variant_validation_errors: list[dict] = []
    variant_validation_error_count = 0
    membership_counts = defaultdict(int)
    demand_role_counts = defaultdict(int)
    qty_status_counts = defaultdict(int)

    with out_path.open("w", encoding="utf-8") as f, \
            audit_path.open("w", encoding="utf-8") as af:
        for i, so in enumerate(selected):
            detail_error = ""
            try:
                d = iss.get_single("Sales Order", so["name"], base, key, secret)
            except Exception as e:  # noqa: BLE001 - 逐单隔离，失败写入审计
                d, detail_status, detail_error = None, "detail_error", type(e).__name__
            else:
                if not d:
                    detail_status = "detail_none"
                elif "items" not in d or d.get("items") is None:
                    detail_status = "items_missing"
                elif not isinstance(d.get("items"), list):
                    detail_status = "items_not_list"
                elif len(d.get("items")) == 0:
                    detail_status = "items_empty"
                else:
                    detail_status = "detail_ok"
            if i + 1 < len(selected):
                time.sleep(0.15)
            so_counts[detail_status] += 1
            if detail_status != "detail_ok":
                failure = {"so": so["name"], "reason": detail_status}
                if detail_error:
                    failure["error"] = detail_error
                so_detail_failures.append(failure)
                continue
            srcs = sorted(matched_sources.get(so["name"], []))
            for line in d.get("items", []):
                item = str(line.get("item_code") or "")
                cust = str(line.get("customer_item_code") or "")
                group = str(line.get("item_group") or "")
                base_item, prefix_role = normalize_item(item)
                line_uom = str(line.get("uom") or "")
                disposition, role, reason_code, evidence_ref, detail = classify_line(
                    item, group, cust, groups, family_items, family_codes,
                    variant_items, variant_mode, date_from, date_to, line_uom)
                if detail:
                    variant_validation_error_count += 1
                    if len(variant_validation_errors) < 200:
                        variant_validation_errors.append({
                            "code": item, "reason": reason_code, "detail": detail})
                customer_hit = cust in PILOT_CODES
                if customer_hit:
                    pilot_seen.add(cust)
                    if disposition == "family":
                        pilot_seen_family.add(cust)
                qty, qty_status = parse_qty(line.get("qty"))
                delivered, delivered_status = parse_qty(line.get("delivered_qty"))
                open_qty = (qty - delivered
                            if qty is not None and delivered is not None else None)
                membership, demand_role, cover_factor, counting_status = derive_demand(
                    disposition, role, cust)
                canonical_ks = (base_item if role in ("PK", "ND")
                                else (item if role == "KS" else None))
                variant_rec = (variant_items.get(item)
                               if prefix_role in ("PK", "ND") else None)
                current_planning_eligible = (
                    variant_rec.get("current_planning_eligible")
                    if isinstance(variant_rec, dict) else None)
                membership_counts[membership] += 1
                demand_role_counts[demand_role] += 1
                qty_status_counts[qty_status] += 1
                common = {
                    "run_id": run_id,
                    "so": so["name"],
                    "line_id": line.get("name"),
                    "so_status": d.get("status"),
                    "so_docstatus": d.get("docstatus"),
                    "transaction_date": d.get("transaction_date"),
                    "delivery_date": d.get("delivery_date"),
                    "item_code": item,
                    "base_item_code": base_item,
                    "prefix_role": prefix_role,
                    "customer_item_code": cust,
                    "customer_hit": customer_hit,
                    "item_group": group,
                    "family_membership": membership,
                    "canonical_ks": canonical_ks,
                    "canonical_ks_status": ("confirmed" if disposition == "family"
                                            else "candidate"),
                    "current_planning_eligible": current_planning_eligible,
                    "demand_role": demand_role,
                    "cover_equivalent_factor": cover_factor,
                    "demand_event_id": (f"{so['name']}:{line.get('name')}"
                                        if disposition == "family" else None),
                    "counting_status": counting_status,
                    "matched_sources": srcs,
                    "warehouse": line.get("warehouse") or d.get("set_warehouse"),
                    "uom": line.get("uom"),
                    "qty": qty,
                    "qty_status": qty_status,
                    "delivered_qty": delivered,
                    "delivered_qty_status": delivered_status,
                    "open_qty": open_qty,
                    "stock_qty": None,
                    "stock_qty_status": "requires_item_uom",
                    "amended_from": d.get("amended_from"),
                    "so_detail_status": detail_status,
                }
                af.write(json.dumps(dict(common, disposition=disposition,
                                         family_role=role,
                                         disposition_reason=reason_code,
                                         evidence_ref=evidence_ref,
                                         validation_detail=detail),
                                    ensure_ascii=False) + "\n")
                audit_rows += 1
                counts[disposition] += 1
                counts[f"role:{role}"] += 1
                disposition_reasons[reason_code] += 1
                if disposition == "pending":
                    pending_reasons[reason_code] += 1
                if disposition != "family":
                    continue
                f.write(json.dumps(dict(common, family_role=role,
                                        disposition_reason=reason_code,
                                        evidence_ref=evidence_ref),
                                    ensure_ascii=False) + "\n")

    detail_ok = so_counts.get("detail_ok", 0)
    hard_detail = any(f["reason"] in ("detail_none", "detail_error")
                      for f in so_detail_failures)
    soft_detail = any(f["reason"] in ("items_missing", "items_not_list", "items_empty")
                      for f in so_detail_failures)
    line_total = sum(counts[d] for d in ("family", "pending", "out_of_scope"))
    audit_checks = {
        "so_accounting": sum(so_counts.values()) == len(selected),
        "line_dispositions": line_total == audit_rows,
    }
    audit_ok = all(audit_checks.values())

    status = "complete"
    status_reasons = []
    if truncated:
        status = "partial"
        status_reasons.append("max_sos_truncated")
    if hard_detail:
        status = "partial"
        status_reasons.append("detail_failed")
    if soft_detail:
        status = "partial"
        status_reasons.append("items_unavailable")
    if any(not s.get("complete") for s in source_stats.values()):
        status = "partial"
        status_reasons.append("source_incomplete")
    if not audit_ok:
        status = "partial"
        status_reasons.append("audit_mismatch")
    manifest = {
        "run_id": run_id,
        "script_sha256": sha256_file(Path(__file__)),
        "family_manifest_version": family["version"],
        "family_manifest_sha256": family["sha256"],
        "family_manifest_counts": {
            "ks_items": len(family["ks_items"]),
            "customer_codes": len(family["customer_codes"]),
            "query_item_codes": len(family["query_item_codes"]),
        },
        "query_item_codes_effective": len(query_codes),
        "variant_query_codes_added": variant_query_added,
        "item_code_query_enabled": not args.no_item_codes,
        "date_basis": date_field,
        "date_from": args.date_from,
        "date_to": args.date_to,
        "groups": groups,
        "sources": source_stats,
        "candidate_sos": candidate_sos,
        "selected_sos": len(selected),
        "truncated": truncated,
        "max_sos": args.max_sos,
        "detail_failed": [f["so"] for f in so_detail_failures],
        "so_audit": {
            "selected": len(selected),
            "detail_ok": detail_ok,
            "detail_none": so_counts.get("detail_none", 0),
            "items_missing": so_counts.get("items_missing", 0),
            "items_not_list": so_counts.get("items_not_list", 0),
            "items_empty": so_counts.get("items_empty", 0),
            "detail_error": so_counts.get("detail_error", 0),
            "failures": so_detail_failures,
        },
        "line_dispositions": {k: v for k, v in counts.items() if not k.startswith("role:")},
        "line_roles": {k[5:]: v for k, v in counts.items() if k.startswith("role:")},
        "family_membership_counts": dict(sorted(membership_counts.items())),
        "demand_role_counts": dict(sorted(demand_role_counts.items())),
        "qty_status_counts": dict(sorted(qty_status_counts.items())),
        "stock_qty_status": "requires_item_uom",
        "disposition_reasons": dict(sorted(disposition_reasons.items())),
        "pending_reasons": dict(sorted(pending_reasons.items())),
        "line_audit": {
            "rows": audit_rows,
            "file": str(audit_path),
            "sha256": sha256_file(audit_path) if audit_path.exists() else "",
            "dispositions": {k: v for k, v in counts.items() if not k.startswith("role:")},
            "pending_reasons": dict(sorted(pending_reasons.items())),
        },
        "audit_checks": audit_checks,
        "audit_ok": audit_ok,
        "audit_sha256": sha256_file(audit_path) if audit_path.exists() else "",
        "variant_mode": variant_mode,
        "variant_manifest": ({
            "path": variant_manifest["path"],
            "schema_version": variant_manifest["schema_version"],
            "version": variant_manifest["version"],
            "source": variant_manifest["source"],
            "site": variant_manifest["site"],
            "as_of": variant_manifest["as_of"],
            "producer": variant_manifest["producer"],
            "rule_version": variant_manifest["rule_version"],
            "payload_sha256": variant_manifest["payload_sha256"],
            "sha256": variant_manifest["sha256"],
            "entries": variant_manifest["entries"],
            "by_evidence": variant_manifest["by_evidence"],
            "by_scope": variant_manifest["by_scope"],
        } if variant_manifest else None),
        "variant_validation_error_count": variant_validation_error_count,
        "variant_validation_errors": variant_validation_errors,
        "adjudication": "pending_remaining" if pending_reasons else "complete",
        "adjudication_reasons": dict(sorted(pending_reasons.items())),
        "pilot_expected": sorted(PILOT_CODES),
        "pilot_seen": sorted(pilot_seen),
        "pilot_seen_family": sorted(pilot_seen_family),
        "pilot_missing": sorted(PILOT_CODES - pilot_seen),
        "status": status,
        "status_reasons": status_reasons,
        "out": str(out_path),
        "out_sha256": sha256_file(out_path),
        "notes": [
            "只读 SO 批量 PoC；不含 DN/WO/JC/SE。",
            "所有来源共用同一日期口径；delivery 与 transaction 不混算。",
            "PK#/ND# 前缀不是证据；无证据表或未命中一律 pending（variant_* 原因码）。",
            "PK=皮壳、ND=内胆；证据方向必须是 KS 成品 BOM 消耗该组件，"
            "ratio=组件 stock 累计/KS 输出 stock，多 BOM 冲突不按 default 消解。",
            "结构证据只证明 family 关系，不证明销量等价；历史窗口需 "
            "historical_verified + 有效期，current_structure_only 一律 pending。",
            "family_membership 与 demand_role 分离；SO 行不等于终端销售，"
            "本 PoC 不做事件级去重、不算终端需求。",
            "qty/stock_qty：缺数/坏值不转 0；stock_qty 需 Item stock UOM，"
            "本 PoC 记 None（requires_item_uom）。",
            "每行 disposition 写入审计 JSONL；so/line 两层守恒由 audit_checks 校验。",
            "status 只表示查询/详情完整性；逐行证据完成度看 adjudication。",
            "候选→选中→详情失败→逐行处置逐层记录。",
        ],
    }
    (out_path.parent / "family_poc_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({k: manifest[k] for k in
                      ("status", "candidate_sos", "selected_sos", "line_dispositions",
                       "pending_reasons", "audit_ok", "adjudication",
                       "pilot_missing")}, ensure_ascii=False))
    return 0 if status == "complete" else 1


if __name__ == "__main__":
    raise SystemExit(main())
