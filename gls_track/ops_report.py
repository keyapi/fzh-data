"""gls_track 运营异常报表生成器（仿 fedex_track.ops_report；口径同 PR#215 共享 classify）。

读 gls_track CLI 产出的 summary.csv + 通途原始表，输出**多工作表 + 配色 + 汇总**的 Excel，
与 FedEx 运营异常表同一套判定语义与版式（Amazon 口径 · 营业日）。

GLS 时点映射（对标 FedEx 建标/收件/交付）：
- 建标 ≈ **数据录入 GLS IT**（history 首条 "data was entered..."）
- 站点收件 ≈ **交接 GLS**（"was handed over to GLS"）
- 交付 = delivered 事件 / arrivalTime

> 定位：GLS 先行版。判定结构/版式仿 fedex_track.ops_report，但**日历/阈值独立**：
> HANDLING_DAYS=2、营业日按波兰 2026 公共假日（起运/交接在 GLS 波兰，不用美国联邦假日）。
> 待 PR#215 (parcel_track) 合入后，GLS adapter 应改走 parcel_track 共享 classify + ops_excel，
> 本文件届时降为薄壳或删除。
"""

from __future__ import annotations

import datetime as _dt
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── GLS 口径（可调）─────────────────────────────────────────
# 迟发允许处理时间：数据录入→交接 GLS，允许 2 个营业日（周末/假日顺延）。
# FedEx 表用 1（美国市场）；GLS 起运/交接都在波兰，仓库+GLS 排程按波兰，定 2。
HANDLING_DAYS = 2
TRANSIT_SLOW_DAYS = 6
TRANSIT_SEVERE_DAYS = 12
STUCK_DAYS = 7
STUCK_SEVERE_DAYS = 14
MISSING_AFTER_DAYS = 3
# 日历：数据录入/交接发生在 GLS 波兰(货主国/起运国) → 用波兰 2026 法定公共假日；
# 周末由 np.busday_count 天然排除，故周日型假日(Easter 4/5、5/3、Pentecost 5/24、11/1、8/15六、12/26六)
# 无需重复列入。12/24(Wigilia)自 2025 起为波兰法定假日，须列入。
# 来源(2026-09-07 复核)：https://www.timeanddate.com/holidays/poland/2026 、
# https://getsix.eu/human-resources-payroll-in-poland/public-holidays-in-poland-in-2026 、
# https://kadry.infor.pl/...（Ustawa o dniach wolnych od pracy 法定清单）。
PL_HOLIDAYS_2026 = [
    "2026-01-01", "2026-01-06", "2026-04-06", "2026-05-01", "2026-05-03",
    "2026-06-04", "2026-08-15", "2026-11-01", "2026-11-11",
    "2026-12-24", "2026-12-25", "2026-12-26",
]


def bizdays(d1, d2, holidays=None):
    """d1→d2 的营业日数（默认波兰 2026 公共假日；周末不计）。解析失败返回 None。"""
    if pd.isna(d1) or pd.isna(d2):
        return None
    try:
        return int(np.busday_count(np.datetime64(d1.date()), np.datetime64(d2.date()),
                                   holidays=holidays or PL_HOLIDAYS_2026))
    except Exception:
        return None


def _late_level(overdue):
    if overdue is None:
        return "迟发"
    if overdue <= 0:
        return "准时"
    if overdue <= 2:
        return "轻度迟发"
    if overdue <= 5:
        return "中度迟发"
    return "重度迟发"


# 身份提取等跨模块复用（与通途列名一致，单一来源；无日历耦合）
from fedex_track.ops_report import _bare_tracking, _to_ts, load_tt_identity  # noqa: E402

COLOR = {"green": "C6EFCE", "yellow": "FFEB9C", "orange": "FCD5B4", "red": "FFC7CE",
         "gray": "D9D9D9", "blue": "DDEBF7", "dark": "404040", "header": "2F5597"}

# key → (中文, 英文, 基础色)。与 parcel_track/classify.CLASS 同词表；承运延误用 generic 显示。
CLASS = {
    "missing_not_handed": ("漏发/未交接", "not_handed", "red"),
    "fresh_no_pickup": ("建标未收件", "label_no_pickup", "yellow"),
    "late_handover": ("迟发", "late_handover", "orange"),
    "carrier_slow": ("承运延误", "carrier_slow", "orange"),
    "stuck": ("卡件", "stuck", "red"),
    "cancelled": ("已取消", "cancelled", "gray"),
    "not_found": ("数据异常/查无", "not_found", "gray"),
    "reused_no_label": ("复用旧票(缺建标)", "reused_no_label", "blue"),
    "delivered_ok": ("正常交付", "delivered_ok", "green"),
    "in_transit": ("在途", "in_transit", "yellow"),
}
LEVEL_COLOR = {
    "漏发/未交接": "red", "卡件": "red", "严重延误": "red", "重度迟发": "red",
    "中度迟发": "orange", "承运延误": "orange", "GLS延误": "orange",
    "轻度迟发": "yellow", "建标未收件": "yellow", "在途": "yellow",
    "已取消": "gray", "数据异常/查无": "gray", "复用旧票(缺建标)": "blue",
    "正常交付": "green", "准时": "green",
}
SHEET_OF_CLASS = {
    "missing_not_handed": "漏发未交接", "fresh_no_pickup": "漏发未交接",
    "late_handover": "迟发", "carrier_slow": "承运异常", "stuck": "承运异常",
    "cancelled": "取消·其他", "not_found": "取消·其他",
}
ACTION = {
    "missing_not_handed": "通知仓库/货代核查是否漏发或未交接；未发出→取消/补发，已交但无扫描→开GLS trace",
    "fresh_no_pickup": "持续关注是否转漏发；与仓确认交接；超过N天按漏发处理",
    "late_handover": "交接SLA复盘；对已延误买家提前告知/安抚；根因治理（仓、货代排程）",
    "carrier_slow": "记录；超过严重阈值→开GLS trace/索赔",
    "stuck": "开GLS trace调查；超14天→索赔+重发/退款",
    "cancelled": "确认取消原因（我方撤单/货代/GLS），按需退款或重下",
    "not_found": "核查通途源数据（非GLS单号/拼接号），改走对应渠道",
    "reused_no_label": "复用号旧票，从建标开始的对应票为准；无需处理",
    "delivered_ok": "—", "in_transit": "—",
}


def _cat(dev, pu, label, ship, now, last_event=None):
    """GLS 分类 key：迟发=数据录入→交接GLS 营业日；承运延误=交接→交付；卡件看最近扫描；延误优先。"""
    if pu is pd.NaT:
        if dev is not pd.NaT:
            return "reused_no_label" if label is pd.NaT else "delivered_ok"
        if ship is not pd.NaT and (now - ship).days > MISSING_AFTER_DAYS:
            return "missing_not_handed"
        return "fresh_no_pickup"
    if label is pd.NaT:
        return "reused_no_label"
    late = False
    bd = bizdays(label, pu)
    if bd is not None and (bd - HANDLING_DAYS) > 0:
        late = True
    if dev is pd.NaT:
        scan = last_event if last_event is not pd.NaT and last_event is not None else label
        if scan is not pd.NaT and (now - scan).days > STUCK_DAYS:
            return "stuck"
        return "late_handover" if late else "in_transit"
    trans = bizdays(pu, dev)
    if trans is not None and trans > TRANSIT_SLOW_DAYS:
        return "carrier_slow"
    if late:
        return "late_handover"
    return "delivered_ok"


def build(summary_csv: str, tt_xlsx: str, out_xlsx: str):
    S = pd.read_csv(summary_csv, dtype=str).fillna("")
    tt = load_tt_identity(tt_xlsx)
    now = pd.Timestamp(_dt.datetime.now())

    COLS = ["跟踪号", "分类(EN)", "分类(中文)", "等级", "订单号", "包裹号", "渠道账号", "渠道",
            "平台SKU", "通途SKU", "产品名称", "发货仓库", "执行发货人", "是否补发货", "销售站点",
            "买家姓名", "国家", "城市", "省/州", "发货日期", "数据录入时间", "交接GLS时间", "交付时间",
            "日历延迟(天)", "营业日延迟(天)", "承运延迟(营业日)", "Amazon是否判迟", "所属Sheet", "建议动作", "处理状态"]

    rows = []
    for _, r in S.iterrows():
        n = _bare_tracking(r["跟踪号"])
        err = str(r.get("错误") or "")
        dev = _to_ts(r.get("交付时间"))
        pu = _to_ts(r.get("交接GLS时间"))
        label = _to_ts(r.get("数据录入时间"))
        last_event = _to_ts(r.get("最近事件时间"))
        status = str(r.get("当前状态") or "").strip().upper()
        ident = tt.get(n, tt.get(r["跟踪号"], {}))
        ch = "{} {}".format(ident.get("渠道", ""), ident.get("销售站点", "")).lower()
        ship = _to_ts(ident.get("发货日期"))

        cat_key = _cat(dev, pu, label, ship, now, last_event=last_event)
        if err:  # rstt028/029 非 200（GLS 无此单/非 GLS 号）
            cat_key = "not_found"
        elif "CANCEL" in status and dev is pd.NaT:
            cat_key = "cancelled"
        elif cat_key in ("delivered_ok",) and status not in ("DELIVERED", "") and dev is not pd.NaT \
                and last_event is not pd.NaT and last_event is not None and last_event > dev:
            # 返件等：先有交付事件后又回退，头条非 DELIVERED → 不标正常交付
            cat_key = "in_transit"
        zh, en, base = CLASS[cat_key]

        overdue = None
        bd = bizdays(label, pu)
        if bd is not None:
            overdue = bd - HANDLING_DAYS
        trans = bizdays(pu, dev) if (dev is not pd.NaT and pu is not pd.NaT) else None
        if cat_key == "late_handover":
            level = _late_level(overdue)
        elif cat_key == "carrier_slow":
            level = "严重延误" if (trans is not None and trans > TRANSIT_SEVERE_DAYS) else "承运延误"
        else:
            level = zh
        # Amazon是否判迟：仅 Amazon 渠道有意义（PR#215 口径）；非 Amazon 渠道不填"是"
        is_amz = any(k in ch for k in ("amazon", "amz", "亚马逊"))
        if overdue is None:
            amazon = ""
        elif is_amz:
            amazon = "是" if overdue > 0 else "否"
        else:
            amazon = "否" if overdue <= 0 else ""
        caldays = int((pu - label).days) if (pu is not pd.NaT and label is not pd.NaT) else ""

        rows.append({
            "跟踪号": r["跟踪号"], "分类(EN)": en, "分类(中文)": zh, "等级": level,
            "订单号": ident.get("订单号", ""), "包裹号": ident.get("包裹号", ""),
            "渠道账号": ident.get("渠道账号", ""), "渠道": ident.get("渠道", ""),
            "平台SKU": ident.get("平台SKU", ""), "通途SKU": ident.get("通途SKU", ""),
            "产品名称": ident.get("产品名称", "")[:40], "发货仓库": ident.get("发货仓库", ""),
            "执行发货人": ident.get("执行发货人", ""), "是否补发货": ident.get("是否补发货", ""),
            "销售站点": ident.get("销售站点", ""), "买家姓名": ident.get("买家姓名", ""),
            "国家": ident.get("国家/地区", ""), "城市": ident.get("城市", ""), "省/州": ident.get("省/州", ""),
            "发货日期": ship.date() if ship is not pd.NaT and not pd.isna(ship) else "",
            "数据录入时间": label, "交接GLS时间": pu, "交付时间": dev,
            "日历延迟(天)": caldays, "营业日延迟(天)": overdue, "承运延迟(营业日)": trans,
            "Amazon是否判迟": amazon, "所属Sheet": SHEET_OF_CLASS.get(cat_key, ""),
            "建议动作": ACTION[cat_key], "处理状态": "",
        })
    df = pd.DataFrame(rows, columns=COLS)
    df["_key"] = df["分类(EN)"].map({v[1]: k for k, v in CLASS.items()})

    anomaly_cats = ["missing_not_handed", "fresh_no_pickup", "late_handover",
                    "carrier_slow", "stuck", "cancelled", "not_found"]

    def pick(*keys):
        return df[df["_key"].isin(keys)]

    wb = Workbook()
    thin = Border(*[Side(style="thin", color="BFBFBF")] * 4)

    # ── 总览 ──
    ws = wb.active; ws.title = "总览"
    ws.cell(row=1, column=1, value="GLS 运营异常总览（Amazon 口径 · 营业日）").font = Font(size=16, bold=True, color="FFFFFF")
    ws.cell(row=1, column=1).fill = PatternFill("solid", fgColor="404040")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    ws.row_dimensions[1].height = 24
    n_total = len(df); n_anom = len(pick(*anomaly_cats))
    kpi = [("总单(票)", n_total), ("异常(票)", n_anom),
           ("漏发/未交接", int(sum(df["_key"] == "missing_not_handed"))),
           ("迟发", int(sum(df["_key"] == "late_handover"))),
           ("承运延误", int(sum(df["_key"] == "carrier_slow"))),
           ("卡件", int(sum(df["_key"] == "stuck"))),
           ("已取消", int(sum(df["_key"] == "cancelled"))),
           ("数据异常", int(sum(df["_key"] == "not_found")))]
    for i, (k, v) in enumerate(kpi):
        col = 1 + i * 2
        ws.cell(row=3, column=col, value=k).font = Font(bold=True)
        c = ws.cell(row=4, column=col, value=v); c.font = Font(size=14, bold=True)
        c.alignment = Alignment(horizontal="center")
        if k in ("漏发/未交接", "卡件"):
            c.fill = PatternFill("solid", fgColor="FFC7CE")
    r = 6
    for c, h in enumerate(("分类(中文)", "数量", "所属Sheet"), 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
    order = ["missing_not_handed", "fresh_no_pickup", "late_handover", "carrier_slow", "stuck",
             "cancelled", "not_found", "reused_no_label", "in_transit", "delivered_ok"]
    for k in order:
        cnt = int(sum(df["_key"] == k))
        r += 1
        zh, _, col = CLASS[k]
        ws.cell(row=r, column=1, value=zh)
        ws.cell(row=r, column=2, value=cnt).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=3, value=SHEET_OF_CLASS.get(k, "全部明细"))
        for c in (1, 2, 3):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=COLOR[col])
    r += 2
    ws.cell(row=r, column=1, value="配色图例").font = Font(bold=True)
    for col, lab, note in [
        ("red", "需立即处理", "漏发/卡件/严重延误/重度迟发"), ("orange", "需关注", "中度迟发/承运延误"),
        ("yellow", "留意", "轻度迟发/建标未收件/在途"), ("blue", "信息", "复用号旧票"),
        ("gray", "已取消/数据异常", "取消 或 查无"), ("green", "正常/准时", "正常交付")]:
        r += 1
        ws.cell(row=r, column=1, value=lab).fill = PatternFill("solid", fgColor=COLOR[col])
        ws.cell(row=r, column=2, value=note).alignment = Alignment(wrap_text=True)
    for c in range(1, 4):
        ws.column_dimensions[get_column_letter(c)].width = 20

    # ── 明细工作表 ──
    detail_plan = [
        ("异常处理", pick(*anomaly_cats), True, "异常子集：漏发/迟发/承运延误/卡件/取消/数据异常，含建议动作"),
        ("漏发未交接", pick("missing_not_handed", "fresh_no_pickup"), True, "建议：通知仓库/货代核查漏发或未交接"),
        ("迟发", pick("late_handover"), True, "含 营业日延迟 与 Amazon是否判迟；请按实际 handling 校准"),
        ("承运异常", pick("carrier_slow", "stuck"), True, "承运延误→记录/严重开trace；卡件→开trace/索赔"),
        ("取消·其他", pick("cancelled", "not_found"), True, "已取消 / 数据异常·查无(如 allegro …U 非 GLS 号)"),
        ("全部明细", df, False, "全量（含正常/在途），审计与追溯"),
    ]
    widths = [16, 12, 14, 12, 14, 12, 12, 10, 12, 12, 24, 12, 10, 8, 10, 12, 8, 10, 8, 12, 18, 18, 18, 10, 10, 10, 10, 12, 58, 10]
    for name, data, colorize, note in detail_plan:
        ws2 = wb.create_sheet(name)
        if note:
            ws2.cell(row=1, column=1, value=note).font = Font(italic=True, color="808080")
            ws2.freeze_panes = "A3"
            start = 3
        else:
            start = 1
        for c, h in enumerate(COLS, 1):
            cell = ws2.cell(row=start, column=c, value=h)
            cell.fill = PatternFill("solid", fgColor="2F5597")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for i, (_, rd) in enumerate(data.iterrows()):
            rr = start + 1 + i
            for c, h in enumerate(COLS, 1):
                cell = ws2.cell(row=rr, column=c, value=rd[h])
                cell.border = thin
            if colorize:
                col = LEVEL_COLOR.get(rd["等级"], "green")
                ws2.cell(row=rr, column=1).fill = PatternFill("solid", fgColor=COLOR[col])
        for c, w in enumerate(widths, 1):
            ws2.column_dimensions[get_column_letter(c)].width = w

    # ── 口径说明 ──
    ws3 = wb.create_sheet("口径说明")
    notes = [
        ("GLS 运营异常报表口径说明（v0.1）", ""),
        ("1. 时点", "建标≈数据录入 GLS IT（history 首条 data entered）；收件≈交接 GLS（handed over）；交付=delivered/arrivalTime。"),
        ("2. 迟发(Amazon口径)", "ship-by=发货日期+处理时间(营业日)；周末与公共假日不计入。营业日延迟=数据录入→交接 营业日数-处理时间(默认2天)。"),
        ("3. 处理时间", f"默认 {HANDLING_DAYS} 个营业日（GLS 定 2；FedEx 表为 1），请按实际改脚本顶部 HANDLING_DAYS。"),
        ("4. 营业日排除", "周六日（np.busday_count 天然排除）+ **2026 波兰法定公共假日**：1/1,1/6,4/6,5/1,5/3,6/4,8/15,11/1,11/11,12/24,12/25,12/26（周日型假日 Easter4/5、Pentecost5/24 已被周末排除）。因数据录入/交接都发生在 GLS 波兰(起运国)，迟发判定用波兰历；美国联邦假日不适用，故未沿用 FedEx 表日历。"),
        ("5. 判定定义", f"漏发/未交接=有发货日期但无交接且>{MISSING_AFTER_DAYS}天；建标未收件=有数据录入但近期无交接；迟发=录入→交接营业日>处理时间({HANDLING_DAYS})；承运延误=交接→交付营业日>{TRANSIT_SLOW_DAYS}；卡件=在途且>{STUCK_DAYS}天无扫描；数据异常=rstt 接口无此单(如 allegro …U 行非 GLS 号)。"),
        ("6. 返件", "少数包裹先 DELIVERED 又回退(如 29626350320 08-10 派送→08-14 回 Strykow)，头条状态 INTRANSIT——本表按“先交付后回退+头条非DELIVERED”归为在途，不误标正常交付；请在通途/仓库侧按返件处理。"),
        ("7. 数据来源", "gls_full_202608.summary.csv（gls_track 公开无鉴权 REST，2026-09-07 跑 8 月通途 GLS 单）+ 通途非FBA订单202608.xlsx。"),
        ("8. 动作", "漏发/建标未收件→通知仓库核查；迟发→SLA复盘+安抚；承运延误→记录/严重开trace；卡件→开GLS trace/索赔；查无→核查源数据。"),
        ("9. 追踪", "GLS 单行先去重（同包裹多订单行）；无邮编只出摘要；明细需目的邮编。"),
        ("10. 日历边界", "欧盟无统一假日；Amazon 判迟本身按各站点国家历(amazon.de→德国、.fr→法国…)，无一套统一的 Amazon-EU 假日。本表口径=波兰(起运/交接国)一份日历作近似；承运延误(交接→交付)跨多国也只按同一历计。要逐站点精确，需按订单目的国/站点分别取假日再重算。"),
        ("11. 日历来源(2026-09-07 复核)", "波兰 2026 法定假日：timeanddate.com/holidays/poland/2026、getsix.eu/…/public-holidays-in-poland-in-2026、kadry.infor.pl（Dni wolne 2026）。12/24(Wigilia)自 2025 起为法定假日。Amazon 判迟为内部近似标记(非 Amazon 官方指标)：Amazon 是否按工作日计迟发口径不一(部分指标按自然日计，如 sellercentral 论坛 GD…)，仅作 ops 参考。"),
    ]
    for i, (k, v) in enumerate(notes, 1):
        ws3.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws3.cell(row=i, column=1).alignment = Alignment(vertical="top")
        ws3.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
        ws3.cell(row=i, column=2).fill = PatternFill("solid", fgColor="DDEBF7")
        ws3.row_dimensions[i].height = 40
        ws3.column_dimensions["A"].width = 22
        ws3.column_dimensions["B"].width = 120
    wb.save(out_xlsx)
    return df


if __name__ == "__main__":
    import argparse
    p = argparse.ArgumentParser()
    p.add_argument("--summary", required=True, help="gls_track query 产出的 summary.csv")
    p.add_argument("--tt", required=True, help="通途原始 xlsx")
    p.add_argument("--out", required=True, help="输出 Excel 路径")
    a = p.parse_args()
    df = build(a.summary, a.tt, a.out)
    print("total rows", len(df))
    print(df["分类(中文)"].value_counts().to_dict())
    print("written", a.out)
