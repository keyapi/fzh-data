"""离线单测：ops_report 分类 + 工作簿结构（合成 summary + 通途身份）。"""

import datetime as _dt

import numpy as np
import pandas as pd

from gls_track.ops_report import PL_HOLIDAYS_2026, build

# 时点相对真实「今天」计算：写死的绝对日期会随时间漂移，令迟发样本一旦越过
# STUCK_DAYS(7) 就必然被判成卡件（测试永久失败）。
NOW_DT = _dt.datetime.now()
NOW = NOW_DT.strftime("%Y-%m-%d")


def _biz_before(base: _dt.date, days: int) -> str:
    """base 之前 days 个营业日（按波兰 2026 假日）的时点串。"""
    d = np.busday_offset(np.datetime64(base), -days, roll="forward", holidays=PL_HOLIDAYS_2026)
    return f"{d} 06:00:00"


def _make_tt(tmp_path):
    df = pd.DataFrame([
        {"跟踪号": "GLSDEL", "订单号": "ORD-1", "发货日期": "2026-08-03", "邮寄方式": "GLS-Poland>>GLS-Poland",
         "渠道": "Mirakl", "渠道账号": "MXXXLDE", "产品名称": "cover", "国家/地区": "DE", "邮编": "12305"},
        {"跟踪号": "GLSRET", "订单号": "ORD-2", "发货日期": "2026-08-03", "邮寄方式": "GLS-Poland>>GLS-Poland",
         "渠道": "亚马逊", "渠道账号": "AMZDE", "产品名称": "cover", "国家/地区": "DE", "邮编": "47179"},
        {"跟踪号": "GLSFRESH", "订单号": "ORD-3", "发货日期": NOW, "邮寄方式": "GLS-Poland>>GLS-Poland",
         "渠道": "allegro", "渠道账号": "AllegroPL", "产品名称": "cover", "国家/地区": "PL", "邮编": "00-001"},
        {"跟踪号": "GLSLATE", "订单号": "ORD-4", "发货日期": "2026-08-03", "邮寄方式": "GLS-Poland>>GLS-Poland",
         "渠道": "Mirakl", "渠道账号": "MXXXLDE", "产品名称": "cover", "国家/地区": "DE", "邮编": "12305"},
        {"跟踪号": "ALLEGROU", "订单号": "ORD-5", "发货日期": "2026-08-03", "邮寄方式": "GLS-Poland>>GLS-Poland",
         "渠道": "allegro", "渠道账号": "AllegroPL", "产品名称": "cover", "国家/地区": "PL", "邮编": ""},
    ])
    path = tmp_path / "tt.xlsx"
    df.to_excel(path, index=False)
    return str(path)


def _make_summary(tmp_path):
    # 未交付的「迟发」样本：录入 → 交接 5 营业日（> HANDLING 3）；交接/最近扫描在 2 天前
    # （< STUCK_DAYS 7）→ 判迟发而非卡件。全程相对 NOW，杜绝时间漂移。
    late_pu_dt = NOW_DT - _dt.timedelta(days=2)
    late_pu = late_pu_dt.strftime("%Y-%m-%d %H:%M:%S")
    late_label = _biz_before(late_pu_dt.date(), 5)
    rows = [
        # delivered ok：录入08-03 → 交接08-04（1营业日，不迟）→ 交付08-05
        {"跟踪号": "GLSDEL", "国家/地区": "DE", "邮编": "12305", "当前状态": "DELIVERED", "状态说明": "Delivered",
         "已交付": "是", "交付时间": "2026-08-05 10:00:00", "数据录入时间": "2026-08-03 06:00:00",
         "交接GLS时间": "2026-08-04 09:00:00", "最近事件时间": "2026-08-05 10:00:00", "客户引用": "P1", "事件数": 8, "错误": ""},
        # 返件：先交付08-10 后回退08-14，头条 INTRANSIT → 不标正常交付
        {"跟踪号": "GLSRET", "国家/地区": "DE", "邮编": "47179", "当前状态": "INTRANSIT", "状态说明": "In transit",
         "已交付": "是", "交付时间": "2026-08-10 09:54:00", "数据录入时间": "2026-08-03 07:00:00",
         "交接GLS时间": "2026-08-04 18:00:00", "最近事件时间": "2026-08-14 04:55:00", "客户引用": "", "事件数": 17, "错误": ""},
        # 建标未收件：录入今天，无交接无交付
        {"跟踪号": "GLSFRESH", "国家/地区": "PL", "邮编": "00-001", "当前状态": "INTRANSIT", "状态说明": "In transit",
         "已交付": "否", "交付时间": "", "数据录入时间": f"{NOW} 06:00:00",
         "交接GLS时间": "", "最近事件时间": f"{NOW} 06:00:00", "客户引用": "", "事件数": 1, "错误": ""},
        # 迟发：录入 → 交接 5 营业日（-处理时间3 = 迟2 天）未交付，最近扫描 2 天前（不卡件）
        {"跟踪号": "GLSLATE", "国家/地区": "DE", "邮编": "12305", "当前状态": "INTRANSIT", "状态说明": "In transit",
         "已交付": "否", "交付时间": "", "数据录入时间": late_label,
         "交接GLS时间": late_pu, "最近事件时间": late_pu, "客户引用": "", "事件数": 6, "错误": ""},
        # 查无：非 GLS 单号
        {"跟踪号": "ALLEGROU", "国家/地区": "PL", "邮编": "", "当前状态": "", "状态说明": "",
         "已交付": "", "交付时间": "", "数据录入时间": "", "交接GLS时间": "", "最近事件时间": "", "客户引用": "", "事件数": 0,
         "错误": "GLS track rstt028/ALLEGROU 失败: ... http=404 category=not_found"},
    ]
    df = pd.DataFrame(rows)
    path = tmp_path / "summary.csv"
    df.to_csv(path, index=False, encoding="utf-8-sig")
    return str(path)


def test_build_workbook(tmp_path):
    out = tmp_path / "gls_ops.xlsx"
    df = build(_make_summary(tmp_path), _make_tt(tmp_path), str(out))
    from openpyxl import load_workbook
    wb = load_workbook(out)
    names = wb.sheetnames
    for expect in ["总览", "异常处理", "漏发未交接", "迟发", "承运异常", "取消·其他", "全部明细", "口径说明"]:
        assert expect in names, expect
    by = df["分类(EN)"].to_dict()
    assert by[0] == "delivered_ok"
    assert by[1] == "in_transit"       # 返件不误标正常
    assert by[2] == "label_no_pickup"   # 建标未收件
    assert by[3] == "late_handover"
    assert by[4] == "not_found"
    assert len(df) == 5
