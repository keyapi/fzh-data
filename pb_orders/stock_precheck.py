#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""SPS New 订单库存预检。

按原始 CSV 的 Detail 行用 ``Vendor Style`` 判断库存；输出仍保持 SPS 的 H/D 结构：

- 全有货 PO：保留 Header + 全部 Detail；
- 部分缺货 PO：保留 Header + 有货 Detail，只剔除缺货 Detail；
- 全部缺货 PO：整组剔除。

同时生成供人工操作 SPS 的 Excel。所有明细表始终并列显示对方 SKU
(``Buyers Catalog or Stock Keeping #``) 与我方 SKU (``Vendor Style``)。
"""

from __future__ import annotations

import os
import re
import shutil
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COL_PO = "PO Number"
COL_LINE = "PO Line #"
COL_RECORD = "Record Type"
COL_QTY = "Qty Ordered"
COL_VENDOR = "Vendor Style"
COL_BUYER = "Buyers Catalog or Stock Keeping #"
REQUIRED_COLUMNS = [COL_PO, COL_LINE, COL_RECORD, COL_QTY, COL_VENDOR, COL_BUYER]
DETAIL_COLUMNS = [COL_PO, COL_LINE, COL_BUYER, COL_VENDOR, COL_QTY]

MIME_CSV = "text/csv"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# 输出文件名里不能出现的字符（Windows 最严）：来源是用户上传的文件名，必须洗过。
_UNSAFE_STEM = re.compile(r'[\\/:*?"<>|\r\n\t]+')
# SPS 命名里的 PO 数记号：`order x21 20260917_0338_456788` 里的 `x21`
_PO_COUNT_TOKEN = re.compile(r"(?i)(?<![0-9A-Za-z])x\d+")

# 网页明细表最多渲染多少行（超了截断并提示，完整内容看 xlsx 产物）
WEB_TABLE_LIMIT = 300

ProgressFn = Callable[[str, str], None]


class StockCheckError(Exception):
    """库存预检的用户可读业务错误。"""

    def __init__(self, message: str, hint: str = "", code: str = "stock_check_error"):
        super().__init__(message)
        self.message = message
        self.hint = hint
        self.code = code

    def to_dict(self) -> dict:
        return {"code": self.code, "message": self.message, "hint": self.hint}


@dataclass
class StockCheckResult:
    checked_csv: Path
    operations_xlsx: Path
    report: dict


def _output_stem(raw: str) -> str:
    """把来源文件名整理成可安全拼进输出名的词干。

    历史命名是 `check0stock …`（待检查）→ `checked0stock …`（已检查），
    所以剥掉开头的 check/checked 记号，避免叠成 `checked0stock check0stock …`。
    用户文件名是外部输入，这里一并洗掉路径分隔符与 Windows 非法字符。
    """
    stem = str(raw or "").strip()
    lowered = stem.lower()
    for marker in ("checked0stock", "check0stock"):
        if lowered.startswith(marker):
            stem = stem[len(marker):].strip(" _-")
            break
    stem = _UNSAFE_STEM.sub("-", stem).strip(" .-")
    return stem or "orders"


def _with_po_count(stem: str, count: int) -> str:
    """把来源名里的 `x21` 换成**筛完剩下的** PO 数（`x18`）。

    SPS 自己重导时命名就是这么变的（数字跟流水号一起换），所以这边也照这个口径 ——
    文件名里的数字应当等于里面装了几个 PO。来源的时间戳/流水号原样留着，
    还能对回是哪一次导出。没有 `x{N}` 记号时原样返回（不去猜用户想叫它什么）。
    """
    replaced, hits = _PO_COUNT_TOKEN.subn(f"x{count}", stem, count=1)
    return replaced if hits else stem


def _noop(_step: str, _message: str) -> None:
    return None


def _normalized_skus(values: list[str]) -> tuple[list[str], set[str]]:
    snapshot: list[str] = []
    seen: set[str] = set()
    for value in values:
        sku = str(value or "").strip()
        key = sku.lower()
        if not sku or key in seen:
            continue
        seen.add(key)
        snapshot.append(sku)
    return snapshot, seen


def _read_csv(path: Path) -> pd.DataFrame:
    try:
        return pd.read_csv(path, dtype=str, keep_default_na=False)
    except Exception as exc:  # pandas 的具体解析异常不直接暴露给页面
        raise StockCheckError(
            "无法读取 SPS 订单 CSV",
            "请确认文件是 SPS 导出的 CSV，且没有被 Excel 另存为其它格式。",
            "csv_read_failed",
        ) from exc


def _source_lines(path: Path, rows: int) -> list[str]:
    """按行取源 CSV 原文，供 checked CSV **逐行照搬**。

    为什么不重新序列化：SPS 的导出本身是参差的 —— 表头 147 列、数据行 146 列，
    且最后一个列名是空的。pandas 读进来会把它命名成 `Unnamed: 146`、并把数据行
    补齐到 147 列，再写出去就不是 SPS 那种格式了（实测比历史 checked0stock 多出
    `Unnamed: 146` 与每行一个尾逗号）。照搬原始行则天然与源文件同构。

    行数对不上说明字段里含换行，此时按行切分不成立，宁可报错也不出格式不同的文件。
    """
    try:
        lines = path.read_text(encoding="utf-8-sig").splitlines()
    except Exception as exc:
        raise StockCheckError(
            "无法按行读取 SPS 订单 CSV",
            "请确认文件是 UTF-8 编码的 SPS 导出。",
            "csv_read_failed",
        ) from exc
    while lines and not lines[-1].strip():  # 容忍文件末尾多出的空行
        lines.pop()
    if len(lines) != rows + 1:
        raise StockCheckError(
            "源 CSV 的行数与解析出的记录数对不上（字段里可能含换行）",
            "请重新从 SPS 的 New 订单列表导出原始 CSV。",
            "ragged_source",
        )
    return lines


def _validate(df: pd.DataFrame) -> pd.DataFrame:
    missing = [column for column in REQUIRED_COLUMNS if column not in df.columns]
    if missing:
        raise StockCheckError(
            f"缺少必需列：{', '.join(missing)}",
            "请重新从 SPS 的 New 订单列表导出原始 CSV。",
            "missing_columns",
        )

    work = df.copy()
    work[COL_RECORD] = work[COL_RECORD].astype(str).str.strip().str.upper()
    unexpected = work.loc[~work[COL_RECORD].isin(["H", "D"]), COL_RECORD].unique().tolist()
    if unexpected:
        raise StockCheckError(
            f"Record Type 含无法识别的值：{', '.join(unexpected[:10])}",
            "库存预检只接受 SPS 的 Header(H) / Detail(D) 订单导出。",
            "invalid_record_type",
        )

    for column in (COL_PO, COL_LINE, COL_VENDOR, COL_BUYER):
        work[column] = work[column].astype(str).str.strip()

    details = work[work[COL_RECORD] == "D"]
    invalid_messages: list[str] = []
    for column in (COL_PO, COL_LINE, COL_VENDOR, COL_BUYER):
        count = int(details[column].eq("").sum())
        if count:
            invalid_messages.append(f"{column} 空值 {count} 行")

    qty = pd.to_numeric(details[COL_QTY], errors="coerce")
    invalid_qty = qty.isna() | (qty <= 0) | (qty % 1 != 0)
    if invalid_qty.any():
        invalid_messages.append(f"Qty Ordered 无效 {int(invalid_qty.sum())} 行")
    if invalid_messages:
        raise StockCheckError(
            "Detail 行数据不完整：" + "；".join(invalid_messages),
            "请在 SPS 核对对应 PO 明细后重新导出。",
            "invalid_detail",
        )

    duplicate = details.duplicated([COL_PO, COL_LINE], keep=False)
    if duplicate.any():
        samples = details.loc[duplicate, [COL_PO, COL_LINE]].drop_duplicates().head(10)
        text = ", ".join(f"{row[COL_PO]}-{row[COL_LINE]}" for _, row in samples.iterrows())
        raise StockCheckError(
            f"PO Line 重复：{text}",
            "同一个 PO 的明细行号必须唯一，请重新从 SPS 导出。",
            "duplicate_po_line",
        )

    po_values = details[COL_PO].drop_duplicates().tolist()
    headers = work[work[COL_RECORD] == "H"]
    header_counts = headers.groupby(COL_PO).size()
    bad_headers = [po for po in po_values if int(header_counts.get(po, 0)) != 1]
    orphan_headers = sorted(set(headers[COL_PO]) - set(po_values))
    if bad_headers or orphan_headers:
        parts = []
        if bad_headers:
            parts.append(f"Header 不是恰好 1 行的 PO：{', '.join(bad_headers[:10])}")
        if orphan_headers:
            parts.append(f"没有 Detail 的 Header：{', '.join(orphan_headers[:10])}")
        raise StockCheckError(
            "；".join(parts),
            "每个 PO 必须有 1 行 Header 和至少 1 行 Detail。",
            "invalid_header_detail",
        )

    work["_qty"] = 0
    work.loc[details.index, "_qty"] = qty.astype(int)
    return work


def _classify(work: pd.DataFrame, wanted: set[str]) -> tuple[pd.DataFrame, dict[str, str]]:
    classified = work.copy()
    detail_mask = classified[COL_RECORD].eq("D")
    classified["_is_no_stock"] = False
    classified.loc[detail_mask, "_is_no_stock"] = (
        classified.loc[detail_mask, COL_VENDOR].str.lower().isin(wanted)
    )

    po_status: dict[str, str] = {}
    for po, group in classified.loc[detail_mask].groupby(COL_PO, sort=False):
        flags = group["_is_no_stock"]
        if bool(flags.all()):
            po_status[po] = "全部缺货"
        elif bool(flags.any()):
            po_status[po] = "部分缺货"
        else:
            po_status[po] = "全部有货"
    classified["_po_status"] = classified[COL_PO].map(po_status)
    return classified, po_status


def _checked_rows(classified: pd.DataFrame, po_status: dict[str, str]) -> pd.DataFrame:
    retained_po = {po for po, status in po_status.items() if status != "全部缺货"}
    keep = classified[COL_PO].isin(retained_po) & (
        classified[COL_RECORD].eq("H")
        | (classified[COL_RECORD].eq("D") & ~classified["_is_no_stock"])
    )
    original_columns = [column for column in classified.columns if not column.startswith("_")]
    return classified.loc[keep, original_columns].copy()


def _detail_table(rows: pd.DataFrame, action: str) -> pd.DataFrame:
    table = rows[DETAIL_COLUMNS].copy()
    # SPS 原始导出里 `Qty Ordered` 是 `1.0` 这样的 float 字面量。照抄进去人在
    # Excel 里看到的就是 `1.0`（坑 8 同源）。这里的数量都是已校验的正整数，
    # 显示成整数；**源 CSV 不动**（那边是逐行照搬的）。
    table[COL_QTY] = [
        str(int(float(value))) if str(value).strip() else ""
        for value in rows[COL_QTY]
    ]
    table.insert(len(DETAIL_COLUMNS), "库存状态", rows["_is_no_stock"].map({True: "缺货", False: "有货"}).values)
    table["PO 库存状态"] = rows["_po_status"].values
    table["SPS 操作"] = action
    return table


def _build_tables(classified: pd.DataFrame, report: dict) -> dict[str, pd.DataFrame]:
    details = classified[classified[COL_RECORD] == "D"].copy()
    in_stock = details[~details["_is_no_stock"]]
    no_stock = details[details["_is_no_stock"]]
    mixed = details[details["_po_status"] == "部分缺货"]
    fully_no_stock = details[details["_po_status"] == "全部缺货"]

    overview = pd.DataFrame([
        ["原始 PO", report["po"]["total"], ""],
        ["全部有货 PO", report["po"]["fully_in_stock"], "可直接生成 ASN"],
        ["部分缺货 PO", report["po"]["partially_out_of_stock"], "不要取消整单；ASN 只勾选有货明细"],
        ["全部缺货 PO", report["po"]["fully_out_of_stock"], "不要生成 ASN；发送新日期通知"],
        ["原始 Detail", report["details"]["total"], ""],
        ["有货 Detail", report["details"]["in_stock"], "在 SPS 勾选对方 SKU 对应行"],
        ["缺货 Detail", report["details"]["out_of_stock"], "不要在 ASN 中勾选"],
    ], columns=["项目", "数量", "操作说明"])

    reconciliation = pd.DataFrame([
        ["PO", report["po"]["total"], report["po"]["fully_in_stock"],
         report["po"]["partially_out_of_stock"], report["po"]["fully_out_of_stock"], 0],
        ["Detail 行", report["details"]["total"], report["details"]["in_stock"], "",
         report["details"]["out_of_stock"], report["details"]["diff"]],
        ["数量", report["quantity"]["total"], report["quantity"]["in_stock"], "",
         report["quantity"]["out_of_stock"], report["quantity"]["diff"]],
    ], columns=["口径", "总数", "有货", "部分缺货", "缺货", "差数"])

    return {
        "操作总览": overview,
        "部分缺货PO-逐行操作": _detail_table(
            mixed,
            "同一 PO 不要整单取消；生成 ASN 时只勾选“有货”行，按对方 SKU 在 SPS 定位",
        ),
        "ASN有货明细": _detail_table(in_stock, "生成 ASN，并在 SPS 勾选此对方 SKU 明细行"),
        "缺货明细": _detail_table(no_stock, "不要在 ASN 中勾选；发送新日期通知"),
        "全部缺货PO": _detail_table(fully_no_stock, "整单不生成 ASN；发送新日期通知"),
        "检查报告": reconciliation,
        "原始剔除行": _detail_table(no_stock, "已从 checked CSV 剔除"),
    }


def _tables_payload(tables: dict[str, pd.DataFrame], limit: int) -> dict:
    """把操作表转成可进 `report_json` 的结构，供网页渲染明细表。

    与写进 xlsx 的是**同一批 DataFrame**，所以网页表与操作表不会各说各话。
    行数超上限时截断并**明确标出**（不静默丢），页面上提示去看完整 xlsx。

    先 `head(limit)` 再转列表：整表转成 Python 二维列表会为每一行都分配对象，
    大表（几万行）在 1G 的 worker 里没必要地吃内存。
    """
    payload: dict[str, dict] = {}
    for name, table in tables.items():
        head = table.head(limit).astype(str)
        payload[name] = {
            "columns": list(head.columns),
            "rows": head.values.tolist(),
            "total": len(table),
            "truncated": len(table) > limit,
        }
    return payload


def _excel_safe(table: pd.DataFrame) -> pd.DataFrame:
    """只给 xlsx 用：把以 `=` 开头的值转义，别让 openpyxl 写成**公式**。

    这些值来自外部 CSV（SPS 导出），`=cmd|…` 这类内容在 Excel 里会当公式执行。
    前面加 `'` 是标准的文本化处理；只影响这个人看的 xlsx，
    checked CSV 仍是逐行照搬的原始字节（那份才是给机器/回传用的）。
    """
    return table.map(lambda value: f"'{value}" if isinstance(value, str) and value.startswith("=") else value)


def _ensure_publishable(*targets: Path) -> None:
    """发布前统一探一遍目标文件：被别的程序占用就整体拒绝，不做半截发布。"""
    for target in targets:
        if not target.exists():
            continue
        try:
            with open(target, "ab"):
                pass
        except OSError as exc:
            raise StockCheckError(
                f"产物被占用，无法覆盖：{target.name}",
                "该文件可能正在 Excel 里打开，请关闭后重跑（或用 --out 换个目录）。",
                "output_locked",
            ) from exc


def _publish(temp: Path, dest: Path) -> bool:
    """把 temp 原子地发布到 dest，返回原来是否已有同名文件。

    用 `os.replace` 而不是 `shutil.move`：Windows 上同名文件 `os.rename` 会失败
    （同一批重跑就撞），`os.replace` 是原子覆盖、两个平台都对。
    """
    existed = dest.exists()
    os.replace(temp, dest)
    return existed


def _style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    no_stock_fill = PatternFill("solid", fgColor="F4CCCC")
    in_stock_fill = PatternFill("solid", fgColor="D9EAD3")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in ws.iter_cols():
            values = [str(cell.value or "") for cell in column[:200]]
            width = min(max(max((len(value) for value in values), default=0) + 2, 10), 48)
            ws.column_dimensions[get_column_letter(column[0].column)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        headers = {cell.value: cell.column for cell in ws[1]}
        status_col = headers.get("库存状态")
        po_status_col = headers.get("PO 库存状态")
        if status_col:
            for row in range(2, ws.max_row + 1):
                status = ws.cell(row, status_col).value
                fill = no_stock_fill if status == "缺货" else in_stock_fill
                for cell in ws[row]:
                    cell.fill = fill
        if po_status_col:
            for row in range(2, ws.max_row + 1):
                if ws.cell(row, po_status_col).value == "部分缺货":
                    ws.cell(row, po_status_col).fill = warning_fill
                    ws.cell(row, po_status_col).font = Font(bold=True, color="9C5700")

    wb.save(path)


def _verify_checked(path: Path, original_columns: list[str], expected_raw: pd.DataFrame) -> None:
    """回读 checked CSV，确认它就是「源文件里被保留的那些行」，一字不改。

    ``expected_raw`` 必须是**源文件原文**（未去空格、未转大小写）：
    输出是逐行照搬的，所以能且只能与原文比。早先误用规范化后的 DataFrame 来比，
    结果源文件里 Record Type 写成小写 `d`、或字段两侧带空格时，
    合法输入会报 `output_verify_failed`（内容其实没写错）。
    """
    actual = _read_csv(path)
    if list(actual.columns) != original_columns:
        raise StockCheckError("checked CSV 列顺序校验失败", code="output_verify_failed")
    if actual.fillna("").astype(str).values.tolist() != expected_raw.fillna("").astype(str).values.tolist():
        raise StockCheckError("checked CSV 回读内容校验失败", code="output_verify_failed")

    # 这两条用规范化后的值判断：源文件里大小写/空白怎么写都不影响结论
    record = actual[COL_RECORD].astype(str).str.strip().str.upper()
    details = actual[record == "D"]
    headers = actual[record == "H"]
    if set(details[COL_PO]) != set(headers[COL_PO]):
        raise StockCheckError("checked CSV 的 Header / Detail 对账失败", code="output_verify_failed")
    if (headers.groupby(COL_PO).size() != 1).any():
        raise StockCheckError("checked CSV 有 PO 不是恰好一个 Header", code="output_verify_failed")


def run_stock_check(
    csv_path,
    no_stock_skus: list[str],
    output_dir,
    progress: ProgressFn | None = None,
    name_stem: str | None = None,
) -> StockCheckResult:
    """检查 SPS New 订单 CSV，生成 checked CSV、操作 Excel 与结构化报告。

    ``name_stem`` 用于输出文件名：网页上传时磁盘名是固定的 ``order.csv``，
    得把用户原始文件名（如 ``check0stock order x21 …``）传进来，
    产物才对得上 SPS 里的那一批。缺省时退回用输入文件的词干。
    """
    progress = progress or _noop
    csv_path = Path(csv_path).resolve()
    output_dir = Path(output_dir).resolve()
    if not csv_path.is_file():
        raise StockCheckError("SPS 订单 CSV 不存在", code="input_missing")

    progress("read", "读取并校验 SPS New 订单 CSV")
    source = _read_csv(csv_path)
    original_columns = list(source.columns)
    work = _validate(source)
    snapshot, wanted = _normalized_skus(no_stock_skus)
    classified, po_status = _classify(work, wanted)
    checked = _checked_rows(classified, po_status)

    details = classified[classified[COL_RECORD] == "D"]
    in_stock = details[~details["_is_no_stock"]]
    no_stock = details[details["_is_no_stock"]]
    status_counts = pd.Series(po_status).value_counts()
    report = {
        "input": {"name": csv_path.name, "rows": len(source), "columns": len(source.columns)},
        "po": {
            "total": len(po_status),
            "fully_in_stock": int(status_counts.get("全部有货", 0)),
            "partially_out_of_stock": int(status_counts.get("部分缺货", 0)),
            "fully_out_of_stock": int(status_counts.get("全部缺货", 0)),
            "retained": int(checked.loc[checked[COL_RECORD] == "H", COL_PO].nunique()),
        },
        "details": {
            "total": len(details),
            "in_stock": len(in_stock),
            "out_of_stock": len(no_stock),
            "retained": int((checked[COL_RECORD] == "D").sum()),
            "diff": len(details) - len(in_stock) - len(no_stock),
        },
        "quantity": {
            "total": int(details["_qty"].sum()),
            "in_stock": int(in_stock["_qty"].sum()),
            "out_of_stock": int(no_stock["_qty"].sum()),
            "diff": int(details["_qty"].sum() - in_stock["_qty"].sum() - no_stock["_qty"].sum()),
        },
        "mixed_po_count": int(status_counts.get("部分缺货", 0)),
        "mixed_po": [po for po, status in po_status.items() if status == "部分缺货"],
        "no_stock_snapshot": snapshot,
        "outputs": [],
    }

    if report["details"]["diff"] or report["quantity"]["diff"]:
        raise StockCheckError("库存分类数量对账失败", code="reconciliation_failed")

    # 操作表只构建一次：xlsx 与网页明细表用的是同一批 DataFrame，不会各说各话
    tables = _build_tables(classified, report)
    report["tables"] = _tables_payload(tables, WEB_TABLE_LIMIT)

    progress("write", "生成 checked CSV 与 SPS 操作 Excel")
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    temp_root = Path(tempfile.mkdtemp(prefix="pb-stock-check-", dir=output_dir.parent))
    try:
        stem = _with_po_count(
            _output_stem(name_stem if name_stem is not None else csv_path.stem),
            report["po"]["retained"],
        )
        checked_name = f"checked0stock {stem}.csv"
        workbook_name = f"SPS库存检查操作表-{stem}.xlsx"
        temp_checked = temp_root / checked_name
        temp_workbook = temp_root / workbook_name
        # 逐行照搬源文件原文（含空列名、参差列数、行尾），
        # 使这份 checked CSV 与 SPS 自己导出的那份是同一种格式；UTF-8 无 BOM / LF。
        source_lines = _source_lines(csv_path, len(source))
        checked_lines = [source_lines[0]] + [source_lines[int(pos) + 1] for pos in checked.index]
        temp_checked.write_text("\n".join(checked_lines) + "\n", encoding="utf-8", newline="")
        # 与源文件原文比（不是规范化后的 DataFrame），见 _verify_checked 的说明
        _verify_checked(temp_checked, original_columns, source.loc[checked.index])

        with pd.ExcelWriter(temp_workbook, engine="openpyxl") as writer:
            for sheet_name, table in tables.items():
                _excel_safe(table).to_excel(writer, sheet_name=sheet_name[:31], index=False)
        _style_workbook(temp_workbook)

        output_dir.mkdir(parents=True, exist_ok=True)
        checked_path = output_dir / checked_name
        workbook_path = output_dir / workbook_name
        # 两个目标先统一探一遍（被 Excel 占用就整体拒绝），再逐个 `os.replace` 覆盖：
        # 避免只发布成功一个、留下「新 CSV + 旧 XLSX」这种半新半旧。
        _ensure_publishable(checked_path, workbook_path)
        replaced = [
            target.name
            for temp, target in ((temp_checked, checked_path), (temp_workbook, workbook_path))
            if _publish(temp, target)
        ]
    except Exception:
        if output_dir.is_dir() and not any(output_dir.iterdir()):
            output_dir.rmdir()
        raise
    finally:
        shutil.rmtree(temp_root, ignore_errors=True)

    report["outputs"] = [
        {"kind": "checked_order", "name": checked_path.name},
        {"kind": "stock_operations", "name": workbook_path.name},
    ]
    report["replaced"] = replaced
    progress("done", "库存预检完成")
    return StockCheckResult(checked_path, workbook_path, report)
