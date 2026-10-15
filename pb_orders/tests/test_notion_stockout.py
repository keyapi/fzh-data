#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""notion_stockout：把「SPS库存检查操作表」算成 Notion 待追加行。

全部用**合成**操作表（openpyxl 现画），不依赖客户数据。
夹具对齐真实 20261012 那批：4 条缺货明细（3×草绿153 + 1×草绿138），全部缺货。
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

import notion_stockout
from notion_stockout import NotionStockoutError, build_payload, load_known_pos, main, norm_vendor

NO_STOCK_COLUMNS = [
    "PO Number",
    "PO Line #",
    "Buyers Catalog or Stock Keeping #",
    "Vendor Style",
    "Qty Ordered",
    "库存状态",
    "PO 库存状态",
    "SPS 操作",
]

SKU_MAP = {
    "CEN961NLINEN-SAGEGREEN-153": {
        "sku": "草绿153", "tongtu_sku": "CEN961NLinen-SageGreen-153", "eta": "2026-10-30",
    },
    "CEN961NLINEN-SAGEGREEN-138": {
        "sku": "草绿138", "tongtu_sku": "CEN961NLinen-SageGreen-138", "eta": "2026-11-17",
    },
    "CENKZ1325-YELLOW-97": {
        "sku": "黄色97", "tongtu_sku": "CENKZ1325-Yellow-97", "eta": None,
    },
}

# 对齐 20261012 真实批次：PO / 行号 / 对方SKU / Vendor Style / 数量 / PO 库存状态
REAL_ROWS = [
    ("138034395", "1", "7329610", "CEN961NLINEN-SAGEGREEN-153", "1", "全部缺货"),
    ("138066286", "1", "7329610", "CEN961NLINEN-SAGEGREEN-153", "1", "全部缺货"),
    ("138067758", "1", "7329610", "CEN961NLINEN-SAGEGREEN-153", "1", "全部缺货"),
    ("138083193", "1", "8147053", "CEN961NLINEN-SAGEGREEN-138", "1", "全部缺货"),
]


def write_operation_xlsx(path, rows=None, fully_rows=None, orig_po=None, columns=None):
    """画一张与 stock_precheck 产出同构的操作表。"""
    rows = REAL_ROWS if rows is None else rows
    columns = columns or NO_STOCK_COLUMNS

    detail = [
        [po, line, buyer, vendor, qty, "缺货", po_status, "整单不生成 ASN；发送新日期通知"]
        for po, line, buyer, vendor, qty, po_status in rows
    ]
    if fully_rows is None:
        fully_rows = [r for r in detail if r[6] == "全部缺货"]

    wb = openpyxl.Workbook()
    overview = wb.active
    overview.title = "操作总览"
    overview.append(["项目", "数量", "操作说明"])
    overview.append(["原始 PO", orig_po if orig_po is not None else len({r[0] for r in detail}), ""])
    overview.append(["全部有货 PO", 0, "可直接生成 ASN"])
    overview.append(["部分缺货 PO", 0, "不要取消整单；ASN 只勾选有货明细"])
    overview.append(["全部缺货 PO", len({r[0] for r in fully_rows}), ""])

    ws = wb.create_sheet("缺货明细")
    ws.append(columns)
    for r in detail:
        ws.append(r)

    wf = wb.create_sheet("全部缺货PO")
    wf.append(columns)
    for r in fully_rows:
        wf.append(r)

    wb.save(path)
    return path


@pytest.fixture
def xlsx(tmp_path):
    # orig_po=28 对齐真实 20261012 那批（操作表「操作总览」里的原始 PO 数）
    return write_operation_xlsx(
        tmp_path / "SPS库存检查操作表-order x24 20261012_0102_199800.xlsx", orig_po=28
    )


def test_norm_vendor_ignores_case_and_spaces():
    assert norm_vendor("cen961nlinen-sagegreen-153") == "CEN961NLINEN-SAGEGREEN-153"
    assert norm_vendor(" CEN961NLinen-SageGreen-153 ") == "CEN961NLINEN-SAGEGREEN-153"
    assert norm_vendor(None) == ""


def test_build_payload_real_batch(xlsx):
    payload = build_payload(xlsx, SKU_MAP, known_pos=set())
    pages = payload["pages"]

    assert payload["summary"]["缺货明细"] == 4
    assert payload["summary"]["新增"] == 4
    assert payload["summary"]["未映射"] == []
    assert payload["summary"]["已存在跳过"] == []
    assert [p["properties"]["PO"] for p in pages] == [
        "138034395", "138066286", "138067758", "138083193",
    ]

    first = pages[0]["properties"]
    assert first["SKU"] == "草绿153"
    assert first["通途SKU"] == "CEN961NLinen-SageGreen-153"
    assert first["对方SKU"] == "7329610"
    assert first["date:到货日期:start"] == "2026-10-30"
    assert first["date:SPS操作日期:start"] == "2026-10-12"

    last = pages[-1]["properties"]
    assert last["SKU"] == "草绿138"
    assert last["date:到货日期:start"] == "2026-11-17"


def test_build_payload_ignores_vendor_style_case(xlsx, tmp_path):
    rows = [("138099999", "1", "9999999", "cen961nlinen-sagegreen-153", "1", "全部缺货")]
    path = write_operation_xlsx(tmp_path / "op 20261012.xlsx", rows=rows)
    payload = build_payload(path, SKU_MAP, known_pos=set())
    assert payload["pages"][0]["properties"]["SKU"] == "草绿153"
    assert payload["summary"]["未映射"] == []


def test_dedup_skips_known_pos(xlsx):
    payload = build_payload(xlsx, SKU_MAP, known_pos={"138034395"})
    assert payload["summary"]["新增"] == 3
    assert payload["summary"]["已存在跳过"] == ["138034395"]
    assert "138034395" not in [p["properties"]["PO"] for p in payload["pages"]]


def test_unmapped_vendor_raises(tmp_path):
    rows = [("138000001", "1", "111", "STYLE-NEW-99", "1", "全部缺货")]
    path = write_operation_xlsx(tmp_path / "op 20261012.xlsx", rows=rows)
    payload = build_payload(path, SKU_MAP, known_pos=set())
    assert payload["summary"]["未映射"] == ["138000001 / STYLE-NEW-99"]

    out = tmp_path / "rows.json"
    rc = main(["--xlsx", str(path), "--sku-map", str(_write_map(tmp_path)),
               "--known-pos", "", "--out", str(out)])
    assert rc == 2
    assert not out.exists(), "有未映射时不该产出半成品"


def test_cross_check_mismatch_raises(tmp_path):
    # 缺货明细里两条 全部缺货，但 全部缺货PO sheet 只放一条 → 对不上
    rows = [
        ("138000001", "1", "111", "CEN961NLINEN-SAGEGREEN-153", "1", "全部缺货"),
        ("138000002", "1", "222", "CEN961NLINEN-SAGEGREEN-153", "1", "全部缺货"),
    ]
    path = write_operation_xlsx(tmp_path / "op 20261012.xlsx", rows=rows)
    wb = openpyxl.load_workbook(path)
    del wb["全部缺货PO"]
    wf = wb.create_sheet("全部缺货PO")
    wf.append(NO_STOCK_COLUMNS)
    wf.append(["138000001", "1", "111", "CEN961NLINEN-SAGEGREEN-153", "1", "缺货", "全部缺货", "x"])
    wb.save(path)
    with pytest.raises(NotionStockoutError, match="对不上"):
        build_payload(path, SKU_MAP, known_pos=set())


def test_properties_encoding(xlsx):
    props = build_payload(xlsx, SKU_MAP, known_pos=set())["pages"][0]["properties"]
    assert props["通途建单"] == "__NO__"          # checkbox 用 __NO__
    assert props["打单"] == "未打单"
    assert props["状态"] == "待到货"
    assert props["invoice"] == "不适用"
    assert props["顾客"] == "正常"
    assert "date:SPS操作日期:start" in props      # 日期用展开列
    assert props["检出批次"].startswith("20261012 check0stock x28")


def test_partial_po_marked_partial(tmp_path):
    rows = [("138000003", "1", "333", "CEN961NLINEN-SAGEGREEN-153", "1", "部分缺货")]
    path = write_operation_xlsx(tmp_path / "op 20261012.xlsx", rows=rows, fully_rows=[])
    props = build_payload(path, SKU_MAP, known_pos=set())["pages"][0]["properties"]
    assert props["打单"] == "部分"


def test_missing_eta_omits_date_and_reports(tmp_path):
    rows = [("138000004", "1", "444", "CENKZ1325-YELLOW-97", "1", "全部缺货")]
    path = write_operation_xlsx(tmp_path / "op 20261012.xlsx", rows=rows)
    payload = build_payload(path, SKU_MAP, known_pos=set())
    assert "date:到货日期:start" not in payload["pages"][0]["properties"]
    assert payload["summary"]["无到货日期"] == ["138000004"]


def test_same_po_multiple_no_stock_lines_reported_once(tmp_path):
    rows = [
        ("138000005", "1", "555", "CEN961NLINEN-SAGEGREEN-153", "1", "部分缺货"),
        ("138000005", "2", "556", "CEN961NLINEN-SAGEGREEN-138", "1", "部分缺货"),
    ]
    path = write_operation_xlsx(tmp_path / "op 20261012.xlsx", rows=rows, fully_rows=[])
    payload = build_payload(path, SKU_MAP, known_pos=set())
    assert payload["summary"]["新增"] == 1
    assert payload["summary"]["同一PO多缺货行"] == ["138000005"]


def test_sps_date_override(xlsx):
    payload = build_payload(xlsx, SKU_MAP, known_pos=set(), sps_date="2026-10-14")
    assert payload["pages"][0]["properties"]["date:SPS操作日期:start"] == "2026-10-14"


def test_filename_without_date_raises(tmp_path):
    rows = [("138000006", "1", "666", "CEN961NLINEN-SAGEGREEN-153", "1", "全部缺货")]
    path = write_operation_xlsx(tmp_path / "SPS库存检查操作表-order x1.xlsx", rows=rows)
    with pytest.raises(NotionStockoutError, match="8 位日期"):
        build_payload(path, SKU_MAP, known_pos=set())


def test_main_writes_json_and_reports_dedup(xlsx, tmp_path):
    out = tmp_path / "rows.json"
    known = tmp_path / "known.json"
    known.write_text(json.dumps({"pos": ["138034395"]}), encoding="utf-8")

    rc = main(["--xlsx", str(xlsx), "--sku-map", str(_write_map(tmp_path)),
               "--known-pos", str(known), "--out", str(out)])
    assert rc == 0

    payload = json.loads(out.read_text(encoding="utf-8"))
    assert payload["parent"]["data_source_id"] == notion_stockout.NOTION_DATA_SOURCE_ID
    assert payload["summary"]["新增"] == 3
    assert payload["summary"]["已存在跳过"] == ["138034395"]
    assert len(payload["pages"]) == 3


def test_load_known_pos_missing_file_is_empty(tmp_path):
    assert load_known_pos(tmp_path / "nope.json") == set()
    assert load_known_pos(None) == set()


def test_shipped_sku_map_is_valid():
    mapping = notion_stockout.load_sku_map(notion_stockout.DEFAULT_SKU_MAP)
    assert mapping["CEN961NLINEN-SAGEGREEN-153"]["sku"] == "草绿153"
    assert mapping["CENKZ1324-SKYBLUE-97"]["tongtu_sku"] == "CENKZ1324-SkyBlue-97"
    # 到货日期允许为空（黄色97 老表里就没记），但键必须齐
    assert set(mapping) >= {
        "CENKZ1324-SKYBLUE-97", "CENKZ1325-YELLOW-97", "CENKZ1325-YELLOW-138",
        "CENKZ1325-YELLOW-194", "CEN961NLINEN-SAGEGREEN-100",
        "CEN961NLINEN-SAGEGREEN-138", "CEN961NLINEN-SAGEGREEN-153",
    }


def _write_map(tmp_path: Path) -> Path:
    path = tmp_path / "sku_map.json"
    path.write_text(json.dumps(SKU_MAP, ensure_ascii=False), encoding="utf-8")
    return path
